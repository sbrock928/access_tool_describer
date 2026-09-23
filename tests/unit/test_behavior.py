"""Regression portfolios for explainable behavior, independent of model inference."""

from pathlib import Path

import pytest

from portfolio_analyzer.models import (
    AnalysisCoverage,
    ArtifactStatus,
    Claim,
    Confidence,
    Datasource,
    ExtractedApplication,
    ExtractedObject,
    InventoryRecord,
    SemanticPortfolioState,
    StagedArtifact,
)
from portfolio_analyzer.semantic.behavior import ui_properties
from portfolio_analyzer.semantic.config import SemanticSettings
from portfolio_analyzer.semantic.pipeline import (
    _citation_validity,
    run_semantic_pipeline,
    semantic_state_is_current,
)


def obj(kind: str, name: str, definition: str = "") -> ExtractedObject:
    return ExtractedObject(object_type=kind, name=name, definition=definition)


def describe(
    objects: list[ExtractedObject],
    datasources: list[Datasource] | None = None,
    claims: list[Claim] | None = None,
    *,
    warning: bool = False,
    prior: SemanticPortfolioState | None = None,
) -> SemanticPortfolioState:
    return run_semantic_pipeline(
        SemanticSettings(),
        None,
        [
            InventoryRecord(
                tool_inventory_id="1",
                tool_name="Example",
                inventory_filename="app.accdb",
                filepath=Path("app.accdb"),
            )
        ],
        [
            StagedArtifact(
                tool_inventory_id="1",
                original_source_path=Path("app.accdb"),
                local_staged_path=Path("staged/app.accdb"),
                filename="app.accdb",
                extension=".accdb",
                sha256="a" * 64,
                status=ArtifactStatus.STAGED,
            )
        ],
        [
            (
                "a" * 64,
                ExtractedApplication(
                    tool_inventory_id="1",
                    staged_path=Path("staged/app.accdb"),
                    extractor_version="test",
                    objects=objects,
                ),
            )
        ],
        [],
        datasources or [],
        [
            AnalysisCoverage(
                tool_inventory_id="1",
                tool_name="Example",
                staging_status="staged",
                extraction_status="complete",
                analysis_status="complete",
                extraction_warning_count=int(warning),
            )
        ],
        claims or [],
        prior_state=prior,
    )


TABLE = obj("table", "Sales")
READ = obj("query", "qrySales", "SELECT * FROM Sales")
REPORT = obj("report", "rptSales", 'Begin Report\nRecordSource = "qrySales"\nEnd')
FORM = obj("form", "frmSales", 'Begin Form\nRecordSource = "Sales"\nAllowEdits = -1\nEnd')
EXPORT = obj(
    "module",
    "Exporter",
    'Sub Export()\nDoCmd.TransferSpreadsheet acExport, 10, "Sales", destination\nEnd Sub',
)
IMPORT = obj(
    "module",
    "Importer",
    'Sub Import()\nDoCmd.TransferSpreadsheet acImport, 10, "Sales", source\nEnd Sub',
)


@pytest.mark.parametrize(
    ("objects", "expected"),
    [
        ([TABLE, READ], "reporting and analytics"),
        ([obj("linked_table", "Sales"), READ], "reporting and analytics"),
        ([TABLE, REPORT], "reporting and analytics"),
        ([TABLE, FORM], "transactional workflow"),
        ([TABLE, FORM, EXPORT], "transactional workflow"),
        ([TABLE, FORM, REPORT], "mixed application"),
        ([TABLE, READ, EXPORT], "reporting and analytics"),
        ([EXPORT], "integration utility"),
        ([TABLE, READ, IMPORT], "integration utility"),
        ([TABLE, obj("query", "UpdateSales", "UPDATE Sales SET Amount = 0")], "batch automation"),
        ([obj("macro", "Nightly", 'Action = "RunSQL"')], "batch automation"),
        ([obj("form", "SaveEntryEdit")], "unknown"),
        ([obj("linked_table", "Sales")], "unknown"),
        (
            [obj("module", "Read", 'Sub Main()\nDoCmd.RunSQL "SELECT * FROM Sales"\nEnd Sub')],
            "reporting and analytics",
        ),
        (
            [obj("module", "HTTP", 'Sub Main()\nSet x = CreateObject("MSXML2.XMLHTTP")\nEnd Sub')],
            "integration utility",
        ),
    ],
)
def test_roles_follow_behavior(objects: list[ExtractedObject], expected: str) -> None:
    state = describe(objects)
    profile = state.applications[0]
    assert profile.primary_archetype == expected
    assert profile.classification_rationale
    assert profile.generation_method == "deterministic"
    assert not any(f.category == "business_capability" for f in profile.findings)
    assert _citation_validity(state) == 1
    if expected != "unknown":
        assert profile.classification_evidence_ids
    else:
        assert profile.confidence == Confidence.LOW


def test_connections_cannot_change_role_and_scope_is_explicit() -> None:
    for platform, kind, expected_scope in [
        ("Access (local)", "table", "local"),
        ("SQL Server", "linked_table", "external"),
        ("Unknown", "", "unresolved"),
    ]:
        objects = [READ] + ([obj(kind, "Sales")] if kind else [])
        state = describe(
            objects,
            [
                Datasource(
                    tool_inventory_id="1", platform=platform, object_name="Sales", operation="READ"
                )
            ],
        )
        assert state.applications[0].primary_archetype == "reporting and analytics"
        reads = [f for f in state.application_irs[0].behavior_facts if f.action == "read"]
        assert reads[0].datasource_scope == expected_scope
        assert expected_scope in state.applications[0].summary


def test_inventory_and_ui_properties_do_not_inflate_code_coverage() -> None:
    state = describe([TABLE, FORM, REPORT, READ])
    ir = state.application_irs[0]
    profile = state.applications[0]
    assert ir.inventory_object_type_counts == {"table": 1, "form": 1, "report": 1, "query": 1}
    assert ir.object_type_counts == {"query": 1}
    assert ir.code_object_count == ir.code_segment_count == 1
    assert len(ir.inventory_source_ids) == 4
    assert profile.semantic_coverage is not None
    assert profile.semantic_coverage.inventory_objects == 4
    assert profile.semantic_coverage.code_objects_inspected == 1
    assert profile.semantic_coverage.complete_code_coverage
    assert "Sales" in profile.inputs
    assert "rptSales" in profile.outputs
    assert any("configured for record entry" in text for text in profile.observed_behavior)


def test_form_metadata_reads_root_and_continuations_only() -> None:
    metadata = ui_properties(
        """Version =20
Begin Form
    RecordSource = "SELECT SalesId "
        "FROM Sales"
    AllowEdits =0
    AllowAdditions =0
    Begin Subform
        RecordSource = "Unrelated"
        DataEntry =-1
    End
End
CodeBehindForm
DataEntry = -1
""",
        {},
    )
    assert metadata == {
        "recordsource": "SELECT SalesId FROM Sales",
        "allowedits": "0",
        "allowadditions": "0",
    }
    assert (
        describe([obj("form", "Entry", "\n".join(f'{k} = "{v}"' for k, v in metadata.items()))])
        .applications[0]
        .primary_archetype
        == "reporting and analytics"
    )


def test_form_code_and_explicit_query_reference_support_updates() -> None:
    form = obj(
        "form",
        "Control",
        'Begin Form\nEnd\nCodeBehindForm\nSub Run()\nDoCmd.OpenQuery "UpdateSales"\nEnd Sub',
    )
    query = obj("query", "UpdateSales", "UPDATE Sales SET Amount = 0")
    state = describe([TABLE, form, query])
    profile = state.applications[0]
    assert profile.primary_archetype == "transactional workflow"
    refs = set(profile.classification_evidence_ids)
    assert {s.object_name for s in state.sources if s.source_id in refs} == {
        "Control",
        "UpdateSales",
    }
    assert not any("then" in text for text in profile.observed_behavior)


def test_comments_strings_and_dynamic_transfer_direction_are_not_invented() -> None:
    module = obj(
        "module",
        "Example",
        """Sub Main()
' DoCmd.OutputTo acOutputReport, "Fake"
Rem DoCmd.TransferSpreadsheet acImport, 10, "Sales"
message = "DoCmd.SendObject and Excel.Application"
End Sub""",
    )
    assert describe([module]).applications[0].primary_archetype == "unknown"
    module.definition = 'Sub Main()\nDoCmd.TransferSpreadsheet importOrExport, 10, "Sales"\nEnd Sub'
    state = describe([module])
    assert {f.action for f in state.application_irs[0].behavior_facts} == {"transfer"}
    assert "direction unresolved" in state.applications[0].summary


def test_owner_purpose_remains_a_claim_and_coverage_qualifies_summary() -> None:
    claim = Claim(
        tool_inventory_id="1",
        field="business_purpose",
        value="Prepare sales forecasts",
        source="owner interview",
    )
    state = describe([TABLE, READ, REPORT], claims=[claim], warning=True)
    profile = state.applications[0]
    assert profile.purpose_provenance == "owner_claim"
    assert profile.purpose_claim_ids == [claim.claim_id]
    assert profile.business_purpose == claim.value
    assert "Owner-stated purpose" in profile.summary
    assert "provisional" in profile.summary
    assert "provisional" in profile.classification_rationale
    assert profile.confidence == Confidence.LOW
    assert not any(f.category == "business_capability" for f in profile.findings)


def test_metadata_changes_invalidate_resume_and_prior_versions_are_stale() -> None:
    state = describe([TABLE, FORM])
    unchanged = describe([TABLE, FORM], prior=state)
    assert unchanged.applications[0].input_fingerprint == state.applications[0].input_fingerprint
    changed_form = FORM.model_copy(update={"definition": FORM.definition.replace("-1", "0")})
    changed = describe([TABLE, changed_form], prior=state)
    assert changed.applications[0].primary_archetype == "reporting and analytics"
    assert changed.application_irs[0].ir_id != state.application_irs[0].ir_id
    state.metadata.semantic_version = "semantic-analysis-v7"
    assert not semantic_state_is_current(state, SemanticSettings(), [], [])


def test_known_transfer_targets_and_secondary_capabilities() -> None:
    profile = describe([TABLE, FORM, EXPORT]).applications[0]
    assert any(
        "Exports spreadsheet data (Sales)" in text for text in profile.secondary_capabilities
    )
    assert "Sales" in profile.inputs
    assert "Exports spreadsheet data (Sales)" in profile.outputs
    imported = describe([IMPORT]).applications[0]
    assert "Sales" in imported.outputs
    assert "Sales" not in imported.inputs


def test_reviewed_empty_capabilities_are_valid_gold_labels() -> None:
    from portfolio_analyzer.semantic.pipeline import evaluate_gold_set

    state = describe([TABLE, READ])
    inventory = [
        InventoryRecord(
            tool_inventory_id="1",
            tool_name="Example",
            inventory_filename="app.accdb",
            filepath=Path("app.accdb"),
        )
    ]
    result = evaluate_gold_set(
        [
            {
                "euc_name": "Example",
                "expected_primary_archetype": "reporting and analytics",
                "expected_business_capabilities": "(none)",
            }
        ],
        inventory,
        state,
    )
    assert result["passed"]
    assert result["capability_macro_f1"] == 1.0


def test_segmented_form_code_retains_metadata_and_all_citations() -> None:
    definition = (
        'Begin Form\nRecordSource = "Sales"\nAllowEdits = -1\nEnd\nCodeBehindForm\n'
        + "\n".join(f"Sub Handler{i}()\nMe.Recordset.Update\nEnd Sub" for i in range(4))
    )
    state = describe([TABLE, obj("form", "SalesForm", definition)])
    ir = state.application_irs[0]
    assert ir.code_segment_count == 4
    assert ir.inventory_object_type_counts == {"table": 1, "form": 1}
    assert state.applications[0].primary_archetype == "transactional workflow"
    assert _citation_validity(state) == 1
    assert len(state.applications[0].classification_evidence_ids) == 4


def test_aggregate_ir_is_not_a_second_independent_confidence_source() -> None:
    state = describe([TABLE, READ])
    assert state.applications[0].confidence == Confidence.MEDIUM


@pytest.mark.parametrize(
    ("objects", "expected"),
    [
        ([TABLE, REPORT, obj("query", "Refresh", "UPDATE Sales SET Amount = 0")],
         {"reporting and analytics", "batch automation"}),
        ([TABLE, FORM, REPORT, EXPORT,
          obj("query", "Refresh", "UPDATE Sales SET Amount = 0")],
         {"transactional workflow", "reporting and analytics", "batch automation", "integration"}),
        ([TABLE, obj("query", "Append", "INSERT INTO Sales SELECT * FROM Archive")],
         {"batch automation"}),
        ([TABLE, FORM, EXPORT], {"transactional workflow", "integration"}),
        ([obj("linked_table", "Sales")], set()),
        ([TABLE, obj("form", "OnlyEntry", 'Begin Form\nRecordSource = "Sales"\n'
                     'AllowEdits = -1\nEnd\nCodeBehindForm\nSub Run()\n'
                     'DoCmd.OpenQuery "Refresh"\nEnd Sub'),
          obj("query", "Refresh", "UPDATE Sales SET Amount = 0")],
         {"transactional workflow"}),
    ],
)
def test_independent_roles_preserve_overlap_without_inventing_roles(
    objects: list[ExtractedObject], expected: set[str],
) -> None:
    state = describe(objects)
    assert {r.role for r in state.applications[0].roles} == expected
    assert all(r.evidence_ids for r in state.applications[0].roles)
    assert _citation_validity(state) == 1.0
    # Roles describe this singleton; they cannot create a shared modernization group.
    assert state.architecture.mappings[0].target_component_ids == []
    assert state.discovered_themes == []


def test_themes_locate_shared_objects_and_qualify_incomplete_evidence() -> None:
    from test_discovery import discovered_portfolio

    from portfolio_analyzer.portfolio.themes import build_portfolio_themes

    state, _ = discovered_portfolio(warning=True)
    themes = build_portfolio_themes(state)
    assert len(themes) == 2
    assert all(t.confidence == Confidence.LOW and "Provisional" in t.coverage_note for t in themes)
    assert any(loc.object_name == "Refresh1" for t in themes for loc in t.locations)
    assert any(loc.object_name == "ReadClaims" for t in themes for loc in t.locations)
    source_ids = {s.source_id for s in state.sources}
    assert all(loc.evidence_id in source_ids for t in themes for loc in t.locations)


def test_role_review_accepts_overlapping_labels_without_legacy_bucket() -> None:
    from portfolio_analyzer.semantic.pipeline import evaluate_gold_set

    state = describe([TABLE, REPORT, obj("query", "Refresh", "UPDATE Sales SET Amount = 0")])
    inventory = [InventoryRecord(tool_inventory_id="1", tool_name="Example",
                                 inventory_filename="app.accdb", filepath=Path("app.accdb"))]
    rows = [{"euc_name": "Example", "expected_roles": "reporting and analytics|batch automation",
             "expected_business_capabilities": "(none)"}]
    result = evaluate_gold_set(rows, inventory, state)
    assert result["passed"]
    assert result["role_macro_f1"] == 1.0
    rows[0]["expected_roles"] = "reporting and analytics"
    assert not evaluate_gold_set(rows, inventory, state)["passed"]


def test_entry_bound_query_is_not_independent_analytics() -> None:
    form = obj("form", "Entry", 'Begin Form\nRecordSource = "qrySales"\nAllowEdits = -1\nEnd')
    state = describe([TABLE, READ, form])
    assert {r.role for r in state.applications[0].roles} == {"transactional workflow"}


def test_themes_and_overlapping_roles_are_consistent_across_reports(tmp_path: Path) -> None:
    import csv
    import json
    import re

    from openpyxl import load_workbook
    from test_discovery import discovered_portfolio

    from portfolio_analyzer.portfolio.themes import (
        THEME_HEADERS,
        THEME_LOCATION_HEADERS,
        build_portfolio_themes,
        theme_location_rows,
        theme_rows,
    )
    from portfolio_analyzer.reporting.intelligence import (
        write_intelligence_html,
        write_semantic_datasets,
    )
    from portfolio_analyzer.reporting.writers import write_csv, write_executive_pdf, write_workbook

    state, inventory = discovered_portfolio()
    themes = build_portfolio_themes(state)
    names = {i.tool_inventory_id: i.tool_name for i in inventory}
    html_path = tmp_path / "Portfolio_Intelligence.html"
    write_intelligence_html(
        html_path, inventory, [], state, semantic_status="current", themes=themes,
    )
    html = html_path.read_text()
    payload = re.search(r'<script id="portfolio-data" type="application/json">(.*?)</script>', html)
    assert payload
    data = json.loads(payload[1])
    assert {r["role"] for r in data["applications"][0]["roles"]} == {
        "reporting and analytics", "batch automation",
    }
    assert {t["title"] for t in data["themes"]} == {t.title for t in themes}
    xlsx = tmp_path / "Portfolio_Analysis.xlsx"
    write_workbook(xlsx, inventory, [], [], [], [], [], semantic=state, themes=themes)
    workbook = load_workbook(xlsx)
    assert workbook["Portfolio Themes"].max_row == 3
    locations = list(workbook["Theme Locations"].values)
    assert any("Refresh1" in row for row in locations)
    assert any("ReadClaims" in row for row in locations)
    app_rows = list(workbook["Application Portfolio"].values)
    roles_column = app_rows[0].index("Supported Roles")
    assert "batch automation" in app_rows[1][roles_column]
    assert "reporting and analytics" in app_rows[1][roles_column]
    write_semantic_datasets(tmp_path, state, names)
    with (tmp_path / "semantic_applications.csv").open() as handle:
        row = next(csv.DictReader(handle))
    assert row["supported_roles"] == app_rows[1][roles_column]
    for filename, headers, rows in (
        ("portfolio_themes.csv", THEME_HEADERS, theme_rows(themes, names)),
        ("theme_locations.csv", THEME_LOCATION_HEADERS, theme_location_rows(themes, names)),
    ):
        write_csv(tmp_path / filename, rows, headers=headers)
        with (tmp_path / filename).open() as handle:
            assert list(csv.DictReader(handle)) == rows
    write_executive_pdf(tmp_path / "Portfolio_Analysis.pdf", inventory, [], [],
                        semantic=state, themes=themes)
    assert (tmp_path / "Portfolio_Analysis.pdf").read_bytes().startswith(b"%PDF")
