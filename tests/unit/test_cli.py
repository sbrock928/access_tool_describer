import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from typer.testing import CliRunner

from portfolio_analyzer.cli import main as cli_main
from portfolio_analyzer.cli.main import app
from portfolio_analyzer.models import ExtractedApplication


def test_stage_and_report_keep_source_and_staged_paths_separate(tmp_path: Path) -> None:
    source = tmp_path / "source" / "tool.accdb"
    source.parent.mkdir()
    source.write_bytes(b"synthetic database")
    inventory = tmp_path / "inventory.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["INVENTORY_ID", "EUCTNAME", "FILE_NAME", "FULLPATH", "DESCRIPTION"])
    sheet.append(["42", "Tool", "tool.accdb", str(source), "Claim"])
    workbook.save(inventory)
    workspace = tmp_path / "workspace"
    runner = CliRunner()

    staged = runner.invoke(
        app, ["stage", "--inventory", str(inventory), "--workspace", str(workspace)]
    )
    reported = runner.invoke(app, ["report", "--workspace", str(workspace)])

    assert staged.exit_code == 0, staged.output
    assert reported.exit_code == 0, reported.output
    state = json.loads((workspace / "analysis" / "staging_state.json").read_text())
    primary = next(item for item in state["artifacts"] if item["is_primary"])
    assert Path(primary["original_source_path"]) == source
    assert Path(primary["local_staged_path"]).is_relative_to(workspace / "staged_tools")
    assert Path(primary["local_staged_path"]).relative_to(workspace / "staged_tools").parts[0] == (
        "Tool"
    )
    assert (workspace / "reports" / "Portfolio_Analysis.xlsx").exists()
    assert (workspace / "reports" / "applications.csv").exists()
    assert (workspace / "reports" / "analysis_coverage.csv").exists()
    assert (workspace / "reports" / "evidence.csv").exists()
    applications_csv = (workspace / "reports" / "applications.csv").read_text()
    assert applications_csv.startswith("euc_name,")
    assert "tool_inventory_id" not in applications_csv
    report = load_workbook(workspace / "reports" / "Portfolio_Analysis.xlsx", read_only=True)
    applications_sheet = report["Applications"]
    assert applications_sheet["A1"].value == "EUC Name"
    assert applications_sheet["A2"].value == "Tool"


def test_extract_uses_euc_name_for_output_directory(tmp_path: Path, monkeypatch) -> None:
    source = tmp_path / "source" / "payments.accdb"
    source.parent.mkdir()
    source.write_bytes(b"synthetic database")
    inventory = tmp_path / "inventory.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["INVENTORY_ID", "EUCTNAME", "FILE_NAME", "FULLPATH", "DESCRIPTION"])
    sheet.append(["42", "Payments EUC", source.name, str(source), "Claim"])
    workbook.save(inventory)
    workspace = tmp_path / "workspace"
    runner = CliRunner()
    staged = runner.invoke(
        app, ["stage", "--inventory", str(inventory), "--workspace", str(workspace)]
    )
    assert staged.exit_code == 0, staged.output

    def fake_extract(artifact, destination, settings, timeout_seconds, on_progress):
        destination.mkdir(parents=True, exist_ok=True)
        return ExtractedApplication(
            tool_inventory_id=artifact.tool_inventory_id,
            staged_path=artifact.local_staged_path,
            extractor_version="fixture",
        )

    monkeypatch.setattr(cli_main.platform, "system", lambda: "Windows")
    monkeypatch.setattr(cli_main, "_extract_with_timeout", fake_extract)

    extracted = runner.invoke(app, ["extract", "--workspace", str(workspace)])

    assert extracted.exit_code == 0, extracted.output
    assert (workspace / "extracted" / "Payments EUC").is_dir()
    assert "[Payments EUC] completed." in extracted.output
