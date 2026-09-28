import json
import shutil
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook
from typer.testing import CliRunner

import portfolio_analyzer.cli.v2 as v2_cli
from portfolio_analyzer.access.windows_extractor import WindowsAccessExtractor
from portfolio_analyzer.cli.v2 import _bundle_matches_stage_application, app
from portfolio_analyzer.config import AnalyzerSettings
from portfolio_analyzer.models import AccessExtractedObject, AccessExtractionResult
from portfolio_analyzer.qwen.pipeline import (
    ApplicationAnalysisResult,
    InterpretationRunStatus,
    PortfolioAnalysisResult,
    PortfolioRunStatus,
)
from portfolio_analyzer.staging.hashing import sha256_file
from portfolio_analyzer.v2.analysis_state import (
    ApplicationAnalysisPayload,
    PortfolioAnalysisPayload,
)
from portfolio_analyzer.v2.fingerprints import (
    evidence_bundle_fingerprint,
    interpretation_fingerprint,
)
from portfolio_analyzer.v2.models import (
    ApplicationEvidenceBundle,
    ApplicationInterpretation,
    ModelProvenance,
    PortfolioAnalysis,
    PortfolioCandidateType,
    PortfolioReportModel,
    ReportOmissionStage,
    ResolutionStatus,
)
from portfolio_analyzer.v2.reporting import (
    ReportPublicationManifest,
    load_report_publication,
)
from portfolio_analyzer.v2.review import ReviewOverlay
from portfolio_analyzer.v2.state import (
    RunPhase,
    RunStatus,
    StateIntegrityError,
    V2StateStore,
)
from portfolio_analyzer.v2.workflow import (
    EXTRACTION_POLICY_VERSION,
    ExtractedArtifactSnapshot,
    StageIndex,
)


def _inventory(path: Path, rows: list[tuple[str, str, Path]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["INVENTORY_ID", "EUCTNAME", "FILE_NAME", "FULLPATH", "DESCRIPTION"])
    for application_id, name, source in rows:
        sheet.append([application_id, name, source.name, str(source), f"Owner claim for {name}"])
    workbook.save(path)


def test_v2_stage_records_unsupported_rows_without_copying_them(tmp_path: Path) -> None:
    source = tmp_path / "source"
    source.mkdir()
    access = source / "Main.ACCDB"
    unsupported = source / "input.xlsx"
    access.write_bytes(b"access")
    unsupported.write_bytes(b"unsupported-canary")
    inventory = tmp_path / "inventory.xlsx"
    _inventory(
        inventory,
        [("app-1", "Example", access), ("app-2", "Spreadsheet", unsupported)],
    )
    workspace = tmp_path / "workspace"

    result = CliRunner().invoke(
        app, ["stage", "--inventory", str(inventory), "--workspace", str(workspace)]
    )

    assert result.exit_code == 0, result.output
    store = V2StateStore(workspace)
    manifest = store.load_current_run(RunPhase.STAGE)
    index = store.load(manifest.input_references[0], StageIndex)
    assert manifest.status == RunStatus.COMPLETE
    assert len(index.artifacts) == 1
    assert index.artifacts[0].access_format == "accdb"
    assert len(index.exclusions) == 1
    assert list((workspace / "staged_tools").rglob("*.xlsx")) == []
    assert unsupported.read_bytes() == b"unsupported-canary"
    generated_json = b"\n".join(
        path.read_bytes()
        for path in (workspace / ".portfolio_analyzer_v2").rglob("*.json")
    )
    assert str(access).encode() not in generated_json
    assert str(unsupported).encode() not in generated_json


def test_v2_stage_all_excluded_is_published_but_exits_nonzero(tmp_path: Path) -> None:
    source = tmp_path / "source.csv"
    source.write_text("a,b\n", encoding="utf-8")
    inventory = tmp_path / "inventory.xlsx"
    _inventory(inventory, [("app-1", "CSV", source)])
    workspace = tmp_path / "workspace"

    result = CliRunner().invoke(
        app, ["stage", "--inventory", str(inventory), "--workspace", str(workspace)]
    )

    assert result.exit_code == 1
    manifest = V2StateStore(workspace).load_current_run(RunPhase.STAGE)
    assert manifest.status == RunStatus.PARTIAL
    assert manifest.applications == ()
    assert not list((workspace / "staged_tools").rglob("*.csv"))


def test_v2_stage_failure_persists_no_original_source_path(tmp_path: Path) -> None:
    missing = tmp_path / "private-source" / "missing.accdb"
    inventory = tmp_path / "inventory.xlsx"
    _inventory(inventory, [("app-1", "Missing", missing)])
    workspace = tmp_path / "workspace"

    result = CliRunner().invoke(
        app, ["stage", "--inventory", str(inventory), "--workspace", str(workspace)]
    )

    assert result.exit_code == 1
    generated_json = b"\n".join(
        path.read_bytes()
        for path in (workspace / ".portfolio_analyzer_v2").rglob("*.json")
    )
    assert str(missing).encode() not in generated_json


def test_v2_extract_uses_staged_paths_and_targeted_force_preserves_other_apps(
    tmp_path: Path, monkeypatch: Any
) -> None:
    source = tmp_path / "source"
    source.mkdir()
    first = source / "first.accdb"
    second = source / "second.mdb"
    first.write_bytes(b"first")
    second.write_bytes(b"second")
    inventory = tmp_path / "inventory.xlsx"
    _inventory(
        inventory,
        [("app-1", "First", first), ("app-2", "Second", second)],
    )
    workspace = tmp_path / "workspace"
    runner = CliRunner()
    staged = runner.invoke(
        app, ["stage", "--inventory", str(inventory), "--workspace", str(workspace)]
    )
    assert staged.exit_code == 0, staged.output
    calls: list[Path] = []

    def fake_extract(
        artifact: Any,
        destination: Path,
        _settings: Any,
        **_kwargs: Any,
    ) -> AccessExtractionResult:
        calls.append(artifact.local_staged_path)
        assert not hasattr(artifact, "original_source_path")
        return AccessExtractionResult(
            tool_inventory_id=artifact.tool_inventory_id,
            artifact_id=artifact.artifact_id,
            staged_path=artifact.local_staged_path,
            extractor_version=WindowsAccessExtractor.version,
            objects=[
                AccessExtractedObject(
                    object_type="query",
                    name="qRead",
                    definition="SELECT * FROM LocalTable",
                    properties={
                        "dao_type_code": "0",
                        "query_kind": "select",
                        "connect": "",
                    },
                )
            ],
        )

    monkeypatch.setattr("portfolio_analyzer.cli.v2.platform.system", lambda: "Windows")
    monkeypatch.setattr("portfolio_analyzer.cli.v2.run_extraction_with_timeout", fake_extract)
    extracted = runner.invoke(app, ["extract", "--workspace", str(workspace)])
    assert extracted.exit_code == 0, extracted.output
    assert len(calls) == 2
    store = V2StateStore(workspace)
    first_run = store.load_current_run(RunPhase.EXTRACT)
    preserved = {
        item.application_id: item.extraction_snapshots for item in first_run.applications
    }
    stage_run = store.load_current_run(RunPhase.STAGE)
    stage_index = store.load(stage_run.input_references[0], StageIndex)
    assert set(
        v2_cli._compatible_previous_extractions(store, first_run, stage_index)
    ) == {"app-1", "app-2"}
    settings = AnalyzerSettings(workspace=workspace)
    library = settings.shared_libraries_dir / "reviewed.accdb"
    library.write_bytes(b"reviewed library")
    library_manifest = settings.source_inventory_dir / "access_libraries.json"
    library_manifest.write_text(
        json.dumps(
            {
                "schema_version": "access-library-manifest-v1",
                "applications": {
                    "app-1": [
                        {
                            "path": library.name,
                            "sha256": sha256_file(library),
                            "reason": "Reviewed VBA dependency",
                        }
                    ]
                },
            }
        ),
        encoding="utf-8",
    )
    assert set(
        v2_cli._compatible_previous_extractions(store, first_run, stage_index)
    ) == {"app-2"}
    library_manifest.unlink()
    incomplete_previous = first_run.model_copy(
        update={
            "applications": tuple(
                record.model_copy(update={"status": RunStatus.PARTIAL})
                for record in first_run.applications
            )
        }
    )
    assert not v2_cli._compatible_previous_extractions(
        store, incomplete_previous, stage_index
    )
    with monkeypatch.context() as policy_change:
        policy_change.setattr(
            v2_cli,
            "EXTRACTION_POLICY_VERSION",
            "access-static-extraction-policy-next",
        )
        assert not v2_cli._compatible_previous_extractions(
            store, first_run, stage_index
        )
    with monkeypatch.context() as extractor_change:
        extractor_change.setattr(
            WindowsAccessExtractor,
            "version",
            "windows-dao-static-next",
        )
        assert not v2_cli._compatible_previous_extractions(
            store, first_run, stage_index
        )

    targeted = runner.invoke(
        app,
        [
            "extract",
            "--workspace",
            str(workspace),
            "--application",
            "app-1",
            "--force",
        ],
    )

    assert targeted.exit_code == 0, targeted.output
    assert len(calls) == 3
    current = store.load_current_run(RunPhase.EXTRACT)
    records = {item.application_id: item for item in current.applications}
    assert records["app-1"].extraction_snapshots != preserved["app-1"]
    assert records["app-2"].extraction_snapshots == preserved["app-2"]
    for record in records.values():
        assert record.status == RunStatus.COMPLETE
        for reference in record.extraction_snapshots:
            snapshot = store.load(reference, ExtractedArtifactSnapshot)
            assert snapshot.application_id == record.application_id
            assert snapshot.extractor_version == WindowsAccessExtractor.version
            assert snapshot.extraction_policy_version == EXTRACTION_POLICY_VERSION


def test_v2_analyze_publishes_evidence_application_and_portfolio_chain(
    tmp_path: Path, monkeypatch: Any
) -> None:
    source = tmp_path / "source.accdb"
    second_source = tmp_path / "source-two.accdb"
    excluded_source = tmp_path / "excluded.xlsx"
    source.write_bytes(b"access-one")
    second_source.write_bytes(b"access-two")
    excluded_source.write_bytes(b"unsupported-primary")
    inventory = tmp_path / "inventory.xlsx"
    _inventory(
        inventory,
        [
            ("app-1", "Example", source),
            ("app-2", "Second", second_source),
            ("app-x", "Excluded", excluded_source),
        ],
    )
    workspace = tmp_path / "workspace"
    runner = CliRunner()
    assert runner.invoke(
        app, ["stage", "--inventory", str(inventory), "--workspace", str(workspace)]
    ).exit_code == 0

    def fake_extract(
        artifact: Any,
        _destination: Path,
        _settings: Any,
        **_kwargs: Any,
    ) -> AccessExtractionResult:
        connection = (
            "ODBC;DRIVER={ODBC Driver 17 for SQL Server};SERVER=SQL01;"
            "DATABASE=CorporateTrust;PWD=SECRET"
            if artifact.tool_inventory_id == "app-1"
            else "ODBC;DSN=MissingPortfolioDsn;PWD=SECRET"
        )
        return AccessExtractionResult(
            tool_inventory_id=artifact.tool_inventory_id,
            artifact_id=artifact.artifact_id,
            staged_path=artifact.local_staged_path,
            extractor_version=WindowsAccessExtractor.version,
            objects=[
                AccessExtractedObject(
                    object_type="query",
                    name="qRead",
                    definition="SELECT * FROM dbo.Deal",
                    properties={
                        "dao_type_code": "112",
                        "query_kind": "pass_through",
                        "connect": connection,
                        "returns_records": "true",
                        "odbc_timeout": "60",
                        "attributes": "0",
                        "parameters": "[]",
                    },
                ),
                AccessExtractedObject(
                    object_type="module",
                    name="modSharedExport",
                    definition=(
                        "Public Sub ExportDeals()\n"
                        'DoCmd.TransferSpreadsheet acExport, , "qRead", '
                        '"C:\\\\Shared\\\\deals.xlsx"\n'
                        "End Sub\n"
                    ),
                ),
            ],
        )

    monkeypatch.setattr("portfolio_analyzer.cli.v2.platform.system", lambda: "Windows")
    monkeypatch.setattr("portfolio_analyzer.cli.v2.run_extraction_with_timeout", fake_extract)
    assert runner.invoke(app, ["extract", "--workspace", str(workspace)]).exit_code == 0

    provenance = ModelProvenance(
        model_manifest_sha256="a" * 64,
        prompt_version="test-prompts-v2",
        output_schema_version="test-contracts-v2",
        inference_library_version="test-transformers",
    )
    monkeypatch.setattr(
        "portfolio_analyzer.cli.v2.resolve_qwen_runtime", lambda *_a, **_k: object()
    )
    monkeypatch.setattr(
        "portfolio_analyzer.cli.v2.LocalQwenProvider",
        lambda _runtime, **_kwargs: object(),
    )
    monkeypatch.setattr("portfolio_analyzer.cli.v2._model_provenance", lambda _provider: provenance)

    def fake_application(bundle: Any, _provider: Any, **_kwargs: Any) -> ApplicationAnalysisResult:
        profile = ApplicationInterpretation(
            application_id=bundle.application_id,
            source_bundle_sha256=evidence_bundle_fingerprint(bundle),
            logical_unit_interpretation_ids=(),
            summary="Reads a local table.",
            business_purpose="Unknown without owner review.",
            evidence_ids=(bundle.evidence[0].evidence_id,),
            provenance=provenance,
        )
        return ApplicationAnalysisResult(
            application_id=bundle.application_id,
            source_bundle_sha256=evidence_bundle_fingerprint(bundle),
            status=InterpretationRunStatus.COMPLETE,
            unit_results=(),
            interpretation=profile,
        )

    def fake_portfolio(
        _candidates: Any,
        profiles: Any,
        _bundles: Any,
        _provider: Any,
        **_kwargs: Any,
    ) -> PortfolioAnalysisResult:
        analysis = PortfolioAnalysis(
            application_interpretation_sha256s=tuple(
                interpretation_fingerprint(item) for item in profiles
            ),
            provenance=provenance,
        )
        return PortfolioAnalysisResult(
            status=PortfolioRunStatus.COMPLETE,
            analysis=analysis,
            batch_count=1,
        )

    monkeypatch.setattr(
        "portfolio_analyzer.cli.v2.analyze_application_two_stage", fake_application
    )
    monkeypatch.setattr(
        "portfolio_analyzer.cli.v2.analyze_portfolio_candidates", fake_portfolio
    )
    result = runner.invoke(
        app,
        [
            "analyze",
            "--workspace",
            str(workspace),
            "--model-dir",
            str(workspace),
            "--verbose",
        ],
    )

    assert result.exit_code == 0, result.output
    assert "Starting application 1/2" in result.output
    assert "evidence ready" in result.output
    assert "Portfolio candidates ready" in result.output
    store = V2StateStore(workspace)
    manifest = store.load_current_run(RunPhase.ANALYZE)
    assert manifest.status == RunStatus.COMPLETE
    assert v2_cli._analysis_manifest_uses_current_model(store, manifest)
    assert manifest.portfolio_analysis is not None
    assert len(manifest.applications) == 2
    stage_manifest = store.load_current_run(RunPhase.STAGE)
    stage_index = store.load(stage_manifest.input_references[0], StageIndex)
    assert {item.application_id for item in stage_index.exclusions} == {"app-x"}
    record = manifest.applications[0]
    assert record.status == RunStatus.COMPLETE
    assert record.evidence_bundle is not None
    assert record.interpretation is not None
    payload = store.load(record.interpretation, ApplicationAnalysisPayload)
    assert payload.interpretation is not None
    historical_provenance = provenance.model_copy(
        update={
            "model_repo_id": "Qwen/Qwen2.5-1.5B-Instruct",
            "model_revision": "989aa7980e4cf806f80c7fef2b1adb7bc71aa306",
        }
    )
    historical_profile = ApplicationInterpretation.model_validate(
        {
            **payload.interpretation.model_dump(mode="python"),
            "provenance": historical_provenance,
            "interpretation_id": "",
        }
    )
    historical_payload = payload.model_copy(
        update={"interpretation": historical_profile}
    )
    historical_record = record.model_copy(
        update={"interpretation": store.put(historical_payload)}
    )
    historical_manifest = manifest.model_copy(
        update={
            "applications": (
                historical_record,
                *manifest.applications[1:],
            )
        }
    )
    assert not v2_cli._analysis_manifest_uses_current_model(store, historical_manifest)
    with pytest.raises(StateIntegrityError, match="obsolete model policy"):
        v2_cli._load_complete_analysis(store, historical_manifest)
    portfolio = store.load(manifest.portfolio_analysis, PortfolioAnalysisPayload)
    assert portfolio.analysis is not None
    candidate_types = {item.candidate_type for item in portfolio.candidates}
    assert PortfolioCandidateType.EXACT_CODE in candidate_types
    assert PortfolioCandidateType.SHARED_FILE in candidate_types
    bundles = tuple(
        store.load(item.evidence_bundle, ApplicationEvidenceBundle)
        for item in manifest.applications
        if item.evidence_bundle is not None
    )
    staged_applications = {
        item.application_id: item for item in stage_index.applications
    }
    first_bundle = next(item for item in bundles if item.application_id == "app-1")
    first_staged_application = staged_applications["app-1"]
    assert _bundle_matches_stage_application(
        first_bundle,
        first_staged_application,
        stage_index,
    )
    assert not _bundle_matches_stage_application(
        first_bundle,
        first_staged_application.model_copy(
            update={"inventory_descriptions": ("changed owner context",)}
        ),
        stage_index,
    )
    assert any(
        connection.resolution_status == ResolutionStatus.UNRESOLVED
        for bundle in bundles
        for connection in bundle.connections
    )
    persisted = b"\n".join(
        path.read_bytes()
        for path in (workspace / ".portfolio_analyzer_v2").rglob("*.json")
    )
    assert b"SECRET" not in persisted

    reported = runner.invoke(app, ["report", "--workspace", str(workspace)])
    assert reported.exit_code == 0, reported.output
    report_manifest = store.load_current_run(RunPhase.REPORT)
    assert report_manifest.status == RunStatus.COMPLETE
    assert report_manifest.parent_run_id == manifest.run_id
    assert report_manifest.report_publication is not None
    assert report_manifest.report_model is not None
    persisted_report = store.load(report_manifest.report_model, PortfolioReportModel)
    assert persisted_report.absence_claims_suppressed
    assert any(
        "semantic-profile similarity candidates are disabled" in warning.casefold()
        for warning in persisted_report.warnings
    )
    assert set(manifest.input_references).issubset(report_manifest.input_references)
    state_publication = store.load(
        report_manifest.report_publication,
        ReportPublicationManifest,
    )
    assert state_publication == load_report_publication(
        workspace / "reports",
        report_manifest.run_id,
    )
    run_directory = workspace / "reports" / "runs" / report_manifest.run_id
    assert (run_directory / "portfolio.html").is_file()
    assert (run_directory / "portfolio.xlsx").is_file()
    assert (run_directory / "portfolio.pdf").is_file()
    assert (run_directory / "portfolio.csv.zip").is_file()
    latest_path = workspace / "reports" / "latest.json"
    assert latest_path.is_file()
    latest_complete = latest_path.read_bytes()

    review_workbook_path = tmp_path / "review.xlsx"
    shutil.copy2(run_directory / "portfolio.xlsx", review_workbook_path)
    review_workbook = load_workbook(review_workbook_path)
    review_queue = review_workbook["Review Queue"]
    assert review_queue.max_row > 1
    review_queue["D2"] = "accept"
    review_queue["F2"] = "portfolio-owner"
    review_workbook.save(review_workbook_path)
    review_workbook.close()
    imported_review = runner.invoke(
        app,
        [
            "import-review",
            "--workspace",
            str(workspace),
            "--workbook",
            str(review_workbook_path),
        ],
    )
    assert imported_review.exit_code == 0, imported_review.output
    reviewed_manifest = store.load_current_run(RunPhase.ANALYZE)
    reviewed_overlay_references = tuple(
        item
        for item in reviewed_manifest.input_references
        if item.schema_name == "review-overlay-v2"
    )
    assert len(reviewed_overlay_references) == 1

    def failing_portfolio(*_args: Any, **_kwargs: Any) -> PortfolioAnalysisResult:
        raise RuntimeError("malformed Qwen output after one repair attempt")

    monkeypatch.setattr(
        "portfolio_analyzer.cli.v2.analyze_portfolio_candidates", failing_portfolio
    )
    failed = runner.invoke(
        app,
        [
            "analyze",
            "--workspace",
            str(workspace),
            "--model-dir",
            str(workspace),
            "--force",
        ],
    )
    assert failed.exit_code == 1
    failed_manifest = store.load_current_run(RunPhase.ANALYZE)
    assert failed_manifest.status == RunStatus.FAILED
    assert reviewed_overlay_references[0] in failed_manifest.input_references
    assert all(
        item.status == RunStatus.COMPLETE for item in failed_manifest.applications
    )
    assert failed_manifest.portfolio_analysis is not None
    failed_portfolio = store.load(
        failed_manifest.portfolio_analysis, PortfolioAnalysisPayload
    )
    assert failed_portfolio.status == "incomplete"
    assert failed_portfolio.candidates
    assert "malformed Qwen output after one repair attempt" in (
        failed_portfolio.reason or ""
    )

    refused = runner.invoke(app, ["report", "--workspace", str(workspace)])
    assert refused.exit_code == 1
    partial = runner.invoke(
        app, ["report", "--workspace", str(workspace), "--allow-partial"]
    )
    assert partial.exit_code == 0, partial.output
    partial_manifest = store.load_current_run(RunPhase.REPORT)
    assert partial_manifest.status == RunStatus.PARTIAL
    assert partial_manifest.report_model is not None
    partial_model = store.load(
        partial_manifest.report_model, PortfolioReportModel
    )
    assert partial_model.candidates == failed_portfolio.candidates
    assert any(
        "Review decisions were omitted" in item.reason
        for item in partial_model.omissions
    )
    partial_directory = workspace / "reports" / "runs" / partial_manifest.run_id
    assert (partial_directory / "portfolio.partial.html").is_file()
    assert latest_path.read_bytes() == latest_complete

    monkeypatch.setattr(
        "portfolio_analyzer.cli.v2.analyze_portfolio_candidates", fake_portfolio
    )
    recovered = runner.invoke(
        app,
        [
            "analyze",
            "--workspace",
            str(workspace),
            "--model-dir",
            str(workspace),
            "--force",
        ],
    )
    assert recovered.exit_code == 0, recovered.output
    recovered_manifest = store.load_current_run(RunPhase.ANALYZE)
    recovered_overlay_references = tuple(
        item
        for item in recovered_manifest.input_references
        if item.schema_name == "review-overlay-v2"
    )
    assert len(recovered_overlay_references) == 1
    recovered_overlay = store.load(
        recovered_overlay_references[0],
        ReviewOverlay,
    )
    assert len(recovered_overlay.decisions) == 1
    assert recovered_overlay.decisions[0].reviewer == "portfolio-owner"
    assert recovered_manifest.portfolio_analysis is not None
    recovered_portfolio = store.load(
        recovered_manifest.portfolio_analysis,
        PortfolioAnalysisPayload,
    )
    assert recovered_portfolio.analysis is not None
    review_records = v2_cli._report_review_records(
        recovered_overlay,
        candidates=recovered_portfolio.candidates,
        findings=recovered_portfolio.analysis.findings,
    )
    assert {item.application_id for item in review_records} == {"app-1", "app-2"}

    reextracted = runner.invoke(
        app,
        [
            "extract",
            "--workspace",
            str(workspace),
            "--application",
            "app-1",
            "--force",
        ],
    )
    assert reextracted.exit_code == 0, reextracted.output
    stale_refused = runner.invoke(app, ["report", "--workspace", str(workspace)])
    assert stale_refused.exit_code == 1
    stale_partial = runner.invoke(
        app, ["report", "--workspace", str(workspace), "--allow-partial"]
    )
    assert stale_partial.exit_code == 0, stale_partial.output
    stale_partial_manifest = store.load_current_run(RunPhase.REPORT)
    assert stale_partial_manifest.report_model is not None
    stale_partial_model = store.load(
        stale_partial_manifest.report_model,
        PortfolioReportModel,
    )
    assert {item.application_id for item in stale_partial_model.applications} == {
        "app-2"
    }
    assert stale_partial_model.excluded_application_ids == ("app-1",)
    application_omissions = tuple(
        item
        for item in stale_partial_model.omissions
        if item.application_id == "app-1" and item.logical_unit_id is None
    )
    assert len(application_omissions) == 1
    assert application_omissions[0].stage == ReportOmissionStage.ANALYZE
    assert any(
        item.application_id == "app-1"
        and item.reason == "analysis is stale for the current extraction"
        for item in stale_partial_model.omissions
    )
    assert latest_path.read_bytes() == latest_complete
