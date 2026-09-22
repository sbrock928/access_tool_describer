from pathlib import Path

import pytest
from openpyxl import Workbook

from portfolio_analyzer.inventory.loader import InventoryValidationError, load_inventory


def test_load_inventory_preserves_description_and_values(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["INVENTORY_ID", "EUCTNAME", "FILE_NAME", "FULLPATH", "DESCRIPTION", "Owner"])
    sheet.append(
        [
            "100",
            "Payment Tool",
            "payment.accdb",
            r"\\server\share\payment.accdb",
            "Old stated description",
            "Finance",
        ]
    )
    path = tmp_path / "inventory.xlsx"
    workbook.save(path)

    record = load_inventory(path)[0]

    assert record.tool_inventory_id == "100"
    assert record.inventory_filename == "payment.accdb"
    assert record.stated_description == "Old stated description"
    assert record.original_values["Owner"] == "Finance"


def test_inventory_requires_expected_columns(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.active.append(["EUCTNAME"])
    path = tmp_path / "bad.xlsx"
    workbook.save(path)

    with pytest.raises(InventoryValidationError, match="missing required columns"):
        load_inventory(path)


def test_inventory_rejects_euc_names_that_collide_as_folders(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["INVENTORY_ID", "EUCTNAME", "FILE_NAME", "FULLPATH", "DESCRIPTION"])
    sheet.append(["1", "Month/End", "one.accdb", "/source/one.accdb", "One"])
    sheet.append(["2", "Month:End", "two.accdb", "/source/two.accdb", "Two"])
    path = tmp_path / "colliding.xlsx"
    workbook.save(path)

    with pytest.raises(InventoryValidationError, match="unique workspace folders"):
        load_inventory(path)


def test_inventory_allows_multiple_files_for_one_inventory_euc(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["INVENTORY_ID", "EUCTNAME", "FILE_NAME", "FULLPATH", "DESCRIPTION"])
    sheet.append(["13", "BelloQ", "BelloQ.accdb", "/source/BelloQ.accdb", "Main"])
    sheet.append(["13", "BelloQ", "IntexLib.accdb", "/source/IntexLib.accdb", "Library"])
    path = tmp_path / "multi-application.xlsx"
    workbook.save(path)

    records = load_inventory(path)

    assert [record.inventory_filename for record in records] == [
        "BelloQ.accdb",
        "IntexLib.accdb",
    ]


def test_inventory_rejects_conflicting_names_for_one_inventory_id(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["INVENTORY_ID", "EUCTNAME", "FILE_NAME", "FULLPATH", "DESCRIPTION"])
    sheet.append(["13", "BelloQ", "BelloQ.accdb", "/source/BelloQ.accdb", "Main"])
    sheet.append(["13", "Other Tool", "other.accdb", "/source/other.accdb", "Other"])
    path = tmp_path / "conflicting-name.xlsx"
    workbook.save(path)

    with pytest.raises(InventoryValidationError, match="conflicting EUCTNAME"):
        load_inventory(path)
