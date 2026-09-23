import tracemalloc
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook

from portfolio_analyzer.models import (
    AnalysisCoverage,
    ApplicationRole,
    ApplicationTargetMapping,
    ArchitectureComponent,
    ArchitectureRelation,
    Confidence,
    Dependency,
    InventoryRecord,
    MigrationWave,
    PortfolioCluster,
    SemanticApplicationIR,
    SemanticApplicationProfile,
    SemanticCoverage,
    SemanticFinding,
    SemanticPortfolioState,
    SemanticRunMetadata,
    SimilarityEdge,
    TargetArchitecture,
)
from portfolio_analyzer.reporting.intelligence import (
    write_architecture_mermaid,
    write_intelligence_html,
    write_report_manifest,
    write_semantic_datasets,
)
from portfolio_analyzer.reporting.writers import write_executive_pdf, write_workbook
from portfolio_analyzer.semantic.config import ClusteringSettings
from portfolio_analyzer.semantic.graph import build_similarity_graph
from portfolio_analyzer.semantic.model_store import APPROVED_MODEL


def _state() -> SemanticPortfolioState:
    finding = SemanticFinding(
        tool_inventory_id="1",
        category="business_capability",
        label="Request intake",
        description="Captures and validates requests.",
        evidence_ids=["ev-1", "ev-2"],
        confidence=Confidence.HIGH,
    )
    profile = SemanticApplicationProfile(
        tool_inventory_id="1",
        tool_name="Request Tracker",
        summary="Tracks customer requests.",
        business_purpose="Request operations",
        primary_archetype="transactional workflow",
        roles=[ApplicationRole(role="transactional workflow",
                               rationale="Form supports entry", evidence_ids=["ev-1"])],
        proposed_disposition="replatform",
        confidence=Confidence.HIGH,
        findings=[finding],
        application_ir_id="ir-test",
        semantic_coverage=SemanticCoverage(
            inventory_objects=12,
            code_objects_available=5,
            code_objects_inspected=5,
            code_segments_available=8,
            code_segments_inspected=8,
            object_type_inventory={"module": 5, "table": 7},
            object_type_inspected={"module": 5},
            complete_code_coverage=True,
        ),
        evidence_ids=["ev-1", "ev-2"],
        artifact_hashes=["a" * 64],
        input_fingerprint="fingerprint",
        semantic_version="semantic-analysis-v2",
        model_repo_id=APPROVED_MODEL.repo_id,
        model_revision=APPROVED_MODEL.revision,
        model_manifest_sha256="b" * 64,
    )
    neutral = ArchitectureComponent(
        component_id="component-requests",
        track="vendor_neutral",
        name="Request domain module",
        component_type="domain_module",
        description="Owns request state and rules.",
        application_ids=["1"],
        cluster_ids=["cluster-1"],
        evidence_ids=["ev-1", "ev-2"],
        confidence=Confidence.HIGH,
    )
    microsoft = ArchitectureComponent(
        component_id="component-power-apps",
        track="microsoft",
        name="Request experience",
        component_type="application_experience",
        description="Candidate Power Apps implementation.",
        platform_service="Power Apps",
        application_ids=["1"],
        evidence_ids=["ev-1"],
        confidence=Confidence.MEDIUM,
    )
    mapping = ApplicationTargetMapping(
        mapping_id="mapping-1",
        tool_inventory_id="1",
        disposition="replatform",
        target_component_ids=["component-requests"],
        wave=1,
        rationale="Contained pilot with complete extraction.",
        evidence_ids=["ev-1", "ev-2"],
        confidence=Confidence.HIGH,
    )
    return SemanticPortfolioState(
        metadata=SemanticRunMetadata(
            semantic_version="semantic-analysis-v2",
            semantic_schema_version="semantic-schema-v2",
            prompt_version="semantic-prompts-v2",
            static_analysis_version="static-analysis-v4",
            deterministic_similarity_version="deterministic-similarity-v1",
            model_repo_id=APPROVED_MODEL.repo_id,
            model_revision=APPROVED_MODEL.revision,
            model_manifest_sha256="b" * 64,
            local_model_identifier=APPROVED_MODEL.local_identifier,
            model_architecture=APPROVED_MODEL.architecture,
            model_license=APPROVED_MODEL.license,
            inference_library="transformers",
            inference_library_version="5.17.0",
            generation_parameters={"temperature": 0.0},
            clustering_parameters={"strong_similarity": 0.88},
            approved_services=["Power Apps"],
            context_hash="context",
            generated_at=datetime.now(UTC),
            input_fingerprint="portfolio",
        ),
        observed_evidence_ids=["ev-1", "ev-2"],
        application_irs=[
            SemanticApplicationIR(
                ir_id="ir-test",
                tool_inventory_id="1",
                ir_version="application-ir-v1",
                input_fingerprint="ir-fingerprint",
                code_object_count=5,
                code_segment_count=8,
                object_type_counts={"module": 5},
                model_input_sha256="c" * 64,
                model_input_characters=1200,
            )
        ],
        applications=[profile],
        similarity_edges=[
            SimilarityEdge(
                source_tool_id="1",
                target_tool_id="1",
                overall_similarity=1.0,
                category_scores={"business_capabilities": 1.0},
                shared_features={"business_capabilities": ["request intake"]},
                shared_capabilities=["Request intake"],
            )
        ],
        clusters=[
            PortfolioCluster(
                cluster_id="cluster-1",
                label="Request operations",
                application_ids=["1"],
                shared_capabilities=["Request intake"],
                rationale="Evidence-grounded request workflow.",
                confidence=Confidence.HIGH,
                evidence_ids=["ev-1", "ev-2"],
            )
        ],
        architecture=TargetArchitecture(
            title="Request platform",
            summary="Modular request architecture.",
            components=[neutral, microsoft],
            relations=[
                ArchitectureRelation(
                    relation_id="relation-1",
                    track="vendor_neutral",
                    source_component_id="component-requests",
                    target_component_id="component-requests",
                    relationship="owns",
                    evidence_ids=["ev-1"],
                    confidence=Confidence.MEDIUM,
                )
            ],
            mappings=[mapping],
            migration_waves=[
                MigrationWave(
                    wave=1,
                    name="Low-coupling pilots",
                    purpose="Validate the target pattern.",
                    application_ids=["1"],
                )
            ],
        ),
    )


def test_semantic_reports_are_offline_traceable_and_reviewable(tmp_path: Path) -> None:
    state = _state()
    inventory = [
        InventoryRecord(
            tool_inventory_id="1",
            tool_name="Request Tracker",
            inventory_filename="request.accdb",
            filepath=Path("request.accdb"),
        )
    ]
    coverage = [
        AnalysisCoverage(
            tool_inventory_id="1",
            tool_name="Request Tracker",
            staging_status="staged",
            extraction_status="complete",
            analysis_status="complete",
        )
    ]
    dependencies = [
        Dependency(
            tool_inventory_id="1",
            source="Request form",
            target="Request table",
            dependency_type="form_record_source",
            operation="READ_WRITE",
            confidence=Confidence.HIGH,
        )
    ]
    workbook_path = tmp_path / "Portfolio_Analysis.xlsx"
    write_workbook(
        workbook_path,
        inventory,
        [],
        [],
        [],
        dependencies,
        [],
        coverage=coverage,
        semantic=state,
        semantic_status="current",
    )
    workbook = load_workbook(workbook_path)
    assert {
        "Application Portfolio",
        "Target Architecture",
        "App-Target Crosswalk",
        "Migration Roadmap",
        "Semantic Findings",
        "Review Queue",
        "Method & Provenance",
    }.issubset(workbook.sheetnames)
    assert workbook["Review Queue"].data_validations.count == 1
    assert len(workbook["Portfolio Summary"]._charts) == 5
    portfolio_headers = [cell.value for cell in workbook["Application Portfolio"][1]]
    assert "Deterministic Code Coverage" in portfolio_headers
    provenance = {
        row[0].value: row[1].value
        for row in workbook["Method & Provenance"].iter_rows(min_row=2, max_col=2)
    }
    assert provenance["Applications with complete code coverage"] == 1

    pdf_path = tmp_path / "Portfolio_Analysis.pdf"
    write_executive_pdf(
        pdf_path,
        inventory,
        [],
        [],
        coverage=coverage,
        semantic=state,
        semantic_status="current",
    )
    assert pdf_path.read_bytes().startswith(b"%PDF")

    html_path = tmp_path / "Portfolio_Intelligence.html"
    write_intelligence_html(
        html_path,
        inventory,
        coverage,
        state,
        semantic_status="current",
        dependencies=dependencies,
    )
    html = html_path.read_text(encoding="utf-8")
    assert "default-src 'none'" in html
    assert "https://" not in html
    assert "Request Tracker" in html
    assert "Observed dependency map" in html
    assert "Request table" in html
    assert "Observed sources" in html
    assert "Deterministic code coverage" in html
    assert "Owner claims" in html
    assert "Semantic proposals" in html
    assert "unknown-evidence" not in html

    outputs = write_semantic_datasets(tmp_path, state, {"1": "Request Tracker"})
    assert (tmp_path / "semantic_application_irs.csv") in outputs
    mermaid_path = tmp_path / "Target_Architecture.md"
    write_architecture_mermaid(mermaid_path, state)
    assert "flowchart LR" in mermaid_path.read_text(encoding="utf-8")
    manifest_path = tmp_path / "report_manifest.json"
    write_report_manifest(
        manifest_path,
        [workbook_path, pdf_path, html_path, mermaid_path, *outputs],
        state,
        semantic_status="current",
    )
    assert '"sha256"' in manifest_path.read_text(encoding="utf-8")


def test_quick_semantic_workbook_is_visibly_test_only(tmp_path: Path) -> None:
    state = _state()
    state.metadata.run_mode = "quick"
    state.metadata.max_objects_per_application = 5
    workbook_path = tmp_path / "Quick_Analysis.xlsx"
    write_workbook(
        workbook_path,
        [],
        [],
        [],
        [],
        [],
        [],
        semantic=state,
        semantic_status="TEST ONLY — QUICK MODE",
    )

    workbook = load_workbook(workbook_path)
    assert "TEST ONLY" in workbook["Portfolio Summary"]["A1"].value
    provenance = {
        row[0].value: row[1].value
        for row in workbook["Method & Provenance"].iter_rows(min_row=2, max_col=2)
    }
    assert provenance["Run mode"] == "quick"
    assert provenance["Maximum semantic objects per application"] == 5


def test_synthetic_500_application_graph_and_reports_stay_bounded(tmp_path: Path) -> None:
    base = _state()
    profiles = []
    inventory = []
    coverage = []
    mappings = []
    for index in range(500):
        tool_id = f"app-{index:03d}"
        profile = base.applications[0].model_copy(deep=True)
        profile.tool_inventory_id = tool_id
        profile.tool_name = f"Synthetic Application {index:03d}"
        profile.input_fingerprint = f"fingerprint-{index}"
        profile.findings[0].tool_inventory_id = tool_id
        profiles.append(profile)
        inventory.append(
            InventoryRecord(
                tool_inventory_id=tool_id,
                tool_name=profile.tool_name,
                inventory_filename=f"{tool_id}.accdb",
                filepath=Path(f"{tool_id}.accdb"),
            )
        )
        coverage.append(
            AnalysisCoverage(
                tool_inventory_id=tool_id,
                tool_name=profile.tool_name,
                staging_status="staged",
                extraction_status="complete",
                analysis_status="complete",
            )
        )
        mapping = base.architecture.mappings[0].model_copy(deep=True)
        mapping.mapping_id = f"mapping-{index:03d}"
        mapping.tool_inventory_id = tool_id
        mappings.append(mapping)

    tracemalloc.start()
    edges, clusters = build_similarity_graph(
        profiles,
        [],
        ClusteringSettings(
            strong_similarity=1.0,
            corroborated_similarity=1.0,
            fixed_seed=11,
        ),
    )
    state = base.model_copy(deep=True)
    state.applications = profiles
    state.similarity_edges = edges
    state.clusters = clusters
    state.architecture.mappings = mappings
    html_path = tmp_path / "Portfolio_Intelligence.html"
    write_intelligence_html(
        html_path,
        inventory,
        coverage,
        state,
        semantic_status="current: 500/500 application profiles",
    )
    workbook_path = tmp_path / "Portfolio_Analysis.xlsx"
    write_workbook(
        workbook_path,
        inventory,
        [],
        [],
        [],
        [],
        [],
        coverage=coverage,
        semantic=state,
        semantic_status="current",
    )
    _, peak = tracemalloc.get_traced_memory()
    tracemalloc.stop()

    assert len(clusters) == 500
    assert not edges
    assert html_path.stat().st_size > 100_000
    assert workbook_path.stat().st_size > 10_000
    assert peak < 100 * 1024 * 1024


def test_behavior_reports_preserve_provenance_and_resolvable_sources(tmp_path: Path) -> None:
    import csv
    import json
    import re

    from portfolio_analyzer.models import BehaviorFact, SemanticSource

    state = _state()
    profile = state.applications[0]
    profile.purpose_provenance = "owner_claim"
    profile.purpose_claim_ids = ["claim-purpose"]
    profile.observed_behavior = ["Reads Requests (local datasource)", "Exports spreadsheet data"]
    profile.inputs = ["Requests"]
    profile.outputs = ["Spreadsheet export"]
    profile.secondary_capabilities = ["Exports spreadsheet data"]
    profile.classification_rationale = "Form contains record editing code."
    profile.classification_evidence_ids = ["src-form"]
    state.application_irs[0].behavior_facts = [
        BehaviorFact(
            action="entry",
            description="Form contains record editing code",
            object_type="form",
            object_name="Requests",
            targets=["Requests"],
            evidence_ids=["src-form"],
        )
    ]
    state.sources = [
        SemanticSource(
            source_id="src-form",
            tool_inventory_id="1",
            artifact_hash="a" * 64,
            object_type="form",
            object_name="Requests",
            content_sha256="b" * 64,
            excerpt="Me.Recordset.Update",
            ui_properties={"recordsource": "Requests"},
        )
    ]
    inventory = [
        InventoryRecord(
            tool_inventory_id="1",
            tool_name="Request Tracker",
            inventory_filename="requests.accdb",
            filepath=Path("requests.accdb"),
        )
    ]
    workbook_path = tmp_path / "behavior.xlsx"
    write_workbook(workbook_path, inventory, [], [], [], [], [], semantic=state)
    workbook = load_workbook(workbook_path)
    rows = list(workbook["Application Portfolio"].values)
    row = dict(zip(rows[0], rows[1], strict=True))
    assert row["Purpose Provenance"] == "owner_claim"
    assert row["Classification Rationale"] == profile.classification_rationale
    assert row["Known Inputs"] == "Requests"
    assert row["Secondary Capabilities"] == "Exports spreadsheet data"
    assert workbook["Behavior Sources"]["B2"].value == "src-form"
    assert workbook["Application Behavior"]["H2"].value == "src-form"

    write_semantic_datasets(tmp_path, state, {"1": "Request Tracker"})
    with (tmp_path / "semantic_applications.csv").open() as handle:
        application = next(csv.DictReader(handle))
    assert application["classification_rationale"] == profile.classification_rationale
    assert application["purpose_provenance"] == "owner_claim"
    assert application["purpose_claim_ids"] == "claim-purpose"
    with (tmp_path / "application_behaviors.csv").open() as handle:
        behavior = next(csv.DictReader(handle))
    with (tmp_path / "behavior_sources.csv").open() as handle:
        source = next(csv.DictReader(handle))
    assert behavior["evidence_ids"] == source["source_id"] == "src-form"
    assert json.loads(source["ui_properties"]) == {"recordsource": "Requests"}

    path = tmp_path / "behavior.html"
    write_intelligence_html(path, inventory, [], state, semantic_status="current")
    html = path.read_text()
    payload = re.search(
        r'<script id="portfolio-data" type="application/json">(.*?)</script>', html, re.S
    )
    assert payload is not None
    application_data = json.loads(payload[1])["applications"][0]
    assert application_data["classification_rationale"] == application["classification_rationale"]
    assert application_data["observed_behavior"] == profile.observed_behavior
    assert application_data["inputs"] == profile.inputs
    assert "Why this classification" in html
    assert "Root UI properties" in html
