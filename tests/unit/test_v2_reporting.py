from __future__ import annotations

import csv
import io
import shutil
import subprocess
import zipfile
from datetime import UTC, datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook

import portfolio_analyzer.v2.reporting as reporting
from portfolio_analyzer.v2.candidates import (
    SEMANTIC_SIMILARITY_DISABLED_WARNING,
    generate_portfolio_candidates,
)
from portfolio_analyzer.v2.fingerprints import (
    evidence_bundle_fingerprint,
    interpretation_fingerprint,
)
from portfolio_analyzer.v2.models import (
    PARTIAL_REPORT_WATERMARK,
    AccessObjectEvidence,
    AccessObjectType,
    AccessTableEvidence,
    ApplicationEvidenceBundle,
    ApplicationProfile,
    ArtifactEvidence,
    Confidence,
    ConnectionEvidence,
    ConnectionKind,
    ConnectionProvenance,
    DataAccess,
    DataOperation,
    DatasourceIdentity,
    DependencyEdge,
    DependencyNode,
    DependencyNodeKind,
    DependencyRelation,
    EvidenceRecord,
    ExtractionCoverage,
    ExtractionStatus,
    InteractionScope,
    InterpretationKind,
    ModelProvenance,
    OwnerClaim,
    PortfolioFinding,
    PortfolioInterpretation,
    PortfolioReportModel,
    QueryEvidence,
    QueryKind,
    ReportOmission,
    ReportOmissionStage,
    ReportReviewRecord,
    ReportStatus,
    ResolutionStatus,
    ReviewStatus,
    UnresolvedReference,
)
from portfolio_analyzer.v2.report_model import (
    IncompleteReportError,
    build_portfolio_report_model,
)
from portfolio_analyzer.v2.reporting import (
    ReportFormat,
    ReportLeakError,
    ReportPublicationError,
    load_latest_report,
    load_report_publication,
    publish_report_run,
    render_report_csv_bundle,
    render_report_html,
    write_report_pdf,
    write_report_xlsx,
)
from portfolio_analyzer.v2.review import REVIEW_HEADERS, import_review_overlay
from portfolio_analyzer.v2.workflow import InventoryExclusion

NOW = datetime(2026, 9, 28, 12, 0, tzinfo=UTC)
MODEL_DIGEST = "a" * 64


def test_builder_is_self_contained_deterministic_and_uses_stable_lineage() -> None:
    first = _bundle("app-1", "=Quarterly Tool", "1")
    second = _bundle("app-2", "Second Tool", "2")
    profiles = (_profile(first), _profile(second))
    candidates = generate_portfolio_candidates((second, first), tuple(reversed(profiles)))
    portfolio = _portfolio(profiles, first, second)
    review = ReportReviewRecord(
        application_id="app-1",
        subject_id=candidates[0].candidate_id,
        status=ReviewStatus.PENDING,
        evidence_ids=candidates[0].evidence_ids,
    )

    report = build_portfolio_report_model(
        (second, first),
        tuple(reversed(profiles)),
        portfolio,
        candidates,
        generated_at=NOW,
        reviews=(review,),
        inventory_exclusions=(_inventory_exclusion(),),
    )
    reordered = build_portfolio_report_model(
        (first, second),
        profiles,
        portfolio,
        tuple(reversed(candidates)),
        generated_at=NOW + timedelta(hours=1),
        reviews=(review,),
        inventory_exclusions=(_inventory_exclusion(),),
    )

    assert report.status == ReportStatus.COMPLETE
    assert SEMANTIC_SIMILARITY_DISABLED_WARNING in report.warnings
    assert report.absence_claims_suppressed
    assert report.report_id == reordered.report_id
    assert report.analysis_fingerprint == reordered.analysis_fingerprint
    assert len(report.observed_facts) == 4
    assert len(report.owner_claims) == 2
    assert len(report.object_registry) == 4
    assert len(report.table_registry) == 2
    assert len(report.query_registry) == 2
    assert len(report.connection_registry) == 4
    assert len(report.datasource_registry) == 1
    assert len(report.data_access) == 2
    assert len(report.dependency_nodes) == 4
    assert len(report.dependency_edges) == 2
    assert len(report.pass_through_queries) == 2
    assert len(report.linked_tables) == 2
    assert len(report.inventory_exclusions) == 1
    assert report.omissions == ()
    pass_through = next(
        item for item in report.pass_through_queries if item.application_id == "app-1"
    )
    assert pass_through.artifact_id
    assert pass_through.query_kind == QueryKind.PASS_THROUGH
    assert pass_through.connection_kind == ConnectionKind.PASS_THROUGH_QUERY
    assert pass_through.direct_provenance == ConnectionProvenance.QUERYDEF_CONNECT
    assert pass_through.resolution_status == ResolutionStatus.RESOLVED
    assert pass_through.resolution_provenance == (
        "declared_connection_metadata",
        "dsn_registry_64_bit",
    )
    assert pass_through.resolution_warnings == ("registry fallback was not required",)
    assert pass_through.dsn == "CorporateTrustDsn"
    assert pass_through.driver == "ODBC Driver 18 for SQL Server"
    target = pass_through.targets[0]
    assert target.operation == DataOperation.READ
    assert target.catalog == "TrustCatalog"
    assert target.database == "CorporateTrust"
    assert target.schema_name == "dbo"
    assert target.object_name == "Deal"
    linked = next(item for item in report.linked_tables if item.application_id == "app-1")
    assert linked.artifact_id
    assert linked.object_kind == AccessObjectType.LINKED_TABLE
    assert linked.connection_kind == ConnectionKind.LINKED_TABLE
    assert linked.direct_provenance == ConnectionProvenance.TABLEDEF_CONNECT
    assert linked.targets[0].operation == DataOperation.UNKNOWN
    assert linked.targets[0].schema_name == "dbo"
    assert linked.targets[0].object_name == "Deal"


def test_optional_extraction_gaps_are_qualified_without_allow_partial(
    tmp_path: Path,
) -> None:
    first = _bundle("app-1", "First", "1", status=ExtractionStatus.PARTIAL)
    second = _bundle("app-2", "Second", "2")
    profiles = (_profile(first), _profile(second))
    report = build_portfolio_report_model(
        (first, second),
        profiles,
        _portfolio(profiles, first, second),
        generate_portfolio_candidates((first, second), profiles),
        generated_at=NOW,
    )

    assert report.status == ReportStatus.COMPLETE
    assert report.omissions == ()
    assert report.partial_watermark is None
    assert report.absence_claims_suppressed is True
    assert any("qualified extraction coverage" in item for item in report.warnings)

    html = render_report_html(report).decode("utf-8")
    assert "qualified extraction coverage" in html
    xlsx_path = tmp_path / "qualified.xlsx"
    write_report_xlsx(xlsx_path, report)
    workbook = load_workbook(xlsx_path, read_only=True)
    try:
        application_rows = tuple(workbook["Applications"].iter_rows(values_only=True))
        assert "qualified extraction coverage" in str(application_rows[1][-1])
    finally:
        workbook.close()
    with zipfile.ZipFile(io.BytesIO(render_report_csv_bundle(report))) as archive:
        assert b"qualified extraction coverage" in archive.read("applications.csv")
    pdf_path = tmp_path / "qualified.pdf"
    write_report_pdf(pdf_path, report)
    assert b"qualified extraction coverage" in pdf_path.read_bytes()


def test_semantic_omissions_require_explicit_partial_and_are_structured() -> None:
    bundle = _bundle("app-1", "First", "1")
    logical_omission = ReportOmission(
        application_id="app-1",
        logical_unit_id=bundle.objects[0].object_id,
        stage=ReportOmissionStage.ANALYZE,
        reason="Logical unit inference failed",
    )
    application_omission = ReportOmission(
        application_id="app-excluded",
        stage=ReportOmissionStage.ANALYZE,
        reason="Analysis is stale for the current extraction",
    )

    with pytest.raises(IncompleteReportError, match="allow_partial"):
        build_portfolio_report_model(
            (bundle,),
            (),
            None,
            (),
            generated_at=NOW,
            omissions=(logical_omission, application_omission),
            excluded_application_ids=("app-excluded",),
        )

    report = build_portfolio_report_model(
        (bundle,),
        (),
        None,
        (),
        generated_at=NOW,
        omissions=(logical_omission, logical_omission, application_omission),
        allow_partial=True,
        excluded_application_ids=("app-excluded",),
        inventory_exclusions=(_inventory_exclusion(),),
    )

    assert report.status == ReportStatus.PARTIAL
    assert report.partial_watermark == PARTIAL_REPORT_WATERMARK
    assert report.absence_claims_suppressed is True
    assert logical_omission.omission_id in {item.omission_id for item in report.omissions}
    assert len({item.omission_id for item in report.omissions}) == len(report.omissions)
    excluded_omissions = tuple(
        item
        for item in report.omissions
        if item.application_id == "app-excluded" and item.logical_unit_id is None
    )
    assert excluded_omissions == (application_omission,)
    assert report.inventory_exclusions


def test_all_renderers_preserve_identity_counts_and_review_queue_round_trips(
    tmp_path: Path,
) -> None:
    report = _complete_report()
    reports_root = tmp_path / "reports"

    manifest = publish_report_run(reports_root, "run-complete", report)

    assert manifest == load_report_publication(reports_root, "run-complete")
    assert manifest == load_latest_report(reports_root)
    assert manifest.report_status == ReportStatus.COMPLETE
    assert manifest.record_count == len(manifest.record_ids)
    assert all(item.record_ids == manifest.record_ids for item in manifest.artifacts)
    lineage_target_ids = {
        target.target_id
        for view in report.pass_through_queries
        for target in view.targets
    } | {
        target.target_id for view in report.linked_tables for target in view.targets
    }
    assert lineage_target_ids.issubset(manifest.record_ids)
    assert all(".partial." not in item.relative_path for item in manifest.artifacts)

    by_format = {
        item.report_format: reports_root / "runs" / "run-complete" / item.relative_path
        for item in manifest.artifacts
    }
    html = by_format[ReportFormat.HTML].read_text(encoding="utf-8")
    assert "default-src 'none'" in html
    assert "<script>alert(1)</script>" not in html
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in html
    assert all(identifier in html for identifier in manifest.record_ids)
    for required in (
        "Artifact ID",
        "Query Kind",
        "Connection Kind",
        "Direct Provenance",
        "Resolution Provenance",
        "Resolution Warnings",
        "Target Database",
        "Target Object",
    ):
        assert f"<th>{required}</th>" in html
    assert "TrustCatalog" in html
    assert "dsn_registry_64_bit" in html

    workbook = load_workbook(by_format[ReportFormat.XLSX], data_only=False)
    try:
        record_ids = {
            str(row[0])
            for row in workbook["Record IDs"].iter_rows(min_row=2, values_only=True)
            if row[0] is not None
        }
        assert record_ids == set(manifest.record_ids)
        assert tuple(cell.value for cell in workbook["Review Queue"][1]) == REVIEW_HEADERS
        queue_rows = tuple(workbook["Review Queue"].iter_rows(min_row=2, values_only=True))
        assert queue_rows
        assert {str(row[0]) for row in queue_rows} == {report.analysis_fingerprint}
        assert workbook["Applications"]["B2"].value == "'=Quarterly Tool"
        pass_headers = tuple(cell.value for cell in workbook["Pass Through"][1])
        pass_rows = tuple(
            dict(zip(pass_headers, row, strict=True))
            for row in workbook["Pass Through"].iter_rows(min_row=2, values_only=True)
        )
        pass_row = next(row for row in pass_rows if row["Application ID"] == "app-1")
        assert pass_row["Application ID"] == "app-1"
        assert str(pass_row["Artifact ID"]).startswith("artifact_")
        assert str(pass_row["Object ID"]).startswith("object_")
        assert pass_row["Query Kind"] == "pass_through"
        assert pass_row["Connection Kind"] == "pass_through_query"
        assert pass_row["Direct Provenance"] == "QueryDef.Connect"
        assert pass_row["Resolution Status"] == "resolved"
        assert "dsn_registry_64_bit" in str(pass_row["Resolution Provenance"])
        assert pass_row["Resolution Warnings"] == "registry fallback was not required"
        assert pass_row["DSN"] == "CorporateTrustDsn"
        assert pass_row["Platform"] == "Microsoft SQL Server"
        assert pass_row["Driver"] == "ODBC Driver 18 for SQL Server"
        assert pass_row["Server"] == "SQL01"
        assert pass_row["Database"] == "CorporateTrust"
        assert pass_row["Catalog"] == "TrustCatalog"
        assert pass_row["Schema"] == "dbo"
        assert pass_row["Target Object"] == "Deal"
        assert pass_row["Operation"] == "READ"
        assert pass_row["Target Evidence IDs"]
        assert all(
            cell.data_type != "f"
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
        )
        workbook["Review Queue"]["D2"] = "accept"
        reviewed_path = tmp_path / "reviewed.xlsx"
        workbook.save(reviewed_path)
    finally:
        workbook.close()

    known_proposals = {
        item.candidate_id: item.evidence_ids for item in report.candidates
    } | {
        item.finding_id: item.evidence_ids for item in report.portfolio_findings
    }
    overlay = import_review_overlay(
        reviewed_path,
        analysis_fingerprint=report.analysis_fingerprint,
        known_proposals=known_proposals,
        imported_at=NOW,
    )
    assert len(overlay.decisions) == 1

    with zipfile.ZipFile(by_format[ReportFormat.CSV_BUNDLE]) as archive:
        identity_rows = csv.DictReader(
            io.StringIO(archive.read("identity.csv").decode("utf-8"))
        )
        assert {row["entity_id"] for row in identity_rows} == set(manifest.record_ids)
        application_rows = tuple(
            csv.DictReader(
                io.StringIO(archive.read("applications.csv").decode("utf-8"))
            )
        )
        assert application_rows[0]["application_name"] == "'=Quarterly Tool"
        pass_rows = tuple(
            csv.DictReader(
                io.StringIO(
                    archive.read("pass_through_queries.csv").decode("utf-8")
                )
            )
        )
        pass_row = next(row for row in pass_rows if row["application_id"] == "app-1")
        assert pass_row["artifact_id"].startswith("artifact_")
        assert pass_row["object_id"].startswith("object_")
        assert pass_row["query_kind"] == "pass_through"
        assert pass_row["connection_kind"] == "pass_through_query"
        assert pass_row["direct_provenance"] == "QueryDef.Connect"
        assert pass_row["resolution_status"] == "resolved"
        assert "dsn_registry_64_bit" in pass_row["resolution_provenance"]
        assert pass_row["resolution_warnings"] == "registry fallback was not required"
        assert pass_row["dsn"] == "CorporateTrustDsn"
        assert pass_row["catalog"] == "TrustCatalog"
        assert pass_row["target_database"] == "CorporateTrust"
        assert pass_row["schema"] == "dbo"
        assert pass_row["target_object"] == "Deal"
        assert pass_row["operation"] == "READ"
        linked_rows = tuple(
            csv.DictReader(io.StringIO(archive.read("linked_tables.csv").decode("utf-8")))
        )
        linked_row = next(row for row in linked_rows if row["application_id"] == "app-1")
        assert linked_row["object_kind"] == "linked_table"
        assert linked_row["connection_kind"] == "linked_table"
        assert linked_row["direct_provenance"] == "TableDef.Connect"
        assert linked_row["resolution_status"] == "resolved"
        assert linked_row["dsn"] == "CorporateTrustDsn"
        assert linked_row["platform"] == "Microsoft SQL Server"
        assert linked_row["driver"] == "ODBC Driver 18 for SQL Server"
        assert linked_row["server"] == "SQL01"
        assert linked_row["database"] == "CorporateTrust"
        assert linked_row["schema"] == "dbo"
        assert linked_row["target_object"] == "Deal"
        assert {
            "data_access.csv",
            "dependency_edges.csv",
            "tables.csv",
            "queries.csv",
            "inventory_exclusions.csv",
        }.issubset(archive.namelist())
    pdf = by_format[ReportFormat.PDF].read_bytes()
    assert pdf.startswith(b"%PDF-")
    for required_value in (
        b"CorporateTrustDsn",
        b"dsn_registry_64_bit",
        b"TrustCatalog",
        b"QueryDef.Connect",
        b"TableDef.Connect",
        b"registry fallback was not required",
    ):
        assert required_value in pdf


def test_pdf_extracted_text_preserves_visible_ids_counts_and_lineage(tmp_path: Path) -> None:
    pdftotext = shutil.which("pdftotext")
    if pdftotext is None:
        pytest.skip("Poppler pdftotext is unavailable")
    report = _complete_report()
    pdf_path = tmp_path / "portfolio.pdf"
    write_report_pdf(pdf_path, report)

    extracted = subprocess.run(
        [pdftotext, str(pdf_path), "-"],
        check=True,
        capture_output=True,
        text=True,
    ).stdout
    compact = "".join(extracted.split())
    for identifier in (
        *(item.application_id for item in report.applications),
        *(item.candidate_id for item in report.candidates),
        *(item.finding_id for item in report.portfolio_findings),
    ):
        assert identifier in compact
    assert "Applications" in extracted
    assert "Candidates" in extracted
    assert "Findings" in extracted
    assert str(len(report.applications)) in extracted
    assert str(len(report.candidates)) in extracted
    assert str(len(report.portfolio_findings)) in extracted
    for lineage_value in (
        "CorporateTrustDsn",
        "QueryDef.Connect",
        "TableDef.Connect",
        "TrustCatalog",
        "dbo",
        "Deal",
    ):
        assert lineage_value in extracted


def test_partial_publication_is_distinct_and_never_replaces_complete_latest(
    tmp_path: Path,
) -> None:
    reports_root = tmp_path / "reports"
    complete = _complete_report()
    complete_manifest = publish_report_run(reports_root, "run-complete", complete)
    partial = build_portfolio_report_model(
        (_bundle("app-1", "First", "1"),),
        (),
        None,
        (),
        generated_at=NOW,
        allow_partial=True,
    )

    partial_manifest = publish_report_run(reports_root, "run-partial", partial)

    assert partial_manifest.report_status == ReportStatus.PARTIAL
    assert all(".partial." in item.relative_path for item in partial_manifest.artifacts)
    assert load_latest_report(reports_root) == complete_manifest
    html_artifact = next(
        item for item in partial_manifest.artifacts if item.report_format == ReportFormat.HTML
    )
    html = (reports_root / "runs" / "run-partial" / html_artifact.relative_path).read_text()
    assert PARTIAL_REPORT_WATERMARK in html


def test_publication_is_stable_hash_idempotent_and_fails_closed(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    report = _complete_report()
    reports_root = tmp_path / "reports"
    first = publish_report_run(reports_root, "run-stable", report)
    regenerated = report.model_copy(update={"generated_at": NOW + timedelta(days=1)})

    assert publish_report_run(reports_root, "run-stable", regenerated) == first

    deferred_root = tmp_path / "deferred-latest"
    deferred = publish_report_run(
        deferred_root, "run-deferred", report, publish_latest=False
    )
    assert not (deferred_root / "latest.json").exists()
    assert publish_report_run(deferred_root, "run-deferred", report) == deferred
    assert load_latest_report(deferred_root) == deferred

    forged = report.model_copy(update={"warnings": ("PWD=cleartext",)})
    with pytest.raises(ReportLeakError):
        render_report_html(forged)
    for credential_alias in (
        "User=user-canary",
        "User Name=user-name-canary",
        "Credential=credential-canary",
        "Account=account-canary",
        "Client ID=client-id-canary",
    ):
        forged_credential = report.model_copy(update={"warnings": (credential_alias,)})
        with pytest.raises(ReportLeakError, match="credential assignment"):
            render_report_html(forged_credential)
    for raw_connection in (
        "Provider=SQLOLEDB;Data Source=SQL01;Initial Catalog=CorporateTrust",
        r"DBQ=C:\private\finance.accdb",
        r"FILEDSN=C:\private\finance.dsn",
        "OLEDB:Provider=SQLOLEDB;Data Source=SQL01",
    ):
        forged_connection = report.model_copy(
            update={"warnings": (raw_connection,)}
        )
        with pytest.raises(ReportLeakError, match="raw connection"):
            render_report_html(forged_connection)

    latest_failure_root = tmp_path / "latest-failure"
    real_atomic_write = reporting._atomic_write_bytes

    def fail_latest(path: Path, payload: bytes) -> None:
        if path.name == "latest.json":
            raise OSError("simulated latest pointer crash")
        real_atomic_write(path, payload)

    monkeypatch.setattr(reporting, "_atomic_write_bytes", fail_latest)
    with pytest.raises(ReportPublicationError, match="could not publish"):
        publish_report_run(latest_failure_root, "run-crash", report)
    assert load_report_publication(latest_failure_root, "run-crash").report_id == report.report_id
    assert not (latest_failure_root / "latest.json").exists()


def test_loader_rejects_tampering_even_when_manifest_is_present(tmp_path: Path) -> None:
    reports_root = tmp_path / "reports"
    manifest = publish_report_run(reports_root, "run-tamper", _complete_report())
    artifact = manifest.artifacts[0]
    path = reports_root / "runs" / "run-tamper" / artifact.relative_path
    path.write_bytes(path.read_bytes() + b"tampered")

    with pytest.raises(ReportPublicationError, match="hash verification"):
        load_report_publication(reports_root, "run-tamper")


def test_all_formats_publish_for_five_hundred_applications(tmp_path: Path) -> None:
    bundles = tuple(
        _bundle(f"app-{index:03d}", f"Application {index:03d}", "a")
        for index in range(500)
    )
    profiles = tuple(_profile(bundle) for bundle in bundles)
    portfolio = PortfolioInterpretation(
        application_interpretation_sha256s=tuple(
            interpretation_fingerprint(profile) for profile in profiles
        ),
        provenance=ModelProvenance(
            model_manifest_sha256=MODEL_DIGEST,
            prompt_version="qwen-portfolio-v2",
            output_schema_version="qwen-portfolio-output-v2",
            inference_library_version="5.17.0",
        ),
    )
    report = build_portfolio_report_model(
        bundles,
        profiles,
        portfolio,
        (),
        generated_at=NOW,
    )

    manifest = publish_report_run(tmp_path / "reports", "run-500", report)

    assert len(report.applications) == 500
    assert {item.report_format for item in manifest.artifacts} == set(ReportFormat)
    assert manifest.record_count == len(manifest.record_ids)


def _complete_report() -> PortfolioReportModel:
    first = _bundle("app-1", "=Quarterly Tool", "1")
    second = _bundle("app-2", "Second Tool", "2")
    profiles = (_profile(first), _profile(second))
    candidates = generate_portfolio_candidates((first, second), profiles)
    portfolio = _portfolio(profiles, first, second)
    return build_portfolio_report_model(
        (first, second),
        profiles,
        portfolio,
        candidates,
        generated_at=NOW,
        reviews=(
            ReportReviewRecord(
                application_id="app-1",
                subject_id=candidates[0].candidate_id,
                status=ReviewStatus.PENDING,
                evidence_ids=candidates[0].evidence_ids,
            ),
        ),
        inventory_exclusions=(_inventory_exclusion(),),
    )


def _bundle(
    application_id: str,
    application_name: str,
    digest_character: str,
    *,
    status: ExtractionStatus = ExtractionStatus.COMPLETE,
) -> ApplicationEvidenceBundle:
    warnings = (
        ("SaveAsText unavailable for optional form metadata",)
        if status == ExtractionStatus.PARTIAL
        else ()
    )
    artifact = ArtifactEvidence(
        application_id=application_id,
        source_locator=rf"\\fileserver\apps\{application_id}.accdb",
        staged_relative_path=f"applications/{application_id}/tool.accdb",
        filename="tool.accdb",
        access_format="accdb",
        sha256=digest_character * 64,
        size_bytes=42,
        extractor_version="windows-com-v2",
        extraction_status=status,
        extracted_at=NOW,
        warnings=warnings,
    )
    query_object = AccessObjectEvidence(
        artifact_id=artifact.artifact_id,
        object_type=AccessObjectType.QUERY,
        name="qryShared",
        sanitized_definition="SELECT * FROM dbo.Deal",
    )
    table_object = AccessObjectEvidence(
        artifact_id=artifact.artifact_id,
        object_type=AccessObjectType.LINKED_TABLE,
        name="lnkDeal",
    )
    query_evidence = EvidenceRecord(
        application_id=application_id,
        artifact_id=artifact.artifact_id,
        artifact_sha256=artifact.sha256,
        object_id=query_object.object_id,
        fact_type="query_definition",
        observation="<script>alert(1)</script> reads approved deal data",
    )
    table_evidence = EvidenceRecord(
        application_id=application_id,
        artifact_id=artifact.artifact_id,
        artifact_sha256=artifact.sha256,
        object_id=table_object.object_id,
        fact_type="linked_table",
        observation="Linked table targets the approved warehouse",
    )
    query_object = query_object.model_copy(
        update={"evidence_ids": (query_evidence.evidence_id,)}
    )
    table_object = table_object.model_copy(
        update={"evidence_ids": (table_evidence.evidence_id,)}
    )
    datasource = DatasourceIdentity(
        platform="Microsoft SQL Server",
        driver="ODBC Driver 18 for SQL Server",
        dsn="CorporateTrustDsn",
        server="SQL01",
        database="CorporateTrust",
        resource=r"\\fileserver\shared\mapping.csv",
    )
    query_connection = ConnectionEvidence(
        application_id=application_id,
        artifact_id=artifact.artifact_id,
        artifact_sha256=artifact.sha256,
        object_id=query_object.object_id,
        connection_kind=ConnectionKind.PASS_THROUGH_QUERY,
        sanitized_summary="SQL Server SQL01 / CorporateTrust",
        direct_provenance=ConnectionProvenance.QUERYDEF_CONNECT,
        resolution_status=ResolutionStatus.RESOLVED,
        resolution_provenance=(
            "declared_connection_metadata",
            "dsn_registry_64_bit",
        ),
        resolution_warnings=("registry fallback was not required",),
        datasource_id=datasource.datasource_id,
        evidence_ids=(query_evidence.evidence_id,),
    )
    table_connection = ConnectionEvidence(
        application_id=application_id,
        artifact_id=artifact.artifact_id,
        artifact_sha256=artifact.sha256,
        object_id=table_object.object_id,
        connection_kind=ConnectionKind.LINKED_TABLE,
        sanitized_summary="SQL Server SQL01 / CorporateTrust",
        direct_provenance=ConnectionProvenance.TABLEDEF_CONNECT,
        resolution_status=ResolutionStatus.RESOLVED,
        resolution_provenance=(
            "declared_connection_metadata",
            "dsn_registry_64_bit",
        ),
        resolution_warnings=("registry fallback was not required",),
        datasource_id=datasource.datasource_id,
        evidence_ids=(table_evidence.evidence_id,),
    )
    interaction = DataAccess(
        source_object_id=query_object.object_id,
        operation=DataOperation.READ,
        scope=InteractionScope.EXTERNAL,
        connection_id=query_connection.connection_id,
        datasource_id=datasource.datasource_id,
        catalog="TrustCatalog",
        database="CorporateTrust",
        schema_name="dbo",
        object_name="Deal",
        evidence_ids=(query_evidence.evidence_id,),
    )
    source_node = DependencyNode(
        application_id=application_id,
        kind=DependencyNodeKind.ACCESS_OBJECT,
        label="qryShared",
        artifact_id=artifact.artifact_id,
        object_id=query_object.object_id,
    )
    target_node = DependencyNode(
        application_id=application_id,
        kind=DependencyNodeKind.DATASOURCE,
        label="CorporateTrust",
        datasource_id=datasource.datasource_id,
    )
    edge = DependencyEdge(
        source_node_id=source_node.node_id,
        target_node_id=target_node.node_id,
        relationship=DependencyRelation.READS,
        operation=DataOperation.READ,
        evidence_ids=(query_evidence.evidence_id,),
    )
    return ApplicationEvidenceBundle(
        application_id=application_id,
        application_name=application_name,
        inventory_record_ids=(f"inventory-{application_id}",),
        generated_at=NOW,
        artifacts=(artifact,),
        objects=(query_object, table_object),
        connections=(query_connection, table_connection),
        tables=(
            AccessTableEvidence(
                object_id=table_object.object_id,
                is_linked=True,
                source_table_name="dbo.Deal",
                connection_id=table_connection.connection_id,
                evidence_ids=(table_evidence.evidence_id,),
            ),
        ),
        queries=(
            QueryEvidence(
                object_id=query_object.object_id,
                query_kind=QueryKind.PASS_THROUGH,
                sanitized_sql="SELECT * FROM dbo.Deal",
                returns_records=True,
                connection_kind=ConnectionKind.PASS_THROUGH_QUERY,
                connection_id=query_connection.connection_id,
                evidence_ids=(query_evidence.evidence_id,),
            ),
        ),
        datasources=(datasource,),
        interactions=(interaction,),
        dependency_nodes=(source_node, target_node),
        dependency_edges=(edge,),
        evidence=(query_evidence, table_evidence),
        owner_claims=(
            OwnerClaim(
                application_id=application_id,
                field="business_owner_note",
                value="@requires review",
                source="owner_context.csv",
            ),
        ),
        unresolved_references=(
            UnresolvedReference(
                source_object_id=query_object.object_id,
                reference="Forms!frmRuntime!Filter",
                reason="Runtime value cannot be resolved statically",
                evidence_ids=(query_evidence.evidence_id,),
            ),
        ),
        coverage=ExtractionCoverage(
            primary_artifact_count=1,
            complete_artifact_count=int(status == ExtractionStatus.COMPLETE),
            partial_artifact_count=int(status == ExtractionStatus.PARTIAL),
            failed_artifact_count=0,
            discovered_object_count=2,
            extracted_object_count=2,
            warning_count=len(warnings),
            unresolved_reference_count=1,
        ),
    )


def _profile(bundle: ApplicationEvidenceBundle) -> ApplicationProfile:
    return ApplicationProfile(
        application_id=bundle.application_id,
        source_bundle_sha256=evidence_bundle_fingerprint(bundle),
        logical_unit_interpretation_ids=(),
        summary=(
            "+Automated investor reporting"
            if bundle.application_id == "app-1"
            else "Produces investor reports"
        ),
        business_purpose="Investor reporting",
        major_workflows=("Prepare report",),
        capabilities=("Investor reporting",),
        evidence_ids=tuple(item.evidence_id for item in bundle.evidence),
        claim_ids=tuple(item.claim_id for item in bundle.owner_claims),
        provenance=ModelProvenance(
            model_manifest_sha256=MODEL_DIGEST,
            prompt_version="qwen-prompts-v2",
            output_schema_version="qwen-output-v2",
            inference_library_version="5.17.0",
        ),
    )


def _portfolio(
    profiles: tuple[ApplicationProfile, ...],
    first: ApplicationEvidenceBundle,
    second: ApplicationEvidenceBundle,
) -> PortfolioInterpretation:
    return PortfolioInterpretation(
        application_interpretation_sha256s=tuple(
            interpretation_fingerprint(item) for item in profiles
        ),
        findings=(
            PortfolioFinding(
                kind=InterpretationKind.PORTFOLIO_OPPORTUNITY,
                title="Shared reporting data",
                narrative="Review the shared endpoint and data object for reuse.",
                application_ids=(first.application_id, second.application_id),
                evidence_ids=(
                    first.evidence[0].evidence_id,
                    second.evidence[0].evidence_id,
                ),
                confidence=Confidence.MEDIUM,
            ),
        ),
        provenance=ModelProvenance(
            model_manifest_sha256=MODEL_DIGEST,
            prompt_version="qwen-portfolio-v2",
            output_schema_version="qwen-portfolio-output-v2",
            inference_library_version="5.17.0",
        ),
    )


def _inventory_exclusion() -> InventoryExclusion:
    return InventoryExclusion(
        application_id="inventory-spreadsheet",
        application_name="Unsupported spreadsheet row",
        filename="input.xlsx",
        source_locator=r"\\fileserver\inventory\input.xlsx",
        reason="unsupported primary format",
    )
