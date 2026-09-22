from pathlib import Path

from openpyxl import load_workbook

from portfolio_analyzer.models import ArtifactStatus, Confidence, Evidence, StagedArtifact
from portfolio_analyzer.reporting.writers import write_csv, write_executive_pdf, write_workbook


def test_report_writers_create_outputs(tmp_path: Path) -> None:
    artifact = StagedArtifact(
        tool_inventory_id="1",
        original_source_path=Path("source.accdb"),
        filename="source.accdb",
        extension=".accdb",
        status=ArtifactStatus.FAILED,
        error="missing",
    )
    workbook = tmp_path / "Portfolio_Analysis.xlsx"
    write_workbook(workbook, [], [artifact], [], [], [], [])
    sheets = load_workbook(workbook).sheetnames
    assert "Portfolio Summary" in sheets
    assert "Applications" in sheets
    assert "Recommendations" in sheets
    csv_path = tmp_path / "evidence.csv"
    write_csv(csv_path, [{"Tool": "1", "Finding": "test"}])
    assert "Finding" in csv_path.read_text()
    pdf_path = tmp_path / "Portfolio_Analysis.pdf"
    write_executive_pdf(pdf_path, [], [artifact], [])
    assert pdf_path.read_bytes().startswith(b"%PDF")


def test_csv_escapes_formula_like_text_and_keeps_empty_headers(tmp_path: Path) -> None:
    unsafe_path = tmp_path / "unsafe.csv"
    write_csv(unsafe_path, [{"value": "=HYPERLINK(\"bad\")"}])
    assert "'=HYPERLINK" in unsafe_path.read_text()

    empty_path = tmp_path / "empty.csv"
    write_csv(empty_path, [], headers=["first", "second"])
    assert empty_path.read_text().strip() == "first,second"


def test_report_writers_remove_illegal_spreadsheet_characters(tmp_path: Path) -> None:
    evidence = Evidence(
        tool_inventory_id="1",
        artifact_path="source.accdb",
        object_type="module",
        object_name="Example",
        text="UPDATE table\x00 SET value = 'kept'\x0b;",
        confidence=Confidence.HIGH,
    )
    workbook_path = tmp_path / "Portfolio_Analysis.xlsx"
    write_workbook(workbook_path, [], [], [evidence], [], [], [])

    workbook = load_workbook(workbook_path)
    assert workbook["Evidence"]["F2"].value == "UPDATE table SET value = 'kept';"

    csv_path = tmp_path / "evidence.csv"
    write_csv(csv_path, [{"evidence": "UPDATE table\x00 SET value = 'kept'\x0b;"}])
    assert "\x00" not in csv_path.read_text()
    assert "\x0b" not in csv_path.read_text()
