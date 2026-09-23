"""Excel, CSV, and PDF outputs from normalized static-analysis results."""

from __future__ import annotations

import csv
import hashlib
import re
from collections import Counter
from collections.abc import Iterable, Sequence
from html import escape
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from reportlab.graphics.shapes import Drawing, Line, Rect, String
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from portfolio_analyzer.models import (
    AnalysisCoverage,
    CapabilityFinding,
    Datasource,
    Dependency,
    Evidence,
    InventoryRecord,
    Recommendation,
    ReviewDecision,
    SemanticPortfolioState,
    StagedArtifact,
)
from portfolio_analyzer.semantic.review import REVIEW_HEADERS

_ILLEGAL_SPREADSHEET_CHARACTERS = re.compile(r"[\x00-\x08\x0b-\x0c\x0e-\x1f]")
_PDF_TABLE_TRUNCATION_SUFFIX = "... [truncated; see workbook]"


def write_csv(
    path: Path,
    rows: Iterable[dict[str, Any]],
    *,
    headers: Sequence[str] | None = None,
) -> None:
    records = list(rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = list(headers or dict.fromkeys(key for row in records for key in row))
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(
            {key: _spreadsheet_safe(value) for key, value in record.items()} for record in records
        )


def write_workbook(
    path: Path,
    inventory: list[InventoryRecord],
    artifacts: list[StagedArtifact],
    evidence: list[Evidence],
    datasources: list[Datasource],
    dependencies: list[Dependency],
    capabilities: list[CapabilityFinding],
    *,
    recommendations: list[Recommendation] | None = None,
    coverage: list[AnalysisCoverage] | None = None,
    semantic: SemanticPortfolioState | None = None,
    semantic_status: str = "not run",
    review_decisions: dict[str, ReviewDecision] | None = None,
) -> None:
    recommendations = recommendations or []
    coverage = coverage or []
    review_decisions = review_decisions or {}
    application_names = {item.tool_inventory_id: item.tool_name for item in inventory}
    workbook = Workbook()
    workbook.remove(workbook.active)
    analyzed = (
        sum(item.analysis_status == "complete" for item in coverage)
        if coverage
        else len({item.tool_inventory_id for item in evidence})
    )
    extraction_attention = (
        sum(item.extraction_status in {"failed", "complete_with_warnings"} for item in coverage)
        if coverage
        else sum(item.is_primary and bool(item.error) for item in artifacts)
    )
    _portfolio_summary_sheet(
        workbook,
        inventory_count=len({item.tool_inventory_id for item in inventory}),
        staged_count=sum(item.is_primary and item.status.value == "staged" for item in artifacts),
        analyzed_count=analyzed,
        extraction_attention=extraction_attention,
        evidence_count=len(evidence),
        datasource_count=len(datasources),
        dependency_count=len(dependencies),
        recommendation_count=len(recommendations),
        coverage=coverage,
        dependencies=dependencies,
        semantic=semantic,
        semantic_status=semantic_status,
    )
    _semantic_workbook_sheets(
        workbook,
        inventory,
        coverage,
        semantic,
        semantic_status=semantic_status,
        review_decisions=review_decisions,
    )
    if coverage:
        _sheet(
            workbook,
            "Analysis Coverage",
            [
                [
                    "EUC Name",
                    "Primary File",
                    "Staging",
                    "Extraction",
                    "Analysis",
                    "Objects",
                    "Warnings",
                    "Evidence",
                    "Data Sources",
                    "Dependencies",
                    "Capabilities",
                    "Notes",
                ],
                *[
                    [
                        item.tool_name,
                        item.inventory_filename,
                        item.staging_status,
                        item.extraction_status,
                        item.analysis_status,
                        item.extracted_object_count,
                        item.extraction_warning_count,
                        item.evidence_count,
                        item.datasource_count,
                        item.dependency_count,
                        item.capability_count,
                        " | ".join(item.notes),
                    ]
                    for item in coverage
                ],
            ],
        )
    _sheet(
        workbook,
        "Applications",
        [
            [
                "EUC Name",
                "Inventory File Name",
                "Stated Description",
                "Original Source Path",
            ],
            *[
                [
                    item.tool_name,
                    item.inventory_filename,
                    item.stated_description or "",
                    str(item.filepath),
                ]
                for item in inventory
            ],
        ],
    )
    _sheet(
        workbook,
        "Supporting Files",
        [
            [
                "EUC Name",
                "Original Path",
                "Local Staged Path",
                "SHA-256",
                "Status",
                "Error",
            ],
            *[
                [
                    _application_name(a.tool_inventory_id, application_names),
                    str(a.original_source_path),
                    str(a.local_staged_path or ""),
                    a.sha256 or "",
                    a.status,
                    a.error or "",
                ]
                for a in artifacts
            ],
        ],
    )
    _sheet(
        workbook,
        "Data Sources",
        [
            [
                "EUC Name",
                "Platform",
                "Server",
                "Database",
                "Schema",
                "Object",
                "Operation",
                "Confidence",
                "Evidence Items",
                "Connection Summary",
            ],
            *[
                [
                    _application_name(d.tool_inventory_id, application_names),
                    d.platform,
                    d.server or "",
                    d.database or "",
                    d.schema_name or "",
                    d.object_name or "",
                    d.operation,
                    d.confidence,
                    len(d.evidence),
                    d.connection_summary or "",
                ]
                for d in datasources
            ],
        ],
    )
    _sheet(
        workbook,
        "Dependencies",
        [
            ["EUC Name", "Source", "Target", "Type", "Operation", "Confidence"],
            *[
                [
                    _application_name(d.tool_inventory_id, application_names),
                    d.source,
                    d.target,
                    d.dependency_type,
                    d.operation,
                    d.confidence,
                ]
                for d in dependencies
            ],
        ],
    )
    _sheet(
        workbook,
        "Application Capabilities",
        [
            ["EUC Name", "Capability", "Layer", "Confidence", "Evidence Items"],
            *[
                [
                    _application_name(c.tool_inventory_id, application_names),
                    c.capability,
                    c.layer,
                    c.confidence,
                    len(c.evidence),
                ]
                for c in capabilities
            ],
        ],
    )
    _sheet(
        workbook,
        "Recommendations",
        [
            [
                "Category",
                "Title",
                "Confidence",
                "Affected EUC Names",
                "Application Count",
                "Evidence Items",
                "Rationale",
            ],
            *[
                [
                    item.category,
                    item.title,
                    item.confidence,
                    ", ".join(
                        _application_name(tool_id, application_names)
                        for tool_id in item.affected_tool_ids
                    ),
                    len(item.affected_tool_ids),
                    len(item.evidence),
                    item.rationale,
                ]
                for item in recommendations
            ],
        ],
    )
    _sheet(
        workbook,
        "Evidence",
        [
            [
                "EUC Name",
                "Artifact",
                "Object Type",
                "Object",
                "Location",
                "Evidence",
                "Inference",
                "Confidence",
            ],
            *[
                [
                    _application_name(e.tool_inventory_id, application_names),
                    e.artifact_path,
                    e.object_type,
                    e.object_name,
                    e.location or "",
                    e.text,
                    e.inference or "",
                    e.confidence,
                ]
                for e in evidence
            ],
        ],
    )
    _sheet(
        workbook,
        "Extraction Errors",
        [
            ["EUC Name", "Original Source Path", "Error"],
            *[
                [
                    _application_name(a.tool_inventory_id, application_names),
                    str(a.original_source_path),
                    a.error or "",
                ]
                for a in artifacts
                if a.error
            ],
        ],
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _portfolio_summary_sheet(
    workbook: Workbook,
    *,
    inventory_count: int,
    staged_count: int,
    analyzed_count: int,
    extraction_attention: int,
    evidence_count: int,
    datasource_count: int,
    dependency_count: int,
    recommendation_count: int,
    coverage: list[AnalysisCoverage],
    dependencies: list[Dependency],
    semantic: SemanticPortfolioState | None,
    semantic_status: str,
) -> None:
    sheet = workbook.create_sheet("Portfolio Summary")
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 75
    if semantic is not None and semantic.metadata.run_mode == "quick":
        sheet.merge_cells("A1:F1")
        sheet["A1"] = "TEST ONLY — QUICK SEMANTIC MODE — NOT FOR PRODUCTION ACCEPTANCE"
        sheet["A1"].fill = PatternFill("solid", fgColor="A33B35")
        sheet["A1"].font = Font(name="Arial", bold=True, color="FFFFFF")
        sheet["A1"].alignment = Alignment(horizontal="center")
    sheet.merge_cells("A2:F2")
    sheet["A2"] = "Access Portfolio Intelligence"
    sheet["A2"].font = Font(name="Arial", size=16, bold=True, color="123047")
    sheet["A3"] = "Observed coverage, semantic interpretation, and modernization proposals"
    sheet["A3"].font = Font(name="Arial", size=10, italic=True, color="53636D")
    sheet["A5"] = "Portfolio measure"
    sheet["B5"] = "Value"
    measures = [
        ("Inventory applications", inventory_count),
        ("Successfully staged primary artifacts", staged_count),
        ("Applications analyzed", analyzed_count),
        ("Extraction failures or warnings", extraction_attention),
        ("Evidence items", evidence_count),
        ("Normalized datasource interactions", datasource_count),
        ("External dependencies", dependency_count),
        ("Evidence-backed recommendations", recommendation_count),
        ("Semantic profiles", len(semantic.applications) if semantic else 0),
        (
            "Target architecture components",
            len(semantic.architecture.components) if semantic else 0,
        ),
    ]
    for row_index, (label, value) in enumerate(measures, start=6):
        sheet.cell(row_index, 1, label)
        sheet.cell(row_index, 2, value)
        sheet.cell(row_index, 2).number_format = "#,##0"

    sheet["D5"] = "Analysis status"
    sheet["E5"] = "Applications"
    coverage_counts = Counter(item.analysis_status for item in coverage)
    if not coverage_counts:
        coverage_counts["not available"] = inventory_count
    for row_index, (status, count) in enumerate(sorted(coverage_counts.items()), start=6):
        sheet.cell(row_index, 4, status.replace("_", " ").title())
        sheet.cell(row_index, 5, count)

    sheet["A18"] = "Semantic analysis"
    sheet["B18"] = semantic_status
    sheet.merge_cells("B18:C18")
    sheet["B18"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[18].height = 30
    sheet["A19"] = "Interpretation boundary"
    sheet["B19"] = (
        "Observed facts, owner claims, and model proposals remain separate. "
        "Proposals require review before adoption."
    )
    sheet.merge_cells("B19:C20")
    sheet["B19"].alignment = Alignment(wrap_text=True, vertical="top")

    archetypes: Counter[str] = Counter()
    if semantic:
        sheet["D18"] = "Application archetype"
        sheet["E18"] = "Applications"
        archetypes = Counter(item.primary_archetype for item in semantic.applications)
        for row_index, (label, count) in enumerate(
            sorted(archetypes.items(), key=lambda item: (-item[1], item[0])), start=19
        ):
            sheet.cell(row_index, 4, label.title())
            sheet.cell(row_index, 5, count)

    application_names = {item.tool_inventory_id: item.tool_name for item in coverage}
    dependency_hotspots = Counter(item.tool_inventory_id for item in dependencies)
    _summary_bar_chart(
        sheet,
        "Analysis coverage",
        {key.replace("_", " ").title(): value for key, value in coverage_counts.items()},
        start_column=24,
        anchor="G5",
    )
    _summary_bar_chart(
        sheet,
        "Application archetypes",
        dict(archetypes) if semantic else {},
        start_column=26,
        anchor="M5",
    )
    _summary_bar_chart(
        sheet,
        "Capability clusters",
        {item.label: len(item.application_ids) for item in semantic.clusters} if semantic else {},
        start_column=28,
        anchor="G23",
    )
    _summary_bar_chart(
        sheet,
        "Dependency hotspots",
        {
            application_names.get(tool_id, tool_id): count
            for tool_id, count in dependency_hotspots.most_common(10)
        },
        start_column=30,
        anchor="M23",
    )
    _summary_bar_chart(
        sheet,
        "Migration waves",
        {
            f"Wave {item.wave}": len(item.application_ids)
            for item in semantic.architecture.migration_waves
        }
        if semantic
        else {},
        start_column=32,
        anchor="G41",
    )

    _style_table_region(sheet, 5, 1, 5 + len(measures), 2)
    _style_table_region(sheet, 5, 4, 5 + len(coverage_counts), 5)
    if semantic:
        _style_table_region(sheet, 18, 4, 18 + len(archetypes), 5)
    sheet.column_dimensions["A"].width = 39
    sheet.column_dimensions["B"].width = 26
    sheet.column_dimensions["C"].width = 3
    sheet.column_dimensions["D"].width = 28
    sheet.column_dimensions["E"].width = 16
    sheet.column_dimensions["F"].width = 3
    sheet.freeze_panes = "A5"


def _summary_bar_chart(
    sheet: Any,
    title: str,
    values: dict[str, int],
    *,
    start_column: int,
    anchor: str,
) -> None:
    if not values:
        return
    ordered = sorted(values.items(), key=lambda item: (-item[1], item[0]))
    sheet.cell(1, start_column, title)
    sheet.cell(1, start_column + 1, "Applications")
    for row_index, (label, value) in enumerate(ordered, start=2):
        sheet.cell(row_index, start_column, label)
        sheet.cell(row_index, start_column + 1, value)
    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = title
    chart.height = 4.4
    chart.width = 8.3
    chart.legend = None
    chart.y_axis.title = "Category"
    chart.x_axis.title = "Applications"
    data = Reference(
        sheet,
        min_col=start_column + 1,
        min_row=1,
        max_row=1 + len(ordered),
    )
    categories = Reference(
        sheet,
        min_col=start_column,
        min_row=2,
        max_row=1 + len(ordered),
    )
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    sheet.add_chart(chart, anchor)
    sheet.column_dimensions[get_column_letter(start_column)].hidden = True
    sheet.column_dimensions[get_column_letter(start_column + 1)].hidden = True


def _semantic_workbook_sheets(
    workbook: Workbook,
    inventory: list[InventoryRecord],
    coverage: list[AnalysisCoverage],
    semantic: SemanticPortfolioState | None,
    *,
    semantic_status: str,
    review_decisions: dict[str, ReviewDecision],
) -> None:
    names = {item.tool_inventory_id: item.tool_name for item in inventory}
    unique_ids = list(dict.fromkeys(item.tool_inventory_id for item in inventory))
    coverage_by_id: dict[str, list[AnalysisCoverage]] = {}
    for item in coverage:
        coverage_by_id.setdefault(item.tool_inventory_id, []).append(item)
    profiles = {item.tool_inventory_id: item for item in semantic.applications} if semantic else {}
    mappings = (
        {item.tool_inventory_id: item for item in semantic.architecture.mappings}
        if semantic
        else {}
    )
    components = (
        {item.component_id: item for item in semantic.architecture.components} if semantic else {}
    )
    _sheet(
        workbook,
        "Application Portfolio",
        [
            [
                "EUC Name",
                "Analysis Coverage",
                "Deterministic Code Coverage",
                "Inspected Code Segments",
                "Business Purpose",
                "Application Archetype",
                "Proposed Disposition",
                "Migration Wave",
                "Semantic Confidence",
                "Review Status",
                "Summary",
                "Open Questions",
            ],
            *[
                _application_portfolio_row(
                    tool_id,
                    names,
                    coverage_by_id,
                    profiles.get(tool_id),
                    mappings.get(tool_id),
                )
                for tool_id in unique_ids
            ],
        ],
    )
    architecture_rows = (
        [
            [
                "Component ID",
                "Track",
                "Component",
                "Type",
                "Platform Service",
                "Description",
                "Applications",
                "Clusters",
                "Confidence",
                "Review Status",
                "Evidence IDs",
                "Claim IDs",
            ],
            *[
                [
                    item.component_id,
                    item.track.replace("_", " ").title(),
                    item.name,
                    item.component_type.replace("_", " ").title(),
                    item.platform_service or "",
                    item.description,
                    ", ".join(
                        names.get(tool_id, "Unknown EUC") for tool_id in item.application_ids
                    ),
                    " | ".join(item.cluster_ids),
                    item.confidence.value,
                    item.review_status,
                    " | ".join(item.evidence_ids),
                    " | ".join(item.claim_ids),
                ]
                for item in semantic.architecture.components
            ],
        ]
        if semantic
        else [["Status", "Detail"], ["Unavailable", semantic_status]]
    )
    _sheet(workbook, "Target Architecture", architecture_rows)
    crosswalk_rows: list[list[object]] = [
        [
            "Mapping ID",
            "EUC Name",
            "Disposition",
            "Target Components",
            "Wave",
            "Rationale",
            "Prerequisites",
            "Confidence",
            "Review Status",
            "Evidence IDs",
            "Claim IDs",
        ]
    ]
    if semantic:
        for mapping_item in semantic.architecture.mappings:
            crosswalk_rows.append(
                [
                    mapping_item.mapping_id,
                    names.get(mapping_item.tool_inventory_id, "Unknown EUC"),
                    mapping_item.disposition,
                    ", ".join(
                        components[component_id].name
                        for component_id in mapping_item.target_component_ids
                        if component_id in components
                    ),
                    mapping_item.wave,
                    mapping_item.rationale,
                    " | ".join(mapping_item.prerequisites),
                    mapping_item.confidence.value,
                    mapping_item.review_status,
                    " | ".join(mapping_item.evidence_ids),
                    " | ".join(mapping_item.claim_ids),
                ]
            )
    else:
        crosswalk_rows = [["Status", "Detail"], ["Unavailable", semantic_status]]
    _sheet(workbook, "App-Target Crosswalk", crosswalk_rows)
    migration_rows: list[list[object]] = (
        [
            ["Wave", "Name", "Purpose", "Applications", "Application Count", "Prerequisites"],
            *[
                [
                    item.wave,
                    item.name,
                    item.purpose,
                    ", ".join(
                        names.get(tool_id, "Unknown EUC") for tool_id in item.application_ids
                    ),
                    len(item.application_ids),
                    " | ".join(item.prerequisites),
                ]
                for item in semantic.architecture.migration_waves
            ],
        ]
        if semantic
        else [["Status", "Detail"], ["Unavailable", semantic_status]]
    )
    _sheet(workbook, "Migration Roadmap", migration_rows)
    finding_rows = (
        [
            [
                "Finding ID",
                "EUC Name",
                "Category",
                "Label",
                "Description",
                "Confidence",
                "Review Status",
                "Evidence IDs",
                "Claim IDs",
            ],
            *[
                [
                    finding.finding_id,
                    names.get(profile.tool_inventory_id, "Unknown EUC"),
                    finding.category.replace("_", " ").title(),
                    finding.label,
                    finding.description,
                    finding.confidence.value,
                    finding.review_status,
                    " | ".join(finding.evidence_ids),
                    " | ".join(finding.claim_ids),
                ]
                for profile in semantic.applications
                for finding in profile.findings
            ],
        ]
        if semantic
        else [["Status", "Detail"], ["Unavailable", semantic_status]]
    )
    _sheet(workbook, "Semantic Findings", finding_rows)
    _review_queue_sheet(workbook, semantic, names, review_decisions, semantic_status)
    _method_sheet(workbook, semantic, semantic_status)


def _application_portfolio_row(
    tool_id: str,
    names: dict[str, str],
    coverage: dict[str, list[AnalysisCoverage]],
    profile: object,
    mapping: object,
) -> list[object]:
    semantic_profile = profile if hasattr(profile, "primary_archetype") else None
    target_mapping = mapping if hasattr(mapping, "disposition") else None
    semantic_coverage = getattr(semantic_profile, "semantic_coverage", None)
    coverage_items = coverage.get(tool_id, [])
    coverage_label = (
        "Complete"
        if coverage_items
        and all(
            item.analysis_status == "complete" and item.extraction_status == "complete"
            for item in coverage_items
        )
        else "Attention required"
    )
    return [
        names.get(tool_id, "Unknown EUC"),
        coverage_label,
        (
            "Complete"
            if getattr(semantic_coverage, "complete_code_coverage", False)
            else "Sampled or unavailable"
        ),
        (
            f"{semantic_coverage.code_segments_inspected}/"
            f"{semantic_coverage.code_segments_available}"
            if semantic_coverage is not None
            else "0/0"
        ),
        getattr(semantic_profile, "business_purpose", "Unknown"),
        getattr(semantic_profile, "primary_archetype", "unknown"),
        getattr(target_mapping, "disposition", "investigate"),
        getattr(target_mapping, "wave", 0),
        getattr(getattr(semantic_profile, "confidence", None), "value", "low"),
        getattr(target_mapping, "review_status", "pending"),
        getattr(semantic_profile, "summary", "Semantic profile unavailable."),
        " | ".join(getattr(semantic_profile, "open_questions", [])),
    ]


def _review_queue_sheet(
    workbook: Workbook,
    semantic: SemanticPortfolioState | None,
    names: dict[str, str],
    decisions: dict[str, ReviewDecision],
    semantic_status: str,
) -> None:
    header: list[object] = list(REVIEW_HEADERS)
    rows: list[list[object]] = [header]
    if semantic:
        for profile in semantic.applications:
            for finding in profile.findings:
                rows.append(
                    _review_row(
                        finding.finding_id,
                        "Semantic finding",
                        names.get(profile.tool_inventory_id, "Unknown EUC"),
                        finding.label,
                        finding.confidence.value,
                        finding.evidence_ids,
                        finding.claim_ids,
                        decisions,
                    )
                )
        for component in semantic.architecture.components:
            rows.append(
                _review_row(
                    component.component_id,
                    "Architecture component",
                    ", ".join(
                        names.get(tool_id, "Unknown EUC") for tool_id in component.application_ids
                    ),
                    component.name,
                    component.confidence.value,
                    component.evidence_ids,
                    component.claim_ids,
                    decisions,
                )
            )
        for mapping in semantic.architecture.mappings:
            rows.append(
                _review_row(
                    mapping.mapping_id,
                    "Application disposition",
                    names.get(mapping.tool_inventory_id, "Unknown EUC"),
                    mapping.disposition,
                    mapping.confidence.value,
                    mapping.evidence_ids,
                    mapping.claim_ids,
                    decisions,
                )
            )
        for question in semantic.architecture.open_questions:
            proposal_id = (
                "open_question_" + hashlib.sha256(question.encode("utf-8")).hexdigest()[:20]
            )
            rows.append(
                _review_row(
                    proposal_id,
                    "Open architecture decision",
                    "Portfolio",
                    question,
                    "low",
                    [],
                    [],
                    decisions,
                )
            )
    else:
        rows.append(["", "Status", "", semantic_status, "", "", "", "", "", "", ""])
    _sheet(workbook, "Review Queue", rows)
    sheet = workbook["Review Queue"]
    if sheet.max_row >= 2:
        validation = DataValidation(type="list", formula1='"Accept,Edit,Reject"', allow_blank=True)
        validation.error = "Choose Accept, Edit, or Reject."
        validation.errorTitle = "Invalid review decision"
        sheet.add_data_validation(validation)
        validation.add(f"H2:H{sheet.max_row}")
        red_fill = PatternFill("solid", fgColor="F7DEDB")
        amber_fill = PatternFill("solid", fgColor="FFF0CF")
        green_fill = PatternFill("solid", fgColor="DCEFE4")
        sheet.conditional_formatting.add(
            f"H2:H{sheet.max_row}", FormulaRule(formula=['H2="Reject"'], fill=red_fill)
        )
        sheet.conditional_formatting.add(
            f"H2:H{sheet.max_row}", FormulaRule(formula=['H2="Edit"'], fill=amber_fill)
        )
        sheet.conditional_formatting.add(
            f"H2:H{sheet.max_row}", FormulaRule(formula=['H2="Accept"'], fill=green_fill)
        )


def _review_row(
    proposal_id: str,
    proposal_type: str,
    euc_name: str,
    proposed_value: str,
    confidence: str,
    evidence_ids: list[str],
    claim_ids: list[str],
    decisions: dict[str, ReviewDecision],
) -> list[object]:
    decision = decisions.get(proposal_id)
    return [
        proposal_id,
        proposal_type,
        euc_name,
        proposed_value,
        confidence,
        " | ".join(evidence_ids),
        " | ".join(claim_ids),
        decision.decision if decision else "",
        decision.edited_value or "" if decision else "",
        decision.reviewer or "" if decision else "",
        decision.notes or "" if decision else "",
    ]


def _method_sheet(
    workbook: Workbook,
    semantic: SemanticPortfolioState | None,
    semantic_status: str,
) -> None:
    rows: list[list[object]] = [
        ["Item", "Value"],
        ["Semantic status", semantic_status],
        ["Interpretation boundary", "AI results are reviewable proposals, not observed facts."],
        ["Network boundary", "Integrity-verified in-process model; offline inference only."],
        [
            "Confidence policy",
            "System-derived from cited evidence, claims, and extraction coverage.",
        ],
    ]
    if semantic:
        metadata = semantic.metadata
        rows.extend(
            [
                ["Semantic version", metadata.semantic_version],
                ["Run mode", metadata.run_mode],
                ["Run status", metadata.run_status],
                [
                    "Maximum semantic objects per application",
                    metadata.max_objects_per_application or "All",
                ],
                ["Schema version", metadata.semantic_schema_version],
                ["Prompt version", metadata.prompt_version],
                ["Static analysis version", metadata.static_analysis_version],
                ["Model repository", metadata.model_repo_id],
                ["Model revision", metadata.model_revision],
                ["Model manifest SHA-256", metadata.model_manifest_sha256],
                ["Model architecture", metadata.model_architecture],
                ["Model license", metadata.model_license],
                [
                    "Inference runtime",
                    f"{metadata.inference_library} {metadata.inference_library_version}",
                ],
                ["Similarity version", metadata.deterministic_similarity_version],
                [
                    "Architecture synthesis",
                    "Local model" if metadata.architecture_model_generation else "Deterministic",
                ],
                ["Generated at", metadata.generated_at.isoformat()],
                ["Semantic errors", len(semantic.errors)],
                ["Deterministic application IRs", len(semantic.application_irs)],
                [
                    "Applications with complete code coverage",
                    sum(
                        profile.semantic_coverage is not None
                        and profile.semantic_coverage.complete_code_coverage
                        for profile in semantic.applications
                    ),
                ],
            ]
        )
    _sheet(workbook, "Method & Provenance", rows)


def _style_table_region(sheet: Any, min_row: int, min_col: int, max_row: int, max_col: int) -> None:
    header_fill = PatternFill("solid", fgColor="123047")
    border = Border(bottom=Side(style="thin", color="D5DFE4"))
    for cell in sheet.iter_cols(min_col=min_col, max_col=max_col, min_row=min_row, max_row=min_row):
        for item in cell:
            item.fill = header_fill
            item.font = Font(name="Arial", bold=True, color="FFFFFF")
            item.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(
        min_row=min_row + 1, max_row=max_row, min_col=min_col, max_col=max_col
    ):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center")


def _sheet(workbook: Workbook, title: str, rows: Sequence[Sequence[object]]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.sheet_view.showGridLines = False
    for row in rows:
        sheet.append([_spreadsheet_safe(value) for value in row])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="123047")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(
            max(len(str(cell.value or "")) for cell in column) + 2, 60
        )
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_executive_pdf(
    path: Path,
    inventory: list[InventoryRecord],
    artifacts: list[StagedArtifact],
    capabilities: list[CapabilityFinding],
    *,
    recommendations: list[Recommendation] | None = None,
    evidence: list[Evidence] | None = None,
    datasources: list[Datasource] | None = None,
    dependencies: list[Dependency] | None = None,
    coverage: list[AnalysisCoverage] | None = None,
    semantic: SemanticPortfolioState | None = None,
    semantic_status: str = "not run",
) -> None:
    """Render an evidence-backed, multi-page executive and architecture report."""
    path.parent.mkdir(parents=True, exist_ok=True)
    recommendations = recommendations or []
    evidence = evidence or []
    datasources = datasources or []
    dependencies = dependencies or []
    coverage = coverage or []
    application_names = {item.tool_inventory_id: item.tool_name for item in inventory}
    successful = sum(a.status.value == "staged" and a.is_primary for a in artifacts)
    failed = sum(a.status.value == "failed" and a.is_primary for a in artifacts)
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="ReportTitle",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=24,
            leading=29,
            textColor=colors.HexColor("#123047"),
            spaceAfter=18,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ReportHeading",
            parent=styles["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=16,
            leading=20,
            textColor=colors.HexColor("#123047"),
            spaceBefore=4,
            spaceAfter=12,
        )
    )
    styles.add(
        ParagraphStyle(name="ReportBody", parent=styles["BodyText"], fontSize=9.5, leading=14)
    )
    styles.add(
        ParagraphStyle(
            name="TestOnly",
            parent=styles["BodyText"],
            fontName="Helvetica-Bold",
            fontSize=11,
            leading=15,
            textColor=colors.HexColor("#A33B35"),
            borderColor=colors.HexColor("#A33B35"),
            borderWidth=1,
            borderPadding=8,
            spaceAfter=14,
        )
    )
    document = SimpleDocTemplate(
        str(path),
        pagesize=letter,
        rightMargin=0.62 * inch,
        leftMargin=0.62 * inch,
        topMargin=0.7 * inch,
        bottomMargin=0.65 * inch,
        title="Access Portfolio Analysis",
        author="Access Portfolio Analyzer",
    )
    story: list[object] = []
    if semantic is not None and semantic.metadata.run_mode == "quick":
        story.append(
            Paragraph(
                "TEST ONLY — QUICK SEMANTIC MODE — NOT FOR PRODUCTION ACCEPTANCE",
                styles["TestOnly"],
            )
        )
    story.extend(
        [
            Paragraph("Access Portfolio Analysis", styles["ReportTitle"]),
            Paragraph(
                "Evidence-backed static analysis and modernization assessment", styles["ReportBody"]
            ),
            Spacer(1, 0.25 * inch),
            Paragraph("Executive Summary", styles["ReportHeading"]),
            _table(
                [
                    ["Measure", "Observed value"],
                    ["Inventory applications", str(len(inventory))],
                    ["Successfully staged primary artifacts", str(successful)],
                    ["Failed primary staging attempts", str(failed)],
                    [
                        "Applications with completed static analysis",
                        str(sum(item.analysis_status == "complete" for item in coverage))
                        if coverage
                        else str(len({item.tool_inventory_id for item in evidence})),
                    ],
                    ["Evidence-backed modernization opportunities", str(len(recommendations))],
                ],
                [2.75 * inch, 3.35 * inch],
            ),
            Spacer(1, 0.18 * inch),
            Paragraph(
                " ".join(
                    (
                        "The stated inventory description is retained as a claim.",
                        (
                            "Conclusions are based on static evidence from verified staged "
                            "copies only."
                        ),
                    )
                ),
                styles["ReportBody"],
            ),
            Spacer(1, 0.12 * inch),
            Paragraph(
                (
                    "Recommendations require business validation; they are not automatic rewrite "
                    "or service decisions."
                ),
                styles["ReportBody"],
            ),
            Spacer(1, 0.12 * inch),
            Paragraph(
                f"<b>Semantic analysis:</b> {escape(semantic_status)}",
                styles["ReportBody"],
            ),
        ]
    )
    story.append(PageBreak())
    story.extend(_coverage_section(coverage, styles))
    if capabilities:
        story.append(PageBreak())
        story.extend(_capability_section(capabilities, styles))
    story.append(PageBreak())
    story.extend(_semantic_portfolio_section(semantic, semantic_status, application_names, styles))
    if semantic:
        story.append(PageBreak())
        story.extend(_target_architecture_section(semantic, styles))
        story.append(PageBreak())
        story.extend(_migration_section(semantic, application_names, styles))
    if datasources or dependencies:
        story.append(PageBreak())
        story.extend(_dependency_section(datasources, dependencies, styles))
    if recommendations:
        story.append(PageBreak())
        story.extend(_recommendation_section(recommendations, application_names, styles))
    story.append(PageBreak())
    story.extend(
        _risk_and_next_steps_section(
            evidence,
            artifacts,
            coverage,
            application_names,
            styles,
        )
    )
    document.build(story, onFirstPage=_footer, onLaterPages=_footer)


def _coverage_section(
    coverage: list[AnalysisCoverage], styles: dict[str, ParagraphStyle]
) -> list[object]:
    content: list[object] = [Paragraph("Analysis Coverage and Confidence", styles["ReportHeading"])]
    if not coverage:
        content.append(
            Paragraph(
                "Pipeline coverage was not available; absence of findings must not be "
                "interpreted as absence of functionality.",
                styles["ReportBody"],
            )
        )
        return content
    counts = Counter(item.analysis_status for item in coverage)
    content.extend(
        [
            Paragraph(
                "Findings are conclusive only for applications with completed analysis. "
                "Extraction warnings may reduce completeness even when analysis ran successfully.",
                styles["ReportBody"],
            ),
            Spacer(1, 0.12 * inch),
            _table(
                [["Analysis state", "Applications"]]
                + [
                    [status.replace("_", " ").title(), str(count)]
                    for status, count in sorted(counts.items())
                ],
                [4.65 * inch, 1.45 * inch],
            ),
        ]
    )
    attention = [
        item
        for item in coverage
        if item.analysis_status != "complete" or item.extraction_warning_count
    ]
    if attention:
        content.extend(
            [
                Spacer(1, 0.18 * inch),
                Paragraph("Applications requiring attention", styles["Heading2"]),
                _table(
                    [["EUC Name", "Primary File", "Extraction", "Analysis", "Warnings"]]
                    + [
                        [
                            item.tool_name,
                            item.inventory_filename,
                            item.extraction_status.replace("_", " "),
                            item.analysis_status.replace("_", " "),
                            str(item.extraction_warning_count),
                        ]
                        for item in attention[:25]
                    ],
                    [1.25 * inch, 1.35 * inch, 1.35 * inch, 1.25 * inch, 0.9 * inch],
                ),
            ]
        )
    return content


def _capability_section(
    capabilities: list[CapabilityFinding], styles: dict[str, ParagraphStyle]
) -> list[object]:
    grouped: dict[str, list[CapabilityFinding]] = {}
    for item in capabilities:
        grouped.setdefault(item.capability, []).append(item)
    confidence_rank = {"high": 3, "medium": 2, "low": 1}
    rows: list[list[str]] = [
        ["Observed technical capability", "Applications", "Evidence", "Confidence"]
    ]
    for name, findings in sorted(
        grouped.items(),
        key=lambda pair: (-len({item.tool_inventory_id for item in pair[1]}), pair[0]),
    ):
        confidence = min(
            (item.confidence.value for item in findings),
            key=lambda value: confidence_rank[value],
        )
        rows.append(
            [
                name,
                str(len({item.tool_inventory_id for item in findings})),
                str(sum(len(item.evidence) for item in findings)),
                confidence.title(),
            ]
        )
    content: list[object] = [Paragraph("Capability Map", styles["ReportHeading"])]
    if grouped:
        content.extend(
            [
                Paragraph(
                    " ".join(
                        (
                            (
                                "This taxonomy is assembled from observed static signals "
                                "in the analyzed corpus."
                            ),
                            "It does not impose business-domain labels in advance.",
                        )
                    ),
                    styles["ReportBody"],
                ),
                Spacer(1, 0.12 * inch),
                _table(rows, [3.25 * inch, 0.95 * inch, 0.85 * inch, 1.05 * inch]),
            ]
        )
    else:
        content.append(
            Paragraph(
                (
                    "No implementation evidence is available yet. Run Windows extraction "
                    "before drawing conclusions."
                ),
                styles["ReportBody"],
            )
        )
    return content


def _semantic_portfolio_section(
    semantic: SemanticPortfolioState | None,
    semantic_status: str,
    application_names: dict[str, str],
    styles: dict[str, ParagraphStyle],
) -> list[object]:
    content: list[object] = [
        Paragraph("Semantic Portfolio Intelligence", styles["ReportHeading"]),
        Paragraph(
            "AI-generated labels and dispositions are grounded proposals. Observed static "
            "facts and owner claims remain separately traceable in the workbook.",
            styles["ReportBody"],
        ),
        Spacer(1, 0.12 * inch),
    ]
    if semantic is None:
        content.append(
            _table(
                [
                    ["Semantic state", "Coverage impact"],
                    [semantic_status, "Deterministic report only"],
                ],
                [2.3 * inch, 3.8 * inch],
            )
        )
        return content

    archetypes = Counter(item.primary_archetype for item in semantic.applications)
    confidence = Counter(item.confidence.value for item in semantic.applications)
    content.extend(
        [
            _table(
                [
                    ["Measure", "Count"],
                    ["Grounded application profiles", str(len(semantic.applications))],
                    ["Consolidation clusters", str(len(semantic.clusters))],
                    ["Open semantic questions", str(len(semantic.architecture.open_questions))],
                    ["Application-level failures", str(len(semantic.errors))],
                ],
                [4.65 * inch, 1.45 * inch],
            ),
            Spacer(1, 0.16 * inch),
            Paragraph("Application archetypes", styles["Heading2"]),
            _table(
                [["Archetype", "Applications"]]
                + [[name, str(count)] for name, count in archetypes.most_common()],
                [4.65 * inch, 1.45 * inch],
            ),
            Spacer(1, 0.16 * inch),
            Paragraph(
                "Profile confidence: "
                + ", ".join(f"{name} {count}" for name, count in sorted(confidence.items())),
                styles["ReportBody"],
            ),
        ]
    )
    if semantic.clusters:
        rows = [["Cluster", "EUC Names", "Shared capabilities", "Confidence"]]
        for cluster in semantic.clusters[:12]:
            names = [
                application_names.get(tool_id, "Unknown EUC") for tool_id in cluster.application_ids
            ]
            rows.append(
                [
                    cluster.label,
                    ", ".join(names),
                    ", ".join(cluster.shared_capabilities) or "Needs owner validation",
                    cluster.confidence.value.title(),
                ]
            )
        content.extend(
            [
                Spacer(1, 0.2 * inch),
                Paragraph("Consolidation candidates", styles["Heading2"]),
                _table(rows, [1.2 * inch, 1.8 * inch, 2.1 * inch, 1 * inch]),
            ]
        )
    return content


def _target_architecture_section(
    semantic: SemanticPortfolioState,
    styles: dict[str, ParagraphStyle],
) -> list[object]:
    content: list[object] = [
        Paragraph("Proposed Target Architecture", styles["ReportHeading"]),
        Paragraph(
            "The design is modular and reviewable. The Microsoft mapping is restricted to the "
            "approved service catalog; unsupported mappings remain open hosting decisions.",
            styles["ReportBody"],
        ),
    ]
    for track, title in (
        ("vendor_neutral", "Vendor-neutral capability architecture"),
        ("microsoft", "Approved Microsoft implementation track"),
    ):
        components = [
            item
            for item in semantic.architecture.components
            if item.track == track and item.review_status != "rejected"
        ]
        content.extend(
            [
                Spacer(1, 0.18 * inch),
                Paragraph(title, styles["Heading2"]),
            ]
        )
        if not components:
            content.append(
                Paragraph("No supported components were proposed.", styles["ReportBody"])
            )
            continue
        content.append(_architecture_drawing(semantic, track))
        rows = [["Component", "Type / service", "Supporting scope", "Status"]]
        for item in components[:12]:
            type_label = item.component_type.replace("_", " ").title()
            if item.platform_service:
                type_label += f" / {item.platform_service}"
            rows.append(
                [
                    item.name,
                    type_label,
                    f"{len(item.application_ids)} applications; {len(item.cluster_ids)} clusters",
                    f"{item.confidence.value.title()} / {item.review_status.title()}",
                ]
            )
        content.extend(
            [
                Spacer(1, 0.08 * inch),
                _table(rows, [1.7 * inch, 1.7 * inch, 1.55 * inch, 1.15 * inch]),
            ]
        )
    if semantic.architecture.open_questions:
        content.extend(
            [
                Spacer(1, 0.18 * inch),
                Paragraph("Open architecture decisions", styles["Heading2"]),
                Paragraph(
                    "<br/>".join(
                        f"- {escape(question)}"
                        for question in semantic.architecture.open_questions[:12]
                    ),
                    styles["ReportBody"],
                ),
            ]
        )
    return content


def _architecture_drawing(semantic: SemanticPortfolioState, track: str) -> Drawing:
    components = [
        item
        for item in semantic.architecture.components
        if item.track == track and item.review_status != "rejected"
    ][:9]
    width = 6.1 * inch
    height = 2.05 * inch
    drawing = Drawing(width, height)
    if not components:
        return drawing
    columns = 3
    box_width = 1.72 * inch
    box_height = 0.48 * inch
    gap_x = 0.3 * inch
    gap_y = 0.2 * inch
    origin_x = 0.1 * inch
    origin_y = height - box_height - 0.12 * inch
    positions: dict[str, tuple[float, float]] = {}
    for index, component in enumerate(components):
        column = index % columns
        row = index // columns
        x = origin_x + column * (box_width + gap_x)
        y = origin_y - row * (box_height + gap_y)
        positions[component.component_id] = (x, y)
    for relation in semantic.architecture.relations:
        if relation.track != track:
            continue
        source = positions.get(relation.source_component_id)
        target = positions.get(relation.target_component_id)
        if source is None or target is None:
            continue
        drawing.add(
            Line(
                source[0] + box_width / 2,
                source[1] + box_height / 2,
                target[0] + box_width / 2,
                target[1] + box_height / 2,
                strokeColor=colors.HexColor("#91A9B5"),
                strokeWidth=0.75,
            )
        )
    for component in components:
        x, y = positions[component.component_id]
        drawing.add(
            Rect(
                x,
                y,
                box_width,
                box_height,
                rx=5,
                ry=5,
                fillColor=colors.HexColor("#E7F0F4"),
                strokeColor=colors.HexColor("#2F6F8F"),
            )
        )
        drawing.add(
            String(
                x + 6,
                y + box_height - 12,
                _pdf_table_text(component.name, box_width)[:28],
                fontName="Helvetica-Bold",
                fontSize=7.5,
                fillColor=colors.HexColor("#123047"),
            )
        )
        subtitle = component.platform_service or component.component_type.replace("_", " ")
        drawing.add(
            String(
                x + 6,
                y + 7,
                _pdf_table_text(subtitle, box_width)[:32],
                fontName="Helvetica",
                fontSize=6.5,
                fillColor=colors.HexColor("#53636D"),
            )
        )
    return drawing


def _migration_section(
    semantic: SemanticPortfolioState,
    application_names: dict[str, str],
    styles: dict[str, ParagraphStyle],
) -> list[object]:
    content: list[object] = [
        Paragraph("Technical Migration Roadmap", styles["ReportHeading"]),
        Paragraph(
            "Waves describe technical sequencing only; they are not estimates of duration, "
            "staffing, budget, or final business priority.",
            styles["ReportBody"],
        ),
        Spacer(1, 0.12 * inch),
    ]
    waves = semantic.architecture.migration_waves
    if waves:
        rows = [["Wave", "Purpose", "EUC scope", "Prerequisites"]]
        for wave in sorted(waves, key=lambda item: item.wave):
            rows.append(
                [
                    f"{wave.wave}: {wave.name}",
                    wave.purpose,
                    _summarize_application_names(wave.application_ids, application_names),
                    "; ".join(wave.prerequisites) or "None recorded",
                ]
            )
        content.append(_table(rows, [1.1 * inch, 1.65 * inch, 1.75 * inch, 1.6 * inch]))
    else:
        content.append(Paragraph("No migration waves were generated.", styles["ReportBody"]))
    mappings = semantic.architecture.mappings
    if mappings:
        content.extend(
            [
                Spacer(1, 0.2 * inch),
                Paragraph("Application-to-target crosswalk", styles["Heading2"]),
                _table(
                    [["EUC Name", "Disposition", "Wave", "Confidence", "Review"]]
                    + [
                        [
                            application_names.get(item.tool_inventory_id, "Unknown EUC"),
                            item.disposition,
                            str(item.wave),
                            item.confidence.value.title(),
                            item.review_status.title(),
                        ]
                        for item in mappings[:25]
                    ],
                    [2.2 * inch, 1.45 * inch, 0.55 * inch, 0.9 * inch, 1 * inch],
                ),
            ]
        )
    return content


def _dependency_section(
    datasources: list[Datasource], dependencies: list[Dependency], styles: dict[str, ParagraphStyle]
) -> list[object]:
    content: list[object] = [Paragraph("Dependency Analysis", styles["ReportHeading"])]
    if datasources:
        content.append(Paragraph("Most-used datasource interactions", styles["ReportBody"]))
        grouped: dict[tuple[str, str, str, str], list[Datasource]] = {}
        for item in datasources:
            key = (
                item.platform,
                " / ".join(value for value in (item.server, item.database) if value),
                ".".join(value for value in (item.schema_name, item.object_name) if value),
                item.operation,
            )
            grouped.setdefault(key, []).append(item)
        rows = [["Platform", "Server / database", "Object", "Operation", "Apps", "Evidence"]]
        ranked = sorted(
            grouped.items(),
            key=lambda pair: (
                -len({item.tool_inventory_id for item in pair[1]}),
                -sum(len(item.evidence) for item in pair[1]),
                pair[0],
            ),
        )
        rows.extend(
            [
                [
                    platform,
                    location,
                    object_name,
                    operation,
                    str(len({item.tool_inventory_id for item in items})),
                    str(sum(len(item.evidence) for item in items)),
                ]
                for (platform, location, object_name, operation), items in ranked[:20]
            ]
        )
        content.extend(
            [
                Spacer(1, 0.1 * inch),
                _table(
                    rows,
                    [0.8 * inch, 1.45 * inch, 1.35 * inch, 0.7 * inch, 0.6 * inch, 0.7 * inch],
                ),
            ]
        )
    else:
        content.append(
            Paragraph("No normalized datasource interactions were extracted.", styles["ReportBody"])
        )
    content.append(Spacer(1, 0.2 * inch))
    if dependencies:
        by_type = Counter(item.dependency_type for item in dependencies)
        content.extend(
            [
                Paragraph("External dependency types", styles["ReportBody"]),
                Spacer(1, 0.08 * inch),
                _table(
                    [["Dependency type", "Relationships"]]
                    + [[name, str(count)] for name, count in by_type.most_common()],
                    [3.4 * inch, 2.7 * inch],
                ),
            ]
        )
    else:
        content.append(
            Paragraph(
                "No external filesystem or application dependencies were extracted.",
                styles["ReportBody"],
            )
        )
    return content


def _recommendation_section(
    recommendations: list[Recommendation],
    application_names: dict[str, str],
    styles: dict[str, ParagraphStyle],
) -> list[object]:
    content: list[object] = [
        Paragraph("Shared Capability and Consolidation Opportunities", styles["ReportHeading"])
    ]
    if not recommendations:
        content.append(
            Paragraph(
                " ".join(
                    (
                        (
                            "No opportunity crossed the threshold for a shared library, platform "
                            "capability, API, or service."
                        ),
                        (
                            "Recommendations are withheld until recurring implementation "
                            "evidence exists."
                        ),
                    )
                ),
                styles["ReportBody"],
            )
        )
        return content
    content.append(
        Paragraph(
            " ".join(
                (
                    "Each opportunity is derived from repeated static evidence.",
                    (
                        "The appropriate boundary depends on ownership, security, transactions, "
                        "and operations."
                    ),
                )
            ),
            styles["ReportBody"],
        )
    )
    for recommendation in recommendations:
        affected = ", ".join(
            _application_name(tool_id, application_names)
            for tool_id in recommendation.affected_tool_ids
        )
        content.extend(
            [
                Spacer(1, 0.14 * inch),
                Paragraph(escape(recommendation.title), styles["Heading2"]),
                Paragraph(
                    f"<b>Candidate boundary:</b> {escape(recommendation.category)}<br/>"
                    f"<b>Confidence:</b> {escape(recommendation.confidence.value.title())}<br/>"
                    f"<b>EUC names affected:</b> {escape(affected)}<br/>"
                    f"<b>Evidence items:</b> {len(recommendation.evidence)}<br/>"
                    f"{escape(recommendation.rationale)}",
                    styles["ReportBody"],
                ),
            ]
        )
    return content


def _risk_and_next_steps_section(
    evidence: list[Evidence],
    artifacts: list[StagedArtifact],
    coverage: list[AnalysisCoverage],
    application_names: dict[str, str],
    styles: dict[str, ParagraphStyle],
) -> list[object]:
    from portfolio_analyzer.portfolio.recommendations import modernization_risks

    content: list[object] = [
        Paragraph("Modernization Risks and Next Steps", styles["ReportHeading"])
    ]
    risks = modernization_risks(evidence)
    if risks:
        rows = [["Observed pattern", "EUC Names", "Modernization implication"]]
        rows.extend(
            [
                [
                    pattern,
                    _summarize_application_names(tools, application_names),
                    implication,
                ]
                for pattern, tools, implication in risks
            ]
        )
        content.extend(
            [
                Paragraph("Observed technical-debt themes", styles["ReportBody"]),
                Spacer(1, 0.1 * inch),
                _table(rows, [1.35 * inch, 1.25 * inch, 3.5 * inch]),
            ]
        )
    else:
        content.append(
            Paragraph(
                "No technical-debt themes can be supported until extraction yields evidence.",
                styles["ReportBody"],
            )
        )
    failed = [item for item in artifacts if item.is_primary and item.error]
    coverage_attention = [
        item
        for item in coverage
        if item.analysis_status != "complete" or item.extraction_warning_count
    ]
    content.extend(
        [
            Spacer(1, 0.22 * inch),
            Paragraph("Recommended Next Steps", styles["Heading2"]),
            Paragraph(
                "1. Validate a representative sample with application owners.",
                styles["ReportBody"],
            ),
            Paragraph(
                "2. Resolve staging and extraction limitations before inferring absence of "
                "functionality.",
                styles["ReportBody"],
            ),
            Paragraph(
                "3. Review opportunities with business and operational constraints before "
                "choosing a boundary.",
                styles["ReportBody"],
            ),
        ]
    )
    if failed or coverage_attention:
        content.append(Spacer(1, 0.12 * inch))
        content.append(
            Paragraph(
                "Manual review queue: "
                f"{len(coverage_attention) if coverage else len(failed)} application(s) have "
                "incomplete or warning-qualified coverage.",
                styles["ReportBody"],
            )
        )
    return content


def _table(rows: list[list[str]], widths: list[float]) -> Table:
    header_style = ParagraphStyle(
        name="TableHeader",
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=10,
        textColor=colors.white,
    )
    body_style = ParagraphStyle(
        name="TableBody",
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        textColor=colors.black,
    )
    wrapped = [
        [
            Paragraph(
                escape(_pdf_table_text(value, widths[column_index])),
                header_style if row_index == 0 else body_style,
            )
            for column_index, value in enumerate(row)
        ]
        for row_index, row in enumerate(rows)
    ]
    table = Table(wrapped, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#123047")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("LEADING", (0, 0), (-1, -1), 10),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C7D1D8")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EDF3F6")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    return table


def _summarize_application_names(
    tool_ids: list[str], names: dict[str, str], *, limit: int = 8
) -> str:
    resolved = [_application_name(tool_id, names) for tool_id in tool_ids]
    if len(resolved) <= limit:
        return ", ".join(resolved)
    return ", ".join(resolved[:limit]) + f", ... (+{len(resolved) - limit} more; see workbook)"


def _pdf_table_text(value: object, width: float) -> str:
    """Bound executive-summary cells so every table row can fit on a page."""
    text = str(value)
    max_characters = max(80, int(width / inch * 160))
    if len(text) <= max_characters:
        return text
    prefix_length = max_characters - len(_PDF_TABLE_TRUNCATION_SUFFIX)
    return text[:prefix_length].rstrip() + _PDF_TABLE_TRUNCATION_SUFFIX


def _footer(canvas: Any, document: Any) -> None:
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#53636D"))
    canvas.drawString(
        0.62 * inch, 0.38 * inch, "Access Portfolio Analyzer - evidence-backed static analysis"
    )
    canvas.drawRightString(7.88 * inch, 0.38 * inch, f"Page {document.page}")
    canvas.restoreState()


def _spreadsheet_safe(value: object) -> object:
    """Return text that is valid and inert in spreadsheet output."""
    if not isinstance(value, str):
        return value
    value = _ILLEGAL_SPREADSHEET_CHARACTERS.sub("", value)
    if value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def _application_name(tool_inventory_id: str, names: dict[str, str]) -> str:
    return names.get(tool_inventory_id, "Unknown EUC")
