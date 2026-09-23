"""Import and apply human decisions to semantic proposals."""

from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path
from typing import Literal, cast

from openpyxl import load_workbook

from portfolio_analyzer.models import ReviewDecision, SemanticPortfolioState

REVIEW_SHEET = "Review Queue"
REVIEW_HEADERS = [
    "Proposal ID",
    "Proposal Type",
    "EUC Name",
    "Proposed Value",
    "Confidence",
    "Evidence IDs",
    "Claim IDs",
    "Decision",
    "Edited Value",
    "Reviewer",
    "Notes",
]
ReviewChoice = Literal["Accept", "Edit", "Reject"]
ReviewStatus = Literal["pending", "accepted", "edited", "rejected"]


def import_review_workbook(path: Path) -> dict[str, ReviewDecision]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    if REVIEW_SHEET not in workbook.sheetnames:
        raise ValueError(f"Workbook has no '{REVIEW_SHEET}' sheet")
    rows = workbook[REVIEW_SHEET].iter_rows(values_only=True)
    try:
        headers = [str(value or "").strip() for value in next(rows)]
    except StopIteration as exc:
        raise ValueError("Review Queue is empty") from exc
    missing = set(REVIEW_HEADERS) - set(headers)
    if missing:
        raise ValueError(f"Review Queue is missing columns: {', '.join(sorted(missing))}")
    positions = {header: headers.index(header) for header in REVIEW_HEADERS}
    decisions: dict[str, ReviewDecision] = {}
    for row_number, row in enumerate(rows, start=2):
        proposal_id = _cell(row, positions["Proposal ID"])
        decision = _cell(row, positions["Decision"])
        if not proposal_id or not decision:
            continue
        if decision not in {"Accept", "Edit", "Reject"}:
            raise ValueError(f"Invalid decision in Review Queue row {row_number}: {decision}")
        edited_value = _cell(row, positions["Edited Value"]) or None
        if decision == "Edit" and not edited_value:
            raise ValueError(f"Review Queue row {row_number} requires an Edited Value")
        decisions[proposal_id] = ReviewDecision(
            proposal_id=proposal_id,
            decision=cast(ReviewChoice, decision),
            edited_value=edited_value,
            reviewer=_cell(row, positions["Reviewer"]) or None,
            notes=_cell(row, positions["Notes"]) or None,
            reviewed_at=datetime.now(UTC),
        )
    return decisions


def apply_review_decisions(
    state: SemanticPortfolioState, decisions: dict[str, ReviewDecision]
) -> SemanticPortfolioState:
    updated = state.model_copy(deep=True)
    valid_dispositions = {
        "retain/remediate",
        "wrap/integrate",
        "replatform",
        "rebuild",
        "consolidate",
        "retire candidate",
        "investigate",
    }
    for profile in updated.applications:
        for finding in profile.findings:
            decision = decisions.get(finding.finding_id)
            if not decision:
                continue
            finding.review_status = _status(decision)
            if decision.decision == "Edit" and decision.edited_value:
                finding.label = decision.edited_value
    for component in updated.architecture.components:
        decision = decisions.get(component.component_id)
        if not decision:
            continue
        component.review_status = _status(decision)
        if decision.decision == "Edit" and decision.edited_value:
            component.name = decision.edited_value
    for mapping in updated.architecture.mappings:
        decision = decisions.get(mapping.mapping_id)
        if not decision:
            continue
        mapping.review_status = _status(decision)
        if decision.decision == "Edit" and decision.edited_value:
            edited = decision.edited_value.casefold()
            if edited not in valid_dispositions:
                raise ValueError(
                    f"Invalid edited disposition for {mapping.mapping_id}: {decision.edited_value}"
                )
            mapping.disposition = edited  # type: ignore[assignment]
        if decision.decision == "Reject":
            mapping.wave = 0
            mapping.prerequisites = sorted(
                {*mapping.prerequisites, "Replace the rejected architecture mapping"}
            )
    return updated


def _cell(row: tuple[object, ...], index: int) -> str:
    value = row[index] if index < len(row) else None
    if isinstance(value, str) and value.startswith("="):
        raise ValueError("Review Queue decision fields must contain values, not formulas")
    return str(value or "").strip()


def _status(decision: ReviewDecision) -> ReviewStatus:
    return cast(
        ReviewStatus,
        {"Accept": "accepted", "Edit": "edited", "Reject": "rejected"}[decision.decision],
    )
