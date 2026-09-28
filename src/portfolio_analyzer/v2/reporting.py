"""Safe renderers and manifest-last publication for the normalized V2 report model."""

from __future__ import annotations

import csv
import hashlib
import io
import os
import re
import shutil
import tempfile
import zipfile
from datetime import datetime
from enum import StrEnum
from html import escape, unescape
from pathlib import Path, PurePosixPath
from typing import Annotated, Any, Literal, Self

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import Field, ValidationError, field_validator, model_validator
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from portfolio_analyzer.v2.fingerprints import report_fingerprint
from portfolio_analyzer.v2.identity import REDACTED, canonical_json_bytes
from portfolio_analyzer.v2.models import (
    NonEmptyString,
    PortfolioReportModel,
    ReportLineageTarget,
    ReportStatus,
    Sha256,
    StrictModel,
)
from portfolio_analyzer.v2.review import REVIEW_HEADERS

REPORT_PUBLICATION_SCHEMA_VERSION = "report-publication-v2"
LATEST_REPORT_POINTER_SCHEMA_VERSION = "latest-report-pointer-v2"
_SAFE_RUN_ID = re.compile(r"[A-Za-z0-9][A-Za-z0-9_.-]{0,127}\Z")
_ILLEGAL_SPREADSHEET_CHARACTERS = re.compile(r"[\x00-\x08\x0b-\x0c\x0e-\x1f]")
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\n")
_SECRET_ASSIGNMENT = re.compile(
    r"(?i)(?:"
    r"[\"']?(?:password|passwd|pwd|user[\s_]*id|userid|uid|username|"
    r"account\s*key|client[\s_]*secret|authentication|integrated[\s_]*security|"
    r"persist[\s_]*security[\s_]*info|trusted[\s_]*connection|api\s*key|"
    r"access\s*token|token|secret)[\"']?\s*(?:=|:)\s*"
    r"(?P<general_value>[^;,\r\n\"']*)|"
    r"[\"']?(?:user[\s_]*name|user|account|credential|client[\s_]*id)"
    r"[\"']?\s*=\s*(?P<equals_value>[^;,\r\n\"']*))"
)
_RAW_CONNECTION = re.compile(
    r"(?i)(?:^|[;\"'])\s*(?:(?:ODBC|OLEDB)\s*[:;]\s*)?"
    r"(?:ODBC|PROVIDER|DRIVER|DSN|FILE\s*DSN|DBQ|SERVER|DATA\s*SOURCE|DATABASE|"
    r"INITIAL\s*CATALOG|TRUSTED_CONNECTION)\s*="
)
_URI_USERINFO = re.compile(r"(?i)\b[a-z][a-z0-9+.-]*://[^/@\s]+(?::[^/@\s]*)?@")

_PASS_THROUGH_LINEAGE_HEADERS = (
    "View ID",
    "Target ID",
    "Application ID",
    "Artifact ID",
    "Object ID",
    "Object Name",
    "Query Kind",
    "Connection ID",
    "Connection Kind",
    "Direct Provenance",
    "Resolution Status",
    "Resolution Provenance",
    "Resolution Warnings",
    "Datasource ID",
    "DSN",
    "Platform",
    "Driver",
    "Server",
    "Database",
    "Interaction ID",
    "Operation",
    "Scope",
    "Catalog",
    "Target Database",
    "Schema",
    "Target Object",
    "Local Target Object ID",
    "View Evidence IDs",
    "Target Evidence IDs",
)
_LINKED_TABLE_LINEAGE_HEADERS = (
    *_PASS_THROUGH_LINEAGE_HEADERS[:6],
    "Object Kind",
    "Source Table",
    *_PASS_THROUGH_LINEAGE_HEADERS[7:],
)


class ReportFormat(StrEnum):
    HTML = "html"
    XLSX = "xlsx"
    PDF = "pdf"
    CSV_BUNDLE = "csv_bundle"


class ReportLeakError(ValueError):
    """Raised when a report model or rendered artifact appears to contain a secret."""


class ReportPublicationError(RuntimeError):
    """Raised when a report run cannot be safely or immutably published."""


class ReportArtifact(StrictModel):
    report_format: ReportFormat
    relative_path: NonEmptyString
    sha256: Sha256
    size_bytes: Annotated[int, Field(ge=0)]
    record_ids: tuple[NonEmptyString, ...]

    @model_validator(mode="after")
    def normalize_and_validate(self) -> Self:
        candidate = PurePosixPath(self.relative_path)
        if candidate.is_absolute() or ".." in candidate.parts:
            raise ValueError("report artifact path must be relative and contained")
        object.__setattr__(self, "record_ids", tuple(sorted(set(self.record_ids))))
        return self


class ReportPublicationManifest(StrictModel):
    schema_version: Literal["report-publication-v2"] = "report-publication-v2"
    run_id: NonEmptyString
    report_id: NonEmptyString
    report_sha256: Sha256
    report_status: ReportStatus
    generated_at: datetime
    application_ids: tuple[NonEmptyString, ...]
    candidate_ids: tuple[NonEmptyString, ...]
    finding_ids: tuple[NonEmptyString, ...]
    record_ids: tuple[NonEmptyString, ...]
    record_count: Annotated[int, Field(ge=0)]
    application_count: Annotated[int, Field(ge=0)]
    candidate_count: Annotated[int, Field(ge=0)]
    finding_count: Annotated[int, Field(ge=0)]
    artifacts: tuple[ReportArtifact, ...]

    @field_validator("run_id")
    @classmethod
    def validate_run_id(cls, value: str) -> str:
        if not _SAFE_RUN_ID.fullmatch(value):
            raise ValueError("report run_id is unsafe")
        return value

    @model_validator(mode="after")
    def validate_identity(self) -> Self:
        for field_name in (
            "application_ids",
            "candidate_ids",
            "finding_ids",
            "record_ids",
        ):
            object.__setattr__(self, field_name, tuple(sorted(set(getattr(self, field_name)))))
        if self.record_count != len(self.record_ids):
            raise ValueError("record count does not match report IDs")
        if self.application_count != len(self.application_ids):
            raise ValueError("application count does not match report IDs")
        if self.candidate_count != len(self.candidate_ids):
            raise ValueError("candidate count does not match report IDs")
        if self.finding_count != len(self.finding_ids):
            raise ValueError("finding count does not match report IDs")
        formats = [item.report_format for item in self.artifacts]
        if set(formats) != set(ReportFormat) or len(formats) != len(ReportFormat):
            raise ValueError("publication must contain exactly one artifact for each format")
        object.__setattr__(
            self,
            "artifacts",
            tuple(sorted(self.artifacts, key=lambda item: item.report_format.value)),
        )
        if any(item.record_ids != self.record_ids for item in self.artifacts):
            raise ValueError("every report format must preserve the same record IDs")
        has_partial_names = all(
            ".partial." in item.relative_path for item in self.artifacts
        )
        if (self.report_status == ReportStatus.PARTIAL) != has_partial_names:
            raise ValueError("report artifact names must distinguish partial publications")
        return self


class LatestReportPointer(StrictModel):
    schema_version: Literal["latest-report-pointer-v2"] = "latest-report-pointer-v2"
    run_id: NonEmptyString
    report_id: NonEmptyString
    report_sha256: Sha256
    manifest_relative_path: NonEmptyString
    manifest_sha256: Sha256

    @model_validator(mode="after")
    def validate_pointer(self) -> Self:
        if not _SAFE_RUN_ID.fullmatch(self.run_id):
            raise ValueError("latest report run_id is unsafe")
        expected = PurePosixPath("runs", self.run_id, "manifest.json").as_posix()
        if self.manifest_relative_path != expected:
            raise ValueError("latest report pointer has an invalid manifest path")
        return self


def render_report_html(report: PortfolioReportModel) -> bytes:
    """Render a script-free, offline HTML report with a deny-by-default CSP."""

    report = _require_report_model(report)
    scan_report_for_leaks(report)
    application_rows = "".join(
        f'<tr data-entity-type="application" data-entity-id="{escape(item.application_id)}">'
        f"<td>{escape(item.application_id)}</td>"
        f"<td>{escape(item.application_name)}</td>"
        f"<td>{escape(item.status.value)}</td>"
        f"<td>{item.object_count}</td><td>{item.query_count}</td>"
        f"<td>{item.datasource_count}</td>"
        f"<td>{escape(item.summary or '')}</td>"
        f"<td>{escape(' | '.join(item.coverage_notes))}</td>"
        "</tr>"
        for item in report.applications
    )
    candidate_rows = "".join(
        f'<tr data-entity-type="candidate" data-entity-id="{escape(item.candidate_id)}">'
        f"<td>{escape(item.candidate_id)}</td>"
        f"<td>{escape(item.candidate_type.value)}</td>"
        f"<td>{escape(', '.join(item.application_ids))}</td>"
        f"<td>{item.score:.6f}</td>"
        f"<td>{escape(', '.join(item.evidence_ids))}</td>"
        "</tr>"
        for item in report.candidates
    )
    finding_rows = "".join(
        f'<tr data-entity-type="finding" data-entity-id="{escape(item.finding_id)}">'
        f"<td>{escape(item.finding_id)}</td>"
        f"<td>{escape(item.kind.value)}</td>"
        f"<td>{escape(item.title)}</td>"
        f"<td>{escape(', '.join(item.application_ids))}</td>"
        f"<td>{escape(item.narrative)}</td>"
        "</tr>"
        for item in report.portfolio_findings
    )
    fact_rows = "".join(
        "<tr>"
        f"<td>{escape(item.evidence_id)}</td>"
        f"<td>{escape(item.application_id)}</td>"
        f"<td>{escape(item.fact_type)}</td>"
        f"<td>{escape(item.confidence.value)}</td>"
        f"<td>{escape(item.observation)}</td>"
        "</tr>"
        for item in report.observed_facts
    )
    claim_rows = "".join(
        "<tr>"
        f"<td>{escape(item.claim_id)}</td>"
        f"<td>{escape(item.application_id)}</td>"
        f"<td>{escape(item.field)}</td>"
        f"<td>{escape(item.value)}</td>"
        f"<td>{escape(item.source)}</td>"
        "</tr>"
        for item in report.owner_claims
    )
    profile_rows = "".join(
        "<tr>"
        f"<td>{escape(item.interpretation_id)}</td>"
        f"<td>{escape(item.application_id)}</td>"
        f"<td>{escape(item.summary)}</td>"
        f"<td>{escape(', '.join(item.evidence_ids))}</td>"
        "</tr>"
        for item in report.application_profiles
    )
    review_rows = "".join(
        "<tr>"
        f"<td>{escape(item.review_id)}</td>"
        f"<td>{escape(item.application_id)}</td>"
        f"<td>{escape(item.subject_id)}</td>"
        f"<td>{escape(item.status.value)}</td>"
        f"<td>{escape(item.notes or '')}</td>"
        "</tr>"
        for item in report.review_records
    )
    unresolved_rows = "".join(
        "<tr>"
        f"<td>{escape(item.entry_id)}</td>"
        f"<td>{escape(item.unresolved.unresolved_id)}</td>"
        f"<td>{escape(item.application_id)}</td>"
        f"<td>{escape(item.unresolved.reference)}</td>"
        f"<td>{escape(item.unresolved.reason)}</td>"
        "</tr>"
        for item in report.unresolved_references
    )
    coverage_rows = "".join(
        "<tr>"
        f"<td>{escape(item.coverage_id)}</td>"
        f"<td>{escape(item.application_id)}</td>"
        f"<td>{item.coverage.primary_artifact_count}</td>"
        f"<td>{item.coverage.extracted_object_count}</td>"
        f"<td>{item.coverage.warning_count}</td>"
        "</tr>"
        for item in report.coverage
    )
    pass_through_rows = _html_table_rows(_pass_through_lineage_rows(report))
    pass_through_headers = _html_table_headers(_PASS_THROUGH_LINEAGE_HEADERS)
    linked_rows = _html_table_rows(_linked_table_lineage_rows(report))
    linked_headers = _html_table_headers(_LINKED_TABLE_LINEAGE_HEADERS)
    omission_rows = "".join(
        "<tr>"
        f"<td>{escape(item.omission_id)}</td>"
        f"<td>{escape(item.application_id or '')}</td>"
        f"<td>{escape(item.logical_unit_id or '')}</td>"
        f"<td>{escape(item.stage.value)}</td>"
        f"<td>{escape(item.reason)}</td>"
        "</tr>"
        for item in report.omissions
    )
    inventory_exclusion_rows = "".join(
        "<tr>"
        f"<td>{escape(item.exclusion_id)}</td>"
        f"<td>{escape(item.application_id)}</td>"
        f"<td>{escape(item.application_name)}</td>"
        f"<td>{escape(item.filename)}</td>"
        f"<td>{escape(item.status)}</td>"
        f"<td>{escape(item.reason)}</td>"
        "</tr>"
        for item in report.inventory_exclusions
    )
    warning_rows = "".join(f"<li>{escape(item)}</li>" for item in report.warnings)
    all_record_ids = _report_record_ids(report)
    inventory_exclusion_count = len(report.inventory_exclusions)
    watermark_html = (
        f'<p class="partial">{escape(report.partial_watermark)}</p>'
        if report.partial_watermark
        else ""
    )
    record_id_rows = "".join(
        f"<li><code>{escape(identifier)}</code></li>" for identifier in all_record_ids
    )
    html = f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta http-equiv="Content-Security-Policy"
 content="default-src 'none'; style-src 'unsafe-inline'; img-src data:;
 base-uri 'none'; form-action 'none'; frame-ancestors 'none'">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Access Portfolio Report {escape(report.report_id)}</title>
<style>
body{{font:14px system-ui,sans-serif;color:#17242d;margin:2rem;line-height:1.45}}
h1,h2{{color:#123047}}.metrics{{display:flex;gap:2rem;flex-wrap:wrap}}
table{{border-collapse:collapse;width:100%;margin:1rem 0 2rem}}
th,td{{border:1px solid #cad4da;padding:.45rem;text-align:left;vertical-align:top}}
th{{background:#123047;color:white}}
.partial{{color:#8a4b08;font-weight:700}}code{{overflow-wrap:anywhere}}
</style>
</head>
<body data-report-id="{escape(report.report_id)}">
<h1>Access Portfolio Report</h1>
{watermark_html}
<p>Report ID: <code>{escape(report.report_id)}</code></p>
<p>Analysis fingerprint: <code>{escape(report.analysis_fingerprint)}</code></p>
<p class="{'partial' if report.status.value == 'partial' else ''}">
Status: {escape(report.status.value)}</p>
<dl class="metrics">
<div><dt>Applications</dt><dd data-count="applications">{len(report.applications)}</dd></div>
<div><dt>Candidates</dt><dd data-count="candidates">{len(report.candidates)}</dd></div>
<div><dt>Findings</dt><dd data-count="findings">{len(report.portfolio_findings)}</dd></div>
<div><dt>Observed facts</dt><dd data-count="facts">{len(report.observed_facts)}</dd></div>
<div><dt>Claims</dt><dd data-count="claims">{len(report.owner_claims)}</dd></div>
<div><dt>Reviews</dt><dd data-count="reviews">{len(report.review_records)}</dd></div>
<div><dt>Omissions</dt><dd data-count="omissions">{len(report.omissions)}</dd></div>
<div><dt>Inventory exclusions</dt>
<dd data-count="inventory-exclusions">{inventory_exclusion_count}</dd></div>
<div><dt>Data access</dt><dd data-count="data-access">{len(report.data_access)}</dd></div>
<div><dt>Dependencies</dt><dd data-count="dependencies">{len(report.dependency_edges)}</dd></div>
<div><dt>Normalized IDs</dt><dd data-count="records">{len(all_record_ids)}</dd></div>
</dl>
<h2>Coverage qualifications and warnings</h2><ul>{warning_rows}</ul>
<h2>Applications</h2><table><thead><tr><th>ID</th><th>Name</th><th>Status</th>
<th>Objects</th><th>Queries</th><th>Datasources</th><th>Summary</th>
<th>Coverage notes</th></tr></thead>
<tbody>{application_rows}</tbody></table>
<h2>Overlap candidates</h2><table><thead><tr><th>ID</th><th>Type</th>
<th>Applications</th><th>Score</th><th>Evidence IDs</th></tr></thead>
<tbody>{candidate_rows}</tbody></table>
<h2>Portfolio findings</h2><table><thead><tr><th>ID</th><th>Kind</th><th>Title</th>
<th>Applications</th><th>Narrative</th></tr></thead><tbody>{finding_rows}</tbody></table>
<h2>Observed facts</h2><table><thead><tr><th>Evidence ID</th><th>Application</th>
<th>Fact type</th><th>Confidence</th><th>Observation</th></tr></thead>
<tbody>{fact_rows}</tbody></table>
<h2>Owner claims</h2><table><thead><tr><th>Claim ID</th><th>Application</th>
<th>Field</th><th>Value</th><th>Source</th></tr></thead><tbody>{claim_rows}</tbody></table>
<h2>Application interpretations</h2><table><thead><tr><th>Interpretation ID</th>
<th>Application</th><th>Summary</th><th>Evidence IDs</th></tr></thead>
<tbody>{profile_rows}</tbody></table>
<h2>Review records</h2><table><thead><tr><th>Review ID</th><th>Application</th>
<th>Subject</th><th>Status</th><th>Notes</th></tr></thead><tbody>{review_rows}</tbody></table>
<h2>Unresolved references</h2><table><thead><tr><th>Entry ID</th>
<th>Unresolved ID</th><th>Application</th><th>Reference</th><th>Reason</th></tr></thead>
<tbody>{unresolved_rows}</tbody></table>
<h2>Extraction coverage</h2><table><thead><tr><th>Coverage ID</th>
<th>Application</th><th>Artifacts</th><th>Extracted objects</th><th>Warnings</th></tr>
</thead><tbody>{coverage_rows}</tbody></table>
<h2>Pass-through queries</h2><table><thead><tr>{pass_through_headers}</tr></thead>
<tbody>{pass_through_rows}</tbody></table>
<h2>Linked tables</h2><table><thead><tr>{linked_headers}</tr></thead>
<tbody>{linked_rows}</tbody></table>
<h2>Structured omissions</h2><table><thead><tr><th>Omission ID</th>
<th>Application</th><th>Logical unit</th><th>Stage</th><th>Reason</th></tr></thead>
<tbody>{omission_rows}</tbody></table>
<h2>Unsupported inventory exclusions</h2><table><thead><tr><th>Exclusion ID</th>
<th>Application</th><th>Name</th><th>Filename</th><th>Status</th><th>Reason</th></tr>
</thead><tbody>{inventory_exclusion_rows}</tbody></table>
<h2>Normalized record identities</h2><p>Every renderer carries this same identity set.</p>
<ul>{record_id_rows}</ul>
</body></html>
"""
    payload = html.encode("utf-8")
    _scan_text_for_leaks(html, "HTML report")
    return payload


def write_report_xlsx(path: Path, report: PortfolioReportModel) -> None:
    """Write a formula-inert workbook from the shared report model."""

    report = _require_report_model(report)
    scan_report_for_leaks(report)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    _append_rows(
        summary,
        (
            ("Measure", "Value"),
            ("Report ID", report.report_id),
            ("Status", report.status.value),
            ("Analysis Fingerprint", report.analysis_fingerprint),
            ("Partial Watermark", report.partial_watermark or ""),
            ("Absence Claims Suppressed", report.absence_claims_suppressed),
            ("Applications", len(report.applications)),
            ("Candidates", len(report.candidates)),
            ("Findings", len(report.portfolio_findings)),
            ("Observed Facts", len(report.observed_facts)),
            ("Owner Claims", len(report.owner_claims)),
            ("Application Profiles", len(report.application_profiles)),
            ("Review Records", len(report.review_records)),
            ("Unresolved References", len(report.unresolved_references)),
            ("Inventory Exclusions", len(report.inventory_exclusions)),
            ("Pass-through Queries", len(report.pass_through_queries)),
            ("Linked Tables", len(report.linked_tables)),
            ("Normalized IDs", len(_report_record_ids(report))),
        ),
    )
    applications = workbook.create_sheet("Applications")
    _append_rows(
        applications,
        (
            (
                "Application ID",
                "Name",
                "Status",
                "Evidence Bundle SHA-256",
                "Interpretation SHA-256",
                "Artifacts",
                "Objects",
                "Tables",
                "Queries",
                "Datasources",
                "Connections",
                "Interactions",
                "Evidence",
                "Unresolved",
                "Summary",
                "Business Purpose",
                "Capabilities",
                "Major Workflows",
                "Coverage Notes",
            ),
            *(
                (
                    item.application_id,
                    item.application_name,
                    item.status.value,
                    item.evidence_bundle_sha256,
                    item.interpretation_sha256 or "",
                    item.artifact_count,
                    item.object_count,
                    item.table_count,
                    item.query_count,
                    item.datasource_count,
                    item.connection_count,
                    item.interaction_count,
                    item.evidence_count,
                    item.unresolved_reference_count,
                    item.summary or "",
                    item.business_purpose or "",
                    " | ".join(item.capabilities),
                    " | ".join(item.major_workflows),
                    " | ".join(item.coverage_notes),
                )
                for item in report.applications
            ),
        ),
    )
    candidates = workbook.create_sheet("Candidates")
    _append_rows(
        candidates,
        (
            ("Candidate ID", "Type", "Score", "Policy", "Applications", "Evidence", "Basis"),
            *(
                (
                    item.candidate_id,
                    item.candidate_type.value,
                    item.score,
                    item.policy_version,
                    " | ".join(item.application_ids),
                    " | ".join(item.evidence_ids),
                    " | ".join(item.basis_ids),
                )
                for item in report.candidates
            ),
        ),
    )
    findings = workbook.create_sheet("Findings")
    _append_rows(
        findings,
        (
            ("Finding ID", "Kind", "Title", "Confidence", "Applications", "Evidence", "Narrative"),
            *(
                (
                    item.finding_id,
                    item.kind.value,
                    item.title,
                    item.confidence.value,
                    " | ".join(item.application_ids),
                    " | ".join(item.evidence_ids),
                    item.narrative,
                )
                for item in report.portfolio_findings
            ),
        ),
    )
    facts = workbook.create_sheet("Observed Facts")
    _append_rows(
        facts,
        (
            (
                "Evidence ID",
                "Application ID",
                "Artifact ID",
                "Object ID",
                "Fact Type",
                "Location",
                "Observation",
                "Confidence",
            ),
            *(
                (
                    item.evidence_id,
                    item.application_id,
                    item.artifact_id,
                    item.object_id or "",
                    item.fact_type,
                    item.location or "",
                    item.observation,
                    item.confidence.value,
                )
                for item in report.observed_facts
            ),
        ),
    )
    claims = workbook.create_sheet("Owner Claims")
    _append_rows(
        claims,
        (
            ("Claim ID", "Application ID", "Field", "Value", "Source"),
            *(
                (item.claim_id, item.application_id, item.field, item.value, item.source)
                for item in report.owner_claims
            ),
        ),
    )
    profiles = workbook.create_sheet("Profiles")
    _append_rows(
        profiles,
        (
            (
                "Interpretation ID",
                "Application ID",
                "Summary",
                "Business Purpose",
                "Evidence IDs",
                "Claim IDs",
            ),
            *(
                (
                    item.interpretation_id,
                    item.application_id,
                    item.summary,
                    item.business_purpose,
                    " | ".join(item.evidence_ids),
                    " | ".join(item.claim_ids),
                )
                for item in report.application_profiles
            ),
        ),
    )
    reviews = workbook.create_sheet("Reviews")
    _append_rows(
        reviews,
        (
            (
                "Review ID",
                "Application ID",
                "Subject ID",
                "Status",
                "Notes",
                "Evidence IDs",
                "Claim IDs",
            ),
            *(
                (
                    item.review_id,
                    item.application_id,
                    item.subject_id,
                    item.status.value,
                    item.notes or "",
                    " | ".join(item.evidence_ids),
                    " | ".join(item.claim_ids),
                )
                for item in report.review_records
            ),
        ),
    )
    unresolved = workbook.create_sheet("Unresolved")
    _append_rows(
        unresolved,
        (
            (
                "Entry ID",
                "Unresolved ID",
                "Application ID",
                "Source Object ID",
                "Reference",
                "Reason",
                "Evidence IDs",
            ),
            *(
                (
                    item.entry_id,
                    item.unresolved.unresolved_id,
                    item.application_id,
                    item.unresolved.source_object_id,
                    item.unresolved.reference,
                    item.unresolved.reason,
                    " | ".join(item.unresolved.evidence_ids),
                )
                for item in report.unresolved_references
            ),
        ),
    )
    coverage = workbook.create_sheet("Coverage")
    _append_rows(
        coverage,
        (
            (
                "Coverage ID",
                "Application ID",
                "Artifacts",
                "Complete",
                "Partial",
                "Failed",
                "Discovered Objects",
                "Extracted Objects",
                "Warnings",
                "Unresolved",
            ),
            *(
                (
                    item.coverage_id,
                    item.application_id,
                    item.coverage.primary_artifact_count,
                    item.coverage.complete_artifact_count,
                    item.coverage.partial_artifact_count,
                    item.coverage.failed_artifact_count,
                    item.coverage.discovered_object_count,
                    item.coverage.extracted_object_count,
                    item.coverage.warning_count,
                    item.coverage.unresolved_reference_count,
                )
                for item in report.coverage
            ),
        ),
    )
    pass_through = workbook.create_sheet("Pass Through")
    _append_rows(
        pass_through,
        (
            _PASS_THROUGH_LINEAGE_HEADERS,
            *_pass_through_lineage_rows(report),
        ),
    )
    linked = workbook.create_sheet("Linked Tables")
    _append_rows(
        linked,
        (
            _LINKED_TABLE_LINEAGE_HEADERS,
            *_linked_table_lineage_rows(report),
        ),
    )
    registry_objects = workbook.create_sheet("Object Registry")
    _append_rows(
        registry_objects,
        (
            ("Entry ID", "Application ID", "Object ID", "Type", "Name", "Evidence IDs"),
            *(
                (
                    item.entry_id,
                    item.application_id,
                    item.access_object.object_id,
                    item.access_object.object_type.value,
                    item.access_object.name,
                    " | ".join(item.access_object.evidence_ids),
                )
                for item in report.object_registry
            ),
        ),
    )
    registry_tables = workbook.create_sheet("Table Registry")
    _append_rows(
        registry_tables,
        (
            (
                "Entry ID",
                "Application ID",
                "Object ID",
                "Linked",
                "Hidden",
                "System",
                "Source Table",
                "Connection ID",
                "Evidence IDs",
            ),
            *(
                (
                    item.entry_id,
                    item.application_id,
                    item.table.object_id,
                    item.table.is_linked,
                    item.table.is_hidden,
                    item.table.is_system,
                    item.table.source_table_name or "",
                    item.table.connection_id or "",
                    " | ".join(item.table.evidence_ids),
                )
                for item in report.table_registry
            ),
        ),
    )
    registry_queries = workbook.create_sheet("Query Registry")
    _append_rows(
        registry_queries,
        (
            (
                "Entry ID",
                "Application ID",
                "Object ID",
                "Kind",
                "Hidden",
                "System",
                "Connection Kind",
                "Connection ID",
                "Returns Records",
                "SQL",
                "Evidence IDs",
            ),
            *(
                (
                    item.entry_id,
                    item.application_id,
                    item.query.object_id,
                    item.query.query_kind.value,
                    item.query.is_hidden,
                    item.query.is_system,
                    item.query.connection_kind.value,
                    item.query.connection_id or "",
                    item.query.returns_records
                    if item.query.returns_records is not None
                    else "",
                    item.query.sanitized_sql or "",
                    " | ".join(item.query.evidence_ids),
                )
                for item in report.query_registry
            ),
        ),
    )
    registry_connections = workbook.create_sheet("Connection Registry")
    _append_rows(
        registry_connections,
        (
            (
                "Connection ID",
                "Application ID",
                "Object ID",
                "Kind",
                "Resolution",
                "Datasource ID",
                "Evidence IDs",
            ),
            *(
                (
                    item.connection_id,
                    item.application_id,
                    item.object_id,
                    item.connection_kind.value,
                    item.resolution_status.value,
                    item.datasource_id or "",
                    " | ".join(item.evidence_ids),
                )
                for item in report.connection_registry
            ),
        ),
    )
    registry_datasources = workbook.create_sheet("Datasource Registry")
    _append_rows(
        registry_datasources,
        (
            ("Datasource ID", "Platform", "Driver", "DSN", "Server", "Database", "Resource"),
            *(
                (
                    item.datasource_id,
                    item.platform,
                    item.driver or "",
                    item.dsn or "",
                    item.server or "",
                    item.database or "",
                    item.resource or "",
                )
                for item in report.datasource_registry
            ),
        ),
    )
    data_access = workbook.create_sheet("Data Access")
    _append_rows(
        data_access,
        (
            (
                "Entry ID",
                "Interaction ID",
                "Application ID",
                "Source Object ID",
                "Operation",
                "Scope",
                "Connection ID",
                "Datasource ID",
                "Target Object",
                "Evidence IDs",
            ),
            *(
                (
                    item.entry_id,
                    item.data_access.interaction_id,
                    item.application_id,
                    item.data_access.source_object_id,
                    item.data_access.operation.value,
                    item.data_access.scope.value,
                    item.data_access.connection_id or "",
                    item.data_access.datasource_id or "",
                    item.data_access.object_name or item.data_access.local_target_object_id or "",
                    " | ".join(item.data_access.evidence_ids),
                )
                for item in report.data_access
            ),
        ),
    )
    dependencies = workbook.create_sheet("Dependencies")
    _append_rows(
        dependencies,
        (
            (
                "Entry ID",
                "Edge ID",
                "Application ID",
                "Source Node ID",
                "Target Node ID",
                "Relationship",
                "Operation",
                "Evidence IDs",
            ),
            *(
                (
                    item.entry_id,
                    item.dependency_edge.edge_id,
                    item.application_id,
                    item.dependency_edge.source_node_id,
                    item.dependency_edge.target_node_id,
                    item.dependency_edge.relationship.value,
                    item.dependency_edge.operation.value,
                    " | ".join(item.dependency_edge.evidence_ids),
                )
                for item in report.dependency_edges
            ),
        ),
    )
    dependency_nodes = workbook.create_sheet("Dependency Nodes")
    _append_rows(
        dependency_nodes,
        (
            (
                "Node ID",
                "Application ID",
                "Kind",
                "Label",
                "Artifact ID",
                "Object ID",
                "Datasource ID",
            ),
            *(
                (
                    item.node_id,
                    item.application_id,
                    item.kind.value,
                    item.label,
                    item.artifact_id or "",
                    item.object_id or "",
                    item.datasource_id or "",
                )
                for item in report.dependency_nodes
            ),
        ),
    )
    omissions = workbook.create_sheet("Omissions")
    _append_rows(
        omissions,
        (
            ("Omission ID", "Application ID", "Logical Unit ID", "Stage", "Reason"),
            *(
                (
                    item.omission_id,
                    item.application_id or "",
                    item.logical_unit_id or "",
                    item.stage.value,
                    item.reason,
                )
                for item in report.omissions
            ),
        ),
    )
    inventory_exclusions = workbook.create_sheet("Inventory Exclusions")
    _append_rows(
        inventory_exclusions,
        (
            (
                "Exclusion ID",
                "Application ID",
                "Application Name",
                "Filename",
                "Source Locator",
                "Status",
                "Reason",
            ),
            *(
                (
                    item.exclusion_id,
                    item.application_id,
                    item.application_name,
                    item.filename,
                    item.source_locator,
                    item.status,
                    item.reason,
                )
                for item in report.inventory_exclusions
            ),
        ),
    )
    warnings = workbook.create_sheet("Warnings")
    _append_rows(
        warnings,
        (
            ("Warning",),
            *((item,) for item in report.warnings),
        ),
    )
    identities = workbook.create_sheet("Record IDs")
    _append_rows(
        identities,
        (
            ("Record ID",),
            *((identifier,) for identifier in _report_record_ids(report)),
        ),
    )
    review_queue = workbook.create_sheet("Review Queue")
    review_by_subject = {item.subject_id: item for item in report.review_records}
    proposals = (
        *((item.candidate_id, item.evidence_ids) for item in report.candidates),
        *((item.finding_id, item.evidence_ids) for item in report.portfolio_findings),
    )
    decision_values = {
        "accepted": "accept",
        "edited": "edit",
        "rejected": "reject",
        "pending": "",
    }
    _append_rows(
        review_queue,
        (
            tuple(REVIEW_HEADERS),
            *(
                (
                    report.analysis_fingerprint,
                    proposal_id,
                    ",".join(evidence_ids),
                    decision_values[
                        review_by_subject[proposal_id].status.value
                    ]
                    if proposal_id in review_by_subject
                    else "",
                    review_by_subject[proposal_id].edited_value or ""
                    if proposal_id in review_by_subject
                    else "",
                    review_by_subject[proposal_id].reviewer or ""
                    if proposal_id in review_by_subject
                    else "",
                    review_by_subject[proposal_id].notes or ""
                    if proposal_id in review_by_subject
                    else "",
                )
                for proposal_id, evidence_ids in proposals
            ),
        ),
    )
    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="123047")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def write_report_pdf(path: Path, report: PortfolioReportModel) -> None:
    """Write a compact offline PDF from the shared report model."""

    report = _require_report_model(report)
    scan_report_for_leaks(report)
    styles = getSampleStyleSheet()
    story: list[object] = [
        Paragraph("Access Portfolio Report", styles["Title"]),
        Paragraph(f"Report ID: {escape(report.report_id)}", styles["BodyText"]),
        Paragraph(
            f"Analysis fingerprint: {escape(report.analysis_fingerprint)}",
            styles["BodyText"],
        ),
        Spacer(1, 0.12 * inch),
        _pdf_table(
            (
                ("Measure", "Count"),
                ("Applications", str(len(report.applications))),
                ("Candidates", str(len(report.candidates))),
                ("Findings", str(len(report.portfolio_findings))),
                ("Observed facts", str(len(report.observed_facts))),
                ("Claims", str(len(report.owner_claims))),
                ("Reviews", str(len(report.review_records))),
                ("Omissions", str(len(report.omissions))),
                ("Inventory exclusions", str(len(report.inventory_exclusions))),
                ("Data access", str(len(report.data_access))),
                ("Dependencies", str(len(report.dependency_edges))),
            )
        ),
        Spacer(1, 0.2 * inch),
        Paragraph("Applications", styles["Heading2"]),
    ]
    if report.partial_watermark is not None:
        story.insert(1, Paragraph(escape(report.partial_watermark), styles["Heading2"]))
    _append_pdf_section(
        story,
        styles,
        "Cross-format verification registry",
        (
            f"APA-N:{len(report.applications)}:{len(report.candidates)}:"
            f"{len(report.portfolio_findings)}",
            *(
                _pdf_entity_marker("application", item.application_id)
                for item in report.applications
            ),
            *(
                _pdf_entity_marker("candidate", item.candidate_id)
                for item in report.candidates
            ),
            *(
                _pdf_entity_marker("finding", item.finding_id)
                for item in report.portfolio_findings
            ),
        ),
    )
    _append_pdf_section(story, styles, "Coverage qualifications and warnings", report.warnings)
    for application in report.applications:
        story.append(
            Paragraph(
                escape(
                    f"{application.application_id} — {application.application_name} — "
                    f"{application.status.value} — objects {application.object_count}, "
                    f"queries {application.query_count}, "
                    f"datasources {application.datasource_count}"
                ),
                styles["BodyText"],
            )
        )
    story.append(Paragraph("Overlap candidates", styles["Heading2"]))
    for candidate in report.candidates:
        story.append(
            Paragraph(
                escape(
                    f"{candidate.candidate_id} — {candidate.candidate_type.value} — "
                    f"{', '.join(candidate.application_ids)} — "
                    f"score {candidate.score:.6f}"
                ),
                styles["BodyText"],
            )
        )
    story.append(Paragraph("Portfolio findings", styles["Heading2"]))
    for finding in report.portfolio_findings:
        story.append(
            Paragraph(
                escape(f"{finding.finding_id} — {finding.title}: {finding.narrative}"),
                styles["BodyText"],
            )
        )
    _append_pdf_section(
        story,
        styles,
        "Observed facts",
        tuple(
            f"{item.evidence_id} — {item.application_id} — {item.fact_type}: "
            f"{item.observation}"
            for item in report.observed_facts
        ),
    )
    _append_pdf_section(
        story,
        styles,
        "Normalized record identities",
        _report_record_ids(report),
    )
    _append_pdf_section(
        story,
        styles,
        "Owner claims",
        tuple(
            f"{item.claim_id} — {item.application_id} — {item.field}: {item.value}"
            for item in report.owner_claims
        ),
    )
    _append_pdf_section(
        story,
        styles,
        "Application interpretations",
        tuple(
            f"{item.interpretation_id} — {item.application_id}: {item.summary} — "
            f"evidence {', '.join(item.evidence_ids)}"
            for item in report.application_profiles
        ),
    )
    _append_pdf_section(
        story,
        styles,
        "Review records",
        tuple(
            f"{item.review_id} — {item.application_id} — {item.subject_id} — "
            f"{item.status.value}"
            for item in report.review_records
        ),
    )
    _append_pdf_section(
        story,
        styles,
        "Unresolved references",
        tuple(
            f"{item.entry_id} — {item.unresolved.unresolved_id} — "
            f"{item.application_id}: {item.unresolved.reference} — {item.unresolved.reason}"
            for item in report.unresolved_references
        ),
    )
    _append_pdf_section(
        story,
        styles,
        "Extraction coverage",
        tuple(
            f"{item.coverage_id} — {item.application_id} — "
            f"artifacts {item.coverage.primary_artifact_count}, "
            f"objects {item.coverage.extracted_object_count}"
            for item in report.coverage
        ),
    )
    _append_pdf_section(
        story,
        styles,
        "Pass-through queries",
        tuple(
            _lineage_row_text(_PASS_THROUGH_LINEAGE_HEADERS, row)
            for row in _pass_through_lineage_rows(report)
        ),
    )
    _append_pdf_section(
        story,
        styles,
        "Linked tables",
        tuple(
            _lineage_row_text(_LINKED_TABLE_LINEAGE_HEADERS, row)
            for row in _linked_table_lineage_rows(report)
        ),
    )
    _append_pdf_section(
        story,
        styles,
        "Data access",
        tuple(
            f"{item.entry_id} — {item.application_id} — "
            f"{item.data_access.operation.value} — {item.data_access.scope.value} — "
            + (
                item.data_access.object_name
                or item.data_access.local_target_object_id
                or "unresolved"
            )
            for item in report.data_access
        ),
    )
    _append_pdf_section(
        story,
        styles,
        "Dependencies",
        tuple(
            f"{item.entry_id} — {item.application_id} — "
            f"{item.dependency_edge.source_node_id} "
            f"{item.dependency_edge.relationship.value} "
            f"{item.dependency_edge.target_node_id}"
            for item in report.dependency_edges
        ),
    )
    _append_pdf_section(
        story,
        styles,
        "Structured omissions",
        tuple(
            f"{item.omission_id} — {item.application_id or 'portfolio'} — "
            f"{item.logical_unit_id or 'all units'} — {item.stage.value}: {item.reason}"
            for item in report.omissions
        ),
    )
    _append_pdf_section(
        story,
        styles,
        "Unsupported inventory exclusions",
        tuple(
            f"{item.exclusion_id} — {item.application_id} — {item.filename} — "
            f"{item.reason}"
            for item in report.inventory_exclusions
        ),
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    document = SimpleDocTemplate(
        str(path),
        pagesize=letter,
        leftMargin=0.65 * inch,
        rightMargin=0.65 * inch,
        topMargin=0.65 * inch,
        bottomMargin=0.65 * inch,
        title=f"Access Portfolio Report {report.report_id}",
        author="Access Portfolio Analyzer",
        keywords=list(_report_record_ids(report)),
        pageCompression=0,
    )
    document.build(story)


def render_report_csv_bundle(report: PortfolioReportModel) -> bytes:
    """Render deterministic normalized CSV tables as a ZIP archive."""

    report = _require_report_model(report)
    scan_report_for_leaks(report)
    tables = _csv_tables(report)
    output = io.BytesIO()
    with zipfile.ZipFile(output, mode="w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, rows in sorted(tables.items()):
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            archive.writestr(info, _csv_bytes(rows))
    payload = output.getvalue()
    _scan_zip_for_leaks(payload, "CSV bundle")
    return payload


def publish_report_run(
    reports_root: Path,
    run_id: str,
    report: PortfolioReportModel,
    *,
    publish_latest: bool = True,
) -> ReportPublicationManifest:
    """Atomically publish every format together; the manifest is the visibility boundary."""

    report = _require_report_model(report)
    if not _SAFE_RUN_ID.fullmatch(run_id):
        raise ReportPublicationError("report run_id is unsafe")
    scan_report_for_leaks(report)
    reports_root = reports_root.resolve()
    report_sha256 = report_fingerprint(report)
    runs_root = reports_root / "runs"
    destination = runs_root / run_id
    manifest_path = destination / "manifest.json"
    if manifest_path.is_file():
        manifest = load_report_publication(reports_root, run_id)
        if (
            manifest.report_id != report.report_id
            or manifest.report_sha256 != report_sha256
            or manifest.report_status != report.status
        ):
            raise ReportPublicationError(
                "report run ID is already published with different content"
            )
        if publish_latest and manifest.report_status == ReportStatus.COMPLETE:
            _publish_latest_pointer(reports_root, manifest)
        return manifest
    if destination.exists():
        raise ReportPublicationError("report run directory exists without a valid manifest")

    runs_root.mkdir(parents=True, exist_ok=True)
    temporary = Path(tempfile.mkdtemp(prefix=f".{run_id}.", dir=runs_root))
    try:
        name_prefix = (
            "portfolio.partial"
            if report.status == ReportStatus.PARTIAL
            else "portfolio"
        )
        html_path = temporary / f"{name_prefix}.html"
        xlsx_path = temporary / f"{name_prefix}.xlsx"
        pdf_path = temporary / f"{name_prefix}.pdf"
        csv_path = temporary / f"{name_prefix}.csv.zip"
        _write_bytes_synced(html_path, render_report_html(report))
        write_report_xlsx(xlsx_path, report)
        write_report_pdf(pdf_path, report)
        _write_bytes_synced(csv_path, render_report_csv_bundle(report))
        rendered = {
            ReportFormat.HTML: html_path,
            ReportFormat.XLSX: xlsx_path,
            ReportFormat.PDF: pdf_path,
            ReportFormat.CSV_BUNDLE: csv_path,
        }
        for report_format, path in rendered.items():
            _fsync_file(path)
            _scan_rendered_artifact(path, report_format)
        record_ids = _report_record_ids(report)
        artifacts = tuple(
            ReportArtifact(
                report_format=report_format,
                relative_path=path.name,
                sha256=_sha256_file(path),
                size_bytes=path.stat().st_size,
                record_ids=record_ids,
            )
            for report_format, path in rendered.items()
        )
        manifest = ReportPublicationManifest(
            run_id=run_id,
            report_id=report.report_id,
            report_sha256=report_sha256,
            report_status=report.status,
            generated_at=report.generated_at,
            application_ids=tuple(item.application_id for item in report.applications),
            candidate_ids=tuple(item.candidate_id for item in report.candidates),
            finding_ids=tuple(item.finding_id for item in report.portfolio_findings),
            record_ids=record_ids,
            record_count=len(record_ids),
            application_count=len(report.applications),
            candidate_count=len(report.candidates),
            finding_count=len(report.portfolio_findings),
            artifacts=artifacts,
        )
        _atomic_write_bytes(
            temporary / "manifest.json", canonical_json_bytes(manifest)
        )
        _fsync_directory(temporary)
        os.replace(temporary, destination)
        _fsync_directory(runs_root)
        published = load_report_publication(reports_root, run_id)
        if publish_latest and published.report_status == ReportStatus.COMPLETE:
            _publish_latest_pointer(reports_root, published)
        return published
    except OSError as exc:
        raise ReportPublicationError(f"could not publish report run: {run_id}") from exc
    finally:
        if temporary.exists():
            shutil.rmtree(temporary)


def load_report_publication(
    reports_root: Path, run_id: str
) -> ReportPublicationManifest:
    """Load a report only through its canonical manifest and verify every artifact."""

    if not _SAFE_RUN_ID.fullmatch(run_id):
        raise ReportPublicationError("report run_id is unsafe")
    root = reports_root.resolve() / "runs" / run_id
    manifest_path = root / "manifest.json"
    try:
        payload = manifest_path.read_bytes()
        manifest = ReportPublicationManifest.model_validate_json(payload)
    except (OSError, ValidationError, ValueError) as exc:
        raise ReportPublicationError("report run is not completely published") from exc
    if canonical_json_bytes(manifest) != payload or manifest.run_id != run_id:
        raise ReportPublicationError("report publication manifest is not canonical")
    for artifact in manifest.artifacts:
        path = _contained_report_path(root, artifact.relative_path)
        try:
            size = path.stat().st_size
        except OSError as exc:
            raise ReportPublicationError("published report artifact is missing") from exc
        if size != artifact.size_bytes or _sha256_file(path) != artifact.sha256:
            raise ReportPublicationError("published report artifact failed hash verification")
        _scan_rendered_artifact(path, artifact.report_format)
        _verify_artifact_record_ids(
            path,
            artifact.report_format,
            manifest.record_ids,
            application_ids=manifest.application_ids,
            candidate_ids=manifest.candidate_ids,
            finding_ids=manifest.finding_ids,
        )
    return manifest


def load_latest_report(reports_root: Path) -> ReportPublicationManifest:
    """Resolve and verify the latest completely published report generation."""

    reports_root = reports_root.resolve()
    pointer_path = reports_root / "latest.json"
    try:
        payload = pointer_path.read_bytes()
        pointer = LatestReportPointer.model_validate_json(payload)
    except (OSError, ValidationError, ValueError) as exc:
        raise ReportPublicationError("no valid latest complete report is published") from exc
    if canonical_json_bytes(pointer) != payload:
        raise ReportPublicationError("latest report pointer is not canonical")
    manifest_path = _contained_report_path(
        reports_root, pointer.manifest_relative_path
    )
    try:
        manifest_sha256 = _sha256_file(manifest_path)
    except OSError as exc:
        raise ReportPublicationError("latest report manifest is missing") from exc
    if manifest_sha256 != pointer.manifest_sha256:
        raise ReportPublicationError("latest report manifest failed hash verification")
    manifest = load_report_publication(reports_root, pointer.run_id)
    if (
        manifest.report_status != ReportStatus.COMPLETE
        or manifest.report_id != pointer.report_id
        or manifest.report_sha256 != pointer.report_sha256
    ):
        raise ReportPublicationError("latest report pointer does not match its manifest")
    return manifest


def _publish_latest_pointer(
    reports_root: Path, manifest: ReportPublicationManifest
) -> LatestReportPointer:
    if manifest.report_status != ReportStatus.COMPLETE:
        raise ReportPublicationError("partial reports cannot become latest complete")
    verified = load_report_publication(reports_root, manifest.run_id)
    if verified != manifest:
        raise ReportPublicationError("latest report manifest changed during publication")
    manifest_relative_path = PurePosixPath(
        "runs", manifest.run_id, "manifest.json"
    ).as_posix()
    manifest_path = _contained_report_path(reports_root, manifest_relative_path)
    pointer = LatestReportPointer(
        run_id=manifest.run_id,
        report_id=manifest.report_id,
        report_sha256=manifest.report_sha256,
        manifest_relative_path=manifest_relative_path,
        manifest_sha256=_sha256_file(manifest_path),
    )
    _atomic_write_bytes(reports_root / "latest.json", canonical_json_bytes(pointer))
    return pointer


def scan_report_for_leaks(report: PortfolioReportModel) -> None:
    """Fail closed if normalized report text still resembles credentials or raw connections."""

    report = _require_report_model(report)
    _scan_text_for_leaks(canonical_json_bytes(report).decode("utf-8"), "report model")


def _require_report_model(report: PortfolioReportModel) -> PortfolioReportModel:
    if type(report) is not PortfolioReportModel:
        raise TypeError("V2 renderers accept only PortfolioReportModel")
    return report


def _pass_through_lineage_rows(
    report: PortfolioReportModel,
) -> tuple[tuple[object, ...], ...]:
    return tuple(
        (
            item.view_id,
            target.target_id if target is not None else "",
            item.application_id,
            item.artifact_id,
            item.query_object_id,
            item.query_name,
            item.query_kind.value,
            item.connection_id,
            item.connection_kind.value,
            item.direct_provenance.value,
            item.resolution_status.value,
            " | ".join(item.resolution_provenance),
            " | ".join(item.resolution_warnings),
            item.datasource_id or "",
            item.dsn or "",
            item.platform or "",
            item.driver or "",
            item.server or "",
            item.database or "",
            target.interaction_id or "" if target is not None else "",
            (
                target.operation.value
                if target is not None
                else " | ".join(value.value for value in item.operations)
            ),
            target.scope.value if target is not None else "",
            (
                target.catalog or ""
                if target is not None
                else " | ".join(item.catalogs)
            ),
            target.database or "" if target is not None else "",
            (
                target.schema_name or ""
                if target is not None
                else " | ".join(item.schemas)
            ),
            (
                target.object_name or ""
                if target is not None
                else " | ".join(item.object_names)
            ),
            target.local_target_object_id or "" if target is not None else "",
            " | ".join(item.evidence_ids),
            " | ".join(target.evidence_ids) if target is not None else "",
        )
        for item in report.pass_through_queries
        for target in _targets_or_placeholder(item.targets)
    )


def _linked_table_lineage_rows(
    report: PortfolioReportModel,
) -> tuple[tuple[object, ...], ...]:
    return tuple(
        (
            item.view_id,
            target.target_id if target is not None else "",
            item.application_id,
            item.artifact_id,
            item.table_object_id,
            item.table_name,
            item.object_kind.value,
            item.source_table_name or "",
            item.connection_id,
            item.connection_kind.value,
            item.direct_provenance.value,
            item.resolution_status.value,
            " | ".join(item.resolution_provenance),
            " | ".join(item.resolution_warnings),
            item.datasource_id or "",
            item.dsn or "",
            item.platform or "",
            item.driver or "",
            item.server or "",
            item.database or "",
            target.interaction_id or "" if target is not None else "",
            (
                target.operation.value
                if target is not None
                else " | ".join(value.value for value in item.operations)
            ),
            target.scope.value if target is not None else "",
            (
                target.catalog or ""
                if target is not None
                else " | ".join(item.catalogs)
            ),
            target.database or "" if target is not None else "",
            (
                target.schema_name or ""
                if target is not None
                else " | ".join(item.schemas)
            ),
            (
                target.object_name or ""
                if target is not None
                else " | ".join(item.object_names)
            ),
            target.local_target_object_id or "" if target is not None else "",
            " | ".join(item.evidence_ids),
            " | ".join(target.evidence_ids) if target is not None else "",
        )
        for item in report.linked_tables
        for target in _targets_or_placeholder(item.targets)
    )


def _targets_or_placeholder(
    targets: tuple[ReportLineageTarget, ...],
) -> tuple[ReportLineageTarget | None, ...]:
    return targets if targets else (None,)


def _html_table_rows(rows: tuple[tuple[object, ...], ...]) -> str:
    return "".join(
        "<tr>" + "".join(f"<td>{escape(str(value))}</td>" for value in row) + "</tr>"
        for row in rows
    )


def _html_table_headers(headers: tuple[str, ...]) -> str:
    return "".join(f"<th>{escape(value)}</th>" for value in headers)


def _lineage_row_text(headers: tuple[str, ...], row: tuple[object, ...]) -> str:
    return " — ".join(
        f"{header}: {value}" for header, value in zip(headers, row, strict=True)
    )


def _csv_headers(headers: tuple[str, ...]) -> tuple[str, ...]:
    return tuple(
        re.sub(r"[^a-z0-9]+", "_", value.casefold()).strip("_")
        for value in headers
    )


def _append_rows(sheet: Worksheet, rows: tuple[tuple[object, ...], ...]) -> None:
    for row in rows:
        sheet.append(tuple(_spreadsheet_safe(value) for value in row))


def _spreadsheet_safe(value: object) -> object:
    if not isinstance(value, str):
        return value
    cleaned = _ILLEGAL_SPREADSHEET_CHARACTERS.sub("", value)
    stripped = cleaned.lstrip(" \u00a0")
    if cleaned.startswith(_FORMULA_PREFIXES) or stripped.startswith(("=", "+", "-", "@")):
        return f"'{cleaned}"
    return cleaned


def _pdf_table(rows: tuple[tuple[str, str], ...]) -> Table:
    table = Table(rows, colWidths=(2.1 * inch, 1.2 * inch), hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#123047")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CAD4DA")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("PADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    return table


def _append_pdf_section(
    story: list[object], styles: Any, title: str, rows: tuple[str, ...]
) -> None:
    story.append(Paragraph(escape(title), styles["Heading2"]))
    for row in rows:
        story.append(Paragraph(escape(row), styles["BodyText"]))


def _pdf_entity_marker(entity_type: str, identifier: str) -> str:
    labels = {"application": "A", "candidate": "C", "finding": "F"}
    digest = hashlib.sha256(identifier.encode("utf-8")).hexdigest()[:32]
    return f"APA-{labels[entity_type]}:{digest}"


def _csv_tables(report: PortfolioReportModel) -> dict[str, tuple[tuple[object, ...], ...]]:
    return {
        "identity.csv": _identity_csv_rows(report),
        "report.csv": (
            (
                "report_id",
                "analysis_fingerprint",
                "status",
                "partial_watermark",
                "absence_claims_suppressed",
                "application_count",
                "candidate_count",
                "finding_count",
                "inventory_exclusion_count",
                "record_count",
            ),
            (
                report.report_id,
                report.analysis_fingerprint,
                report.status.value,
                report.partial_watermark or "",
                report.absence_claims_suppressed,
                len(report.applications),
                len(report.candidates),
                len(report.portfolio_findings),
                len(report.inventory_exclusions),
                len(_report_record_ids(report)),
            ),
        ),
        "applications.csv": (
            (
                "application_id",
                "application_name",
                "status",
                "evidence_bundle_sha256",
                "interpretation_sha256",
                "artifact_count",
                "object_count",
                "table_count",
                "query_count",
                "datasource_count",
                "connection_count",
                "interaction_count",
                "evidence_count",
                "unresolved_reference_count",
                "summary",
                "business_purpose",
                "major_workflows",
                "capabilities",
                "coverage_notes",
            ),
            *(
                (
                    item.application_id,
                    item.application_name,
                    item.status.value,
                    item.evidence_bundle_sha256,
                    item.interpretation_sha256 or "",
                    item.artifact_count,
                    item.object_count,
                    item.table_count,
                    item.query_count,
                    item.datasource_count,
                    item.connection_count,
                    item.interaction_count,
                    item.evidence_count,
                    item.unresolved_reference_count,
                    item.summary or "",
                    item.business_purpose or "",
                    " | ".join(item.major_workflows),
                    " | ".join(item.capabilities),
                    " | ".join(item.coverage_notes),
                )
                for item in report.applications
            ),
        ),
        "observed_facts.csv": (
            (
                "evidence_id",
                "application_id",
                "artifact_id",
                "artifact_sha256",
                "object_id",
                "origin",
                "fact_type",
                "location",
                "observation",
                "confidence",
            ),
            *(
                (
                    item.evidence_id,
                    item.application_id,
                    item.artifact_id,
                    item.artifact_sha256,
                    item.object_id or "",
                    item.origin.value,
                    item.fact_type,
                    item.location or "",
                    item.observation,
                    item.confidence.value,
                )
                for item in report.observed_facts
            ),
        ),
        "owner_claims.csv": (
            ("claim_id", "application_id", "field", "value", "source"),
            *(
                (item.claim_id, item.application_id, item.field, item.value, item.source)
                for item in report.owner_claims
            ),
        ),
        "application_profiles.csv": (
            (
                "interpretation_id",
                "application_id",
                "source_bundle_sha256",
                "summary",
                "business_purpose",
                "major_workflows",
                "capabilities",
                "modernization_concerns",
                "uncertainties",
            ),
            *(
                (
                    item.interpretation_id,
                    item.application_id,
                    item.source_bundle_sha256,
                    item.summary,
                    item.business_purpose,
                    " | ".join(item.major_workflows),
                    " | ".join(item.capabilities),
                    " | ".join(item.modernization_concerns),
                    " | ".join(item.uncertainties),
                )
                for item in report.application_profiles
            ),
        ),
        "profile_evidence.csv": (
            ("interpretation_id", "evidence_id"),
            *(
                (item.interpretation_id, evidence_id)
                for item in report.application_profiles
                for evidence_id in item.evidence_ids
            ),
        ),
        "profile_claims.csv": (
            ("interpretation_id", "claim_id"),
            *(
                (item.interpretation_id, claim_id)
                for item in report.application_profiles
                for claim_id in item.claim_ids
            ),
        ),
        "reviews.csv": (
            (
                "review_id",
                "application_id",
                "subject_id",
                "status",
                "edited_value",
                "reviewer",
                "notes",
            ),
            *(
                (
                    item.review_id,
                    item.application_id,
                    item.subject_id,
                    item.status.value,
                    item.edited_value or "",
                    item.reviewer or "",
                    item.notes or "",
                )
                for item in report.review_records
            ),
        ),
        "review_evidence.csv": (
            ("review_id", "evidence_id"),
            *(
                (item.review_id, evidence_id)
                for item in report.review_records
                for evidence_id in item.evidence_ids
            ),
        ),
        "review_claims.csv": (
            ("review_id", "claim_id"),
            *(
                (item.review_id, claim_id)
                for item in report.review_records
                for claim_id in item.claim_ids
            ),
        ),
        "unresolved_references.csv": (
            (
                "entry_id",
                "unresolved_id",
                "application_id",
                "source_object_id",
                "reference",
                "reason",
            ),
            *(
                (
                    item.entry_id,
                    item.unresolved.unresolved_id,
                    item.application_id,
                    item.unresolved.source_object_id,
                    item.unresolved.reference,
                    item.unresolved.reason,
                )
                for item in report.unresolved_references
            ),
        ),
        "coverage.csv": (
            (
                "coverage_id",
                "application_id",
                "primary_artifact_count",
                "complete_artifact_count",
                "partial_artifact_count",
                "failed_artifact_count",
                "discovered_object_count",
                "extracted_object_count",
                "warning_count",
                "unresolved_reference_count",
            ),
            *(
                (
                    item.coverage_id,
                    item.application_id,
                    item.coverage.primary_artifact_count,
                    item.coverage.complete_artifact_count,
                    item.coverage.partial_artifact_count,
                    item.coverage.failed_artifact_count,
                    item.coverage.discovered_object_count,
                    item.coverage.extracted_object_count,
                    item.coverage.warning_count,
                    item.coverage.unresolved_reference_count,
                )
                for item in report.coverage
            ),
        ),
        "pass_through_queries.csv": (
            _csv_headers(_PASS_THROUGH_LINEAGE_HEADERS),
            *_pass_through_lineage_rows(report),
        ),
        "linked_tables.csv": (
            _csv_headers(_LINKED_TABLE_LINEAGE_HEADERS),
            *_linked_table_lineage_rows(report),
        ),
        "objects.csv": (
            (
                "entry_id",
                "application_id",
                "object_id",
                "artifact_id",
                "object_type",
                "name",
                "evidence_ids",
            ),
            *(
                (
                    item.entry_id,
                    item.application_id,
                    item.access_object.object_id,
                    item.access_object.artifact_id,
                    item.access_object.object_type.value,
                    item.access_object.name,
                    " | ".join(item.access_object.evidence_ids),
                )
                for item in report.object_registry
            ),
        ),
        "tables.csv": (
            (
                "entry_id",
                "application_id",
                "object_id",
                "is_linked",
                "is_hidden",
                "is_system",
                "source_table_name",
                "connection_id",
                "evidence_ids",
            ),
            *(
                (
                    item.entry_id,
                    item.application_id,
                    item.table.object_id,
                    item.table.is_linked,
                    item.table.is_hidden,
                    item.table.is_system,
                    item.table.source_table_name or "",
                    item.table.connection_id or "",
                    " | ".join(item.table.evidence_ids),
                )
                for item in report.table_registry
            ),
        ),
        "queries.csv": (
            (
                "entry_id",
                "application_id",
                "object_id",
                "query_kind",
                "is_hidden",
                "is_system",
                "connection_kind",
                "connection_id",
                "returns_records",
                "sanitized_sql",
                "evidence_ids",
            ),
            *(
                (
                    item.entry_id,
                    item.application_id,
                    item.query.object_id,
                    item.query.query_kind.value,
                    item.query.is_hidden,
                    item.query.is_system,
                    item.query.connection_kind.value,
                    item.query.connection_id or "",
                    item.query.returns_records
                    if item.query.returns_records is not None
                    else "",
                    item.query.sanitized_sql or "",
                    " | ".join(item.query.evidence_ids),
                )
                for item in report.query_registry
            ),
        ),
        "connections.csv": (
            (
                "connection_id",
                "application_id",
                "object_id",
                "connection_kind",
                "direct_provenance",
                "resolution_status",
                "sanitized_summary",
                "datasource_id",
                "evidence_ids",
            ),
            *(
                (
                    item.connection_id,
                    item.application_id,
                    item.object_id,
                    item.connection_kind.value,
                    item.direct_provenance.value,
                    item.resolution_status.value,
                    item.sanitized_summary,
                    item.datasource_id or "",
                    " | ".join(item.evidence_ids),
                )
                for item in report.connection_registry
            ),
        ),
        "datasources.csv": (
            (
                "datasource_id",
                "platform",
                "driver",
                "dsn",
                "server",
                "database",
                "resource",
            ),
            *(
                (
                    item.datasource_id,
                    item.platform,
                    item.driver or "",
                    item.dsn or "",
                    item.server or "",
                    item.database or "",
                    item.resource or "",
                )
                for item in report.datasource_registry
            ),
        ),
        "data_access.csv": (
            (
                "entry_id",
                "interaction_id",
                "application_id",
                "source_object_id",
                "operation",
                "scope",
                "connection_id",
                "datasource_id",
                "local_target_object_id",
                "catalog",
                "database",
                "schema_name",
                "object_name",
                "evidence_ids",
            ),
            *(
                (
                    item.entry_id,
                    item.data_access.interaction_id,
                    item.application_id,
                    item.data_access.source_object_id,
                    item.data_access.operation.value,
                    item.data_access.scope.value,
                    item.data_access.connection_id or "",
                    item.data_access.datasource_id or "",
                    item.data_access.local_target_object_id or "",
                    item.data_access.catalog or "",
                    item.data_access.database or "",
                    item.data_access.schema_name or "",
                    item.data_access.object_name or "",
                    " | ".join(item.data_access.evidence_ids),
                )
                for item in report.data_access
            ),
        ),
        "dependency_nodes.csv": (
            (
                "node_id",
                "application_id",
                "kind",
                "label",
                "artifact_id",
                "object_id",
                "datasource_id",
            ),
            *(
                (
                    item.node_id,
                    item.application_id,
                    item.kind.value,
                    item.label,
                    item.artifact_id or "",
                    item.object_id or "",
                    item.datasource_id or "",
                )
                for item in report.dependency_nodes
            ),
        ),
        "dependency_edges.csv": (
            (
                "entry_id",
                "edge_id",
                "application_id",
                "source_node_id",
                "target_node_id",
                "relationship",
                "operation",
                "confidence",
                "evidence_ids",
            ),
            *(
                (
                    item.entry_id,
                    item.dependency_edge.edge_id,
                    item.application_id,
                    item.dependency_edge.source_node_id,
                    item.dependency_edge.target_node_id,
                    item.dependency_edge.relationship.value,
                    item.dependency_edge.operation.value,
                    item.dependency_edge.confidence.value,
                    " | ".join(item.dependency_edge.evidence_ids),
                )
                for item in report.dependency_edges
            ),
        ),
        "omissions.csv": (
            (
                "omission_id",
                "application_id",
                "logical_unit_id",
                "stage",
                "reason",
            ),
            *(
                (
                    item.omission_id,
                    item.application_id or "",
                    item.logical_unit_id or "",
                    item.stage.value,
                    item.reason,
                )
                for item in report.omissions
            ),
        ),
        "inventory_exclusions.csv": (
            (
                "exclusion_id",
                "application_id",
                "application_name",
                "filename",
                "source_locator",
                "status",
                "reason",
            ),
            *(
                (
                    item.exclusion_id,
                    item.application_id,
                    item.application_name,
                    item.filename,
                    item.source_locator,
                    item.status,
                    item.reason,
                )
                for item in report.inventory_exclusions
            ),
        ),
        "warnings.csv": (
            ("warning",),
            *((item,) for item in report.warnings),
        ),
        "candidates.csv": (
            ("candidate_id", "candidate_type", "score", "policy_version"),
            *(
                (
                    item.candidate_id,
                    item.candidate_type.value,
                    item.score,
                    item.policy_version,
                )
                for item in report.candidates
            ),
        ),
        "candidate_applications.csv": (
            ("candidate_id", "application_id"),
            *(
                (item.candidate_id, application_id)
                for item in report.candidates
                for application_id in item.application_ids
            ),
        ),
        "candidate_evidence.csv": (
            ("candidate_id", "evidence_id"),
            *(
                (item.candidate_id, evidence_id)
                for item in report.candidates
                for evidence_id in item.evidence_ids
            ),
        ),
        "candidate_basis.csv": (
            ("candidate_id", "basis_id"),
            *(
                (item.candidate_id, basis_id)
                for item in report.candidates
                for basis_id in item.basis_ids
            ),
        ),
        "findings.csv": (
            ("finding_id", "kind", "title", "narrative", "confidence"),
            *(
                (
                    item.finding_id,
                    item.kind.value,
                    item.title,
                    item.narrative,
                    item.confidence.value,
                )
                for item in report.portfolio_findings
            ),
        ),
        "finding_applications.csv": (
            ("finding_id", "application_id"),
            *(
                (item.finding_id, application_id)
                for item in report.portfolio_findings
                for application_id in item.application_ids
            ),
        ),
        "finding_evidence.csv": (
            ("finding_id", "evidence_id"),
            *(
                (item.finding_id, evidence_id)
                for item in report.portfolio_findings
                for evidence_id in item.evidence_ids
            ),
        ),
    }


def _identity_csv_rows(
    report: PortfolioReportModel,
) -> tuple[tuple[object, ...], ...]:
    return (
        ("entity_type", "entity_id"),
        *(("record", identifier) for identifier in _report_record_ids(report)),
    )


def _csv_bytes(rows: tuple[tuple[object, ...], ...]) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.writer(output, lineterminator="\n")
    for row in rows:
        writer.writerow(tuple(_spreadsheet_safe(value) for value in row))
    return output.getvalue().encode("utf-8")


def _scan_text_for_leaks(value: str, source: str) -> None:
    decoded = unescape(value)
    for match in _SECRET_ASSIGNMENT.finditer(decoded):
        assigned_value = match.group("general_value") or match.group("equals_value") or ""
        if REDACTED not in assigned_value.casefold():
            raise ReportLeakError(f"possible credential assignment found in {source}")
    if _RAW_CONNECTION.search(decoded):
        raise ReportLeakError(f"possible raw connection string found in {source}")
    if _URI_USERINFO.search(decoded):
        raise ReportLeakError(f"possible URI credentials found in {source}")


def _scan_zip_for_leaks(payload: bytes, source: str) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            for name in archive.namelist():
                if name.endswith("/"):
                    continue
                raw = archive.read(name)
                _scan_text_for_leaks(raw.decode("utf-8", errors="ignore"), f"{source}:{name}")
    except zipfile.BadZipFile as exc:
        raise ReportLeakError(f"invalid archive while scanning {source}") from exc


def _scan_rendered_artifact(path: Path, report_format: ReportFormat) -> None:
    payload = path.read_bytes()
    if report_format in {ReportFormat.XLSX, ReportFormat.CSV_BUNDLE}:
        _scan_zip_for_leaks(payload, path.name)
    elif report_format == ReportFormat.PDF:
        # Our ReportLab policy disables page compression. Rejoin adjacent PDF text-show
        # operands before scanning so a safe marker rendered as several runs (for example
        # ``user=<redacted>``) is evaluated as the visible text rather than PDF operators.
        decoded = payload.decode("latin-1", errors="ignore")
        decoded = re.sub(r"\)\s*Tj\s+T\*\s*\(", "\n", decoded)
        decoded = re.sub(r"\)\s*Tj\s*\(", "", decoded)
        _scan_text_for_leaks(decoded, path.name)
    else:
        _scan_text_for_leaks(payload.decode("latin-1", errors="ignore"), path.name)


def _verify_artifact_record_ids(
    path: Path,
    report_format: ReportFormat,
    expected_ids: tuple[str, ...],
    *,
    application_ids: tuple[str, ...],
    candidate_ids: tuple[str, ...],
    finding_ids: tuple[str, ...],
) -> None:
    expected = set(expected_ids)
    expected_entities = {
        "application": set(application_ids),
        "candidate": set(candidate_ids),
        "finding": set(finding_ids),
    }
    expected_counts = {
        "applications": len(application_ids),
        "candidates": len(candidate_ids),
        "findings": len(finding_ids),
    }
    try:
        if report_format == ReportFormat.HTML:
            html = path.read_text(encoding="utf-8")
            actual = {
                unescape(item)
                for item in re.findall(r"<li><code>([^<]+)</code></li>", html)
            }
            actual_entities = {
                entity_type: {
                    unescape(item)
                    for item in re.findall(
                        rf'data-entity-type="{entity_type}"\s+'
                        r'data-entity-id="([^"]+)"',
                        html,
                    )
                }
                for entity_type in expected_entities
            }
            actual_counts = {
                name: int(match.group(1))
                for name in expected_counts
                if (
                    match := re.search(
                        rf'data-count="{name}">(\d+)</dd>',
                        html,
                    )
                )
            }
        elif report_format == ReportFormat.XLSX:
            workbook = load_workbook(path, read_only=True, data_only=False)
            try:
                if "Record IDs" not in workbook.sheetnames:
                    raise ReportPublicationError("XLSX report has no identity registry")
                actual = {
                    str(row[0])
                    for row in workbook["Record IDs"].iter_rows(
                        min_row=2, values_only=True
                    )
                    if row and row[0] is not None
                }
                actual_entities = {
                    "application": _xlsx_first_column(workbook, "Applications"),
                    "candidate": _xlsx_first_column(workbook, "Candidates"),
                    "finding": _xlsx_first_column(workbook, "Findings"),
                }
                actual_counts = {
                    str(row[0]).casefold(): int(row[1])
                    for row in workbook["Summary"].iter_rows(
                        min_row=2,
                        values_only=True,
                    )
                    if row
                    and len(row) >= 2
                    and row[0] in {"Applications", "Candidates", "Findings"}
                    and row[1] is not None
                }
            finally:
                workbook.close()
            expected = {str(_spreadsheet_safe(item)) for item in expected}
            expected_entities = {
                name: {str(_spreadsheet_safe(item)) for item in values}
                for name, values in expected_entities.items()
            }
        elif report_format == ReportFormat.CSV_BUNDLE:
            with zipfile.ZipFile(path) as archive:
                rows = csv.reader(
                    io.StringIO(archive.read("identity.csv").decode("utf-8"))
                )
                header = next(rows, None)
                if header != ["entity_type", "entity_id"]:
                    raise ReportPublicationError(
                        "CSV report has no canonical identity registry"
                    )
                actual = {row[1] for row in rows if len(row) == 2 and row[0] == "record"}
                actual_entities = {
                    "application": _csv_first_column(archive, "applications.csv"),
                    "candidate": _csv_first_column(archive, "candidates.csv"),
                    "finding": _csv_first_column(archive, "findings.csv"),
                }
                report_rows = list(
                    csv.DictReader(
                        io.StringIO(archive.read("report.csv").decode("utf-8"))
                    )
                )
                if len(report_rows) != 1:
                    raise ReportPublicationError(
                        "CSV report has no singular canonical report row"
                    )
                actual_counts = {
                    "applications": int(report_rows[0]["application_count"]),
                    "candidates": int(report_rows[0]["candidate_count"]),
                    "findings": int(report_rows[0]["finding_count"]),
                }
            expected = {str(_spreadsheet_safe(item)) for item in expected}
            expected_entities = {
                name: {str(_spreadsheet_safe(item)) for item in values}
                for name, values in expected_entities.items()
            }
        else:
            payload = path.read_bytes()
            actual = {
                identifier
                for identifier in expected
                if identifier.encode("ascii") in payload
            }
            text = payload.decode("latin-1")
            actual_entities = {
                entity_type: {
                    item
                    for item in re.findall(rf"APA-{label}:([0-9a-f]{{32}})", text)
                }
                for entity_type, label in {
                    "application": "A",
                    "candidate": "C",
                    "finding": "F",
                }.items()
            }
            expected_entities = {
                entity_type: {
                    hashlib.sha256(item.encode("utf-8")).hexdigest()[:32]
                    for item in values
                }
                for entity_type, values in expected_entities.items()
            }
            count_match = re.search(
                r"APA-N:(\d+):(\d+):(\d+)",
                text,
            )
            actual_counts = (
                {
                    "applications": int(count_match.group(1)),
                    "candidates": int(count_match.group(2)),
                    "findings": int(count_match.group(3)),
                }
                if count_match is not None
                else {}
            )
    except (OSError, UnicodeError, ValueError, KeyError, zipfile.BadZipFile) as exc:
        raise ReportPublicationError(
            f"could not verify {report_format.value} report identities"
        ) from exc
    if actual != expected:
        raise ReportPublicationError(
            f"{report_format.value} report does not preserve normalized record identities"
        )
    if actual_entities != expected_entities or actual_counts != expected_counts:
        raise ReportPublicationError(
            f"{report_format.value} report has inconsistent application/candidate/finding "
            "identities or counts"
        )


def _xlsx_first_column(workbook: Any, sheet_name: str) -> set[str]:
    if sheet_name not in workbook.sheetnames:
        raise ReportPublicationError(f"XLSX report has no {sheet_name} table")
    return {
        str(row[0])
        for row in workbook[sheet_name].iter_rows(min_row=2, values_only=True)
        if row and row[0] is not None
    }


def _csv_first_column(archive: zipfile.ZipFile, name: str) -> set[str]:
    rows = csv.reader(io.StringIO(archive.read(name).decode("utf-8")))
    if next(rows, None) is None:
        raise ReportPublicationError(f"CSV report table is empty: {name}")
    return {row[0] for row in rows if row}


def _report_record_ids(report: PortfolioReportModel) -> tuple[str, ...]:
    return tuple(
        sorted(
            {
                report.report_id,
                *(item.application_id for item in report.applications),
                *(item.candidate_id for item in report.candidates),
                *(item.finding_id for item in report.portfolio_findings),
                *(item.evidence_id for item in report.observed_facts),
                *(item.claim_id for item in report.owner_claims),
                *(item.interpretation_id for item in report.application_profiles),
                *(item.review_id for item in report.review_records),
                *(item.entry_id for item in report.unresolved_references),
                *(item.unresolved.unresolved_id for item in report.unresolved_references),
                *(item.coverage_id for item in report.coverage),
                *(item.view_id for item in report.pass_through_queries),
                *(
                    target.target_id
                    for item in report.pass_through_queries
                    for target in item.targets
                ),
                *(item.view_id for item in report.linked_tables),
                *(
                    target.target_id
                    for item in report.linked_tables
                    for target in item.targets
                ),
                *(item.entry_id for item in report.object_registry),
                *(item.access_object.object_id for item in report.object_registry),
                *(item.entry_id for item in report.table_registry),
                *(item.entry_id for item in report.query_registry),
                *(item.connection_id for item in report.connection_registry),
                *(item.datasource_id for item in report.datasource_registry),
                *(item.entry_id for item in report.data_access),
                *(item.data_access.interaction_id for item in report.data_access),
                *(item.node_id for item in report.dependency_nodes),
                *(item.entry_id for item in report.dependency_edges),
                *(item.dependency_edge.edge_id for item in report.dependency_edges),
                *(item.exclusion_id for item in report.inventory_exclusions),
                *(item.omission_id for item in report.omissions),
            }
        )
    )


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _write_bytes_synced(path: Path, payload: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("wb") as handle:
        handle.write(payload)
        handle.flush()
        os.fsync(handle.fileno())


def _atomic_write_bytes(path: Path, payload: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    file_descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{path.name}.", dir=path.parent
    )
    temporary_path = Path(temporary_name)
    try:
        with os.fdopen(file_descriptor, "wb") as handle:
            handle.write(payload)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary_path, path)
        _fsync_directory(path.parent)
    finally:
        temporary_path.unlink(missing_ok=True)


def _fsync_file(path: Path) -> None:
    with path.open("rb") as handle:
        os.fsync(handle.fileno())


def _fsync_directory(path: Path) -> None:
    file_descriptor = os.open(path, os.O_RDONLY)
    try:
        os.fsync(file_descriptor)
    finally:
        os.close(file_descriptor)


def _contained_report_path(root: Path, relative_path: str) -> Path:
    candidate = PurePosixPath(relative_path)
    if candidate.is_absolute() or ".." in candidate.parts:
        raise ReportPublicationError("report artifact path escapes its run")
    resolved_root = root.resolve()
    resolved = (resolved_root / Path(*candidate.parts)).resolve()
    if not resolved.is_relative_to(resolved_root):
        raise ReportPublicationError("report artifact path escapes its run")
    return resolved
