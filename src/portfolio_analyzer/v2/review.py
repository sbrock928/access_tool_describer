"""Strict human-review overlay import for immutable V2 analyses."""

from __future__ import annotations

from datetime import UTC, datetime
from enum import StrEnum
from pathlib import Path
from typing import Literal, Self

from openpyxl import load_workbook
from pydantic import field_validator, model_validator

from portfolio_analyzer.v2.identity import stable_id
from portfolio_analyzer.v2.models import NonEmptyString, Sha256, StrictModel

REVIEW_OVERLAY_SCHEMA_VERSION = "review-overlay-v2"
REVIEW_SHEET = "Review Queue"
REVIEW_HEADERS = (
    "Analysis Fingerprint",
    "Proposal ID",
    "Evidence IDs",
    "Decision",
    "Edited Value",
    "Reviewer",
    "Notes",
)


class ReviewDecisionKind(StrEnum):
    ACCEPT = "accept"
    EDIT = "edit"
    REJECT = "reject"


class ReviewedProposalDecision(StrictModel):
    proposal_id: NonEmptyString
    evidence_ids: tuple[NonEmptyString, ...]
    decision: ReviewDecisionKind
    edited_value: str | None = None
    reviewer: str | None = None
    notes: str | None = None
    reviewed_at: datetime
    decision_id: str = ""

    @field_validator("reviewed_at")
    @classmethod
    def require_aware_timestamp(cls, value: datetime) -> datetime:
        if value.tzinfo is None or value.utcoffset() is None:
            raise ValueError("reviewed_at must include a timezone")
        return value

    @model_validator(mode="after")
    def validate_decision(self) -> Self:
        evidence_ids = tuple(sorted(set(self.evidence_ids)))
        object.__setattr__(self, "evidence_ids", evidence_ids)
        if not evidence_ids:
            raise ValueError("review decisions must retain proposal evidence identities")
        if self.decision == ReviewDecisionKind.EDIT and not self.edited_value:
            raise ValueError("edited decisions require an edited value")
        if self.decision != ReviewDecisionKind.EDIT and self.edited_value is not None:
            raise ValueError("only edited decisions may contain an edited value")
        expected = stable_id(
            "review_decision",
            self.proposal_id,
            evidence_ids,
            self.decision,
            self.edited_value,
        )
        if self.decision_id and self.decision_id != expected:
            raise ValueError("decision_id does not match stable proposal evidence")
        object.__setattr__(self, "decision_id", expected)
        return self


class ReviewOverlay(StrictModel):
    schema_version: Literal["review-overlay-v2"] = "review-overlay-v2"
    analysis_fingerprint: Sha256
    imported_at: datetime
    decisions: tuple[ReviewedProposalDecision, ...]
    overlay_id: str = ""

    @field_validator("imported_at")
    @classmethod
    def require_aware_timestamp(cls, value: datetime) -> datetime:
        if value.tzinfo is None or value.utcoffset() is None:
            raise ValueError("imported_at must include a timezone")
        return value

    @model_validator(mode="after")
    def validate_overlay(self) -> Self:
        by_proposal = {item.proposal_id: item for item in self.decisions}
        if len(by_proposal) != len(self.decisions):
            raise ValueError("review overlay contains duplicate proposal decisions")
        decisions = tuple(sorted(self.decisions, key=lambda item: item.proposal_id))
        object.__setattr__(self, "decisions", decisions)
        expected = stable_id(
            "review_overlay",
            self.analysis_fingerprint,
            tuple((item.proposal_id, item.evidence_ids, item.decision_id) for item in decisions),
        )
        if self.overlay_id and self.overlay_id != expected:
            raise ValueError("overlay_id does not match the review contents")
        object.__setattr__(self, "overlay_id", expected)
        return self


def import_review_overlay(
    workbook_path: Path,
    *,
    analysis_fingerprint: str,
    known_proposals: dict[str, tuple[str, ...]],
    imported_at: datetime | None = None,
) -> ReviewOverlay:
    """Import decisions only when analysis, proposal, and evidence identities still match."""

    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    if REVIEW_SHEET not in workbook.sheetnames:
        raise ValueError(f"Workbook has no '{REVIEW_SHEET}' sheet")
    rows = workbook[REVIEW_SHEET].iter_rows(values_only=True)
    try:
        headers = tuple(_plain_cell(value, row_number=1) for value in next(rows))
    except StopIteration as exc:
        raise ValueError("Review Queue is empty") from exc
    missing = set(REVIEW_HEADERS) - set(headers)
    if missing:
        raise ValueError(f"Review Queue is missing columns: {', '.join(sorted(missing))}")
    positions = {header: headers.index(header) for header in REVIEW_HEADERS}
    decisions: list[ReviewedProposalDecision] = []
    seen: set[str] = set()
    timestamp = imported_at or datetime.now(UTC)
    for row_number, row in enumerate(rows, start=2):
        proposal_id = _row_cell(row, positions["Proposal ID"], row_number)
        choice = _row_cell(row, positions["Decision"], row_number).casefold()
        if not proposal_id and not choice:
            continue
        if not proposal_id:
            raise ValueError(f"Review Queue row {row_number} is incomplete")
        if proposal_id in seen:
            raise ValueError(f"Duplicate review proposal '{proposal_id}'")
        seen.add(proposal_id)
        row_fingerprint = _row_cell(
            row, positions["Analysis Fingerprint"], row_number
        ).casefold()
        if row_fingerprint != analysis_fingerprint.casefold():
            raise ValueError(
                f"Review Queue row {row_number} is stale for the current analysis fingerprint"
            )
        expected_evidence = known_proposals.get(proposal_id)
        if expected_evidence is None:
            raise ValueError(f"Unknown proposal ID in Review Queue row {row_number}")
        evidence_ids = tuple(
            sorted(
                {
                    item.strip()
                    for item in _row_cell(
                        row, positions["Evidence IDs"], row_number
                    ).split(",")
                    if item.strip()
                }
            )
        )
        if evidence_ids != tuple(sorted(set(expected_evidence))):
            raise ValueError(
                f"Review Queue row {row_number} has stale or conflicting evidence identities"
            )
        if not choice:
            continue
        try:
            decision = ReviewDecisionKind(choice)
        except ValueError as exc:
            raise ValueError(
                f"Invalid decision in Review Queue row {row_number}: {choice}"
            ) from exc
        edited_value = _row_cell(row, positions["Edited Value"], row_number) or None
        decisions.append(
            ReviewedProposalDecision(
                proposal_id=proposal_id,
                evidence_ids=evidence_ids,
                decision=decision,
                edited_value=edited_value,
                reviewer=_row_cell(row, positions["Reviewer"], row_number) or None,
                notes=_row_cell(row, positions["Notes"], row_number) or None,
                reviewed_at=timestamp,
            )
        )
    if not decisions:
        raise ValueError("Review Queue contains no decisions")
    return ReviewOverlay(
        analysis_fingerprint=analysis_fingerprint,
        imported_at=timestamp,
        decisions=tuple(decisions),
    )


def carry_forward_decisions(
    overlay: ReviewOverlay,
    *,
    known_proposals: dict[str, tuple[str, ...]],
) -> tuple[ReviewedProposalDecision, ...]:
    """Carry decisions only while stable proposal and evidence identities remain identical."""

    return tuple(
        item
        for item in overlay.decisions
        if tuple(sorted(set(known_proposals.get(item.proposal_id, ())))) == item.evidence_ids
    )


def _row_cell(row: tuple[object, ...], index: int, row_number: int) -> str:
    value = row[index] if index < len(row) else None
    return _plain_cell(value, row_number=row_number)


def _plain_cell(value: object, *, row_number: int) -> str:
    if isinstance(value, str) and value.lstrip(" \t\r\n\u00a0").startswith(
        ("=", "+", "-", "@")
    ):
        raise ValueError(f"Review Queue row {row_number} contains a formula-like value")
    return str(value or "").strip()
