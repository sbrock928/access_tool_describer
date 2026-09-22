"""Read inventory workbooks without interpreting their descriptions as fact."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from portfolio_analyzer.models import InventoryRecord
from portfolio_analyzer.naming import euc_directory_name

REQUIRED_COLUMNS = ("INVENTORY_ID", "EUCTNAME", "FILE_NAME", "FULLPATH", "DESCRIPTION")


class InventoryValidationError(ValueError):
    pass


def load_inventory(path: Path) -> list[InventoryRecord]:
    """Load and preserve workbook values; reject missing required headings."""
    workbook = load_workbook(path, read_only=True, data_only=False)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    try:
        headers = tuple("" if value is None else str(value).strip() for value in next(rows))
    except StopIteration as exc:
        raise InventoryValidationError("Inventory workbook has no header row") from exc
    missing = [name for name in REQUIRED_COLUMNS if name not in headers]
    if missing:
        raise InventoryValidationError(f"Inventory missing required columns: {', '.join(missing)}")
    records: list[InventoryRecord] = []
    for row_number, row in enumerate(rows, start=2):
        values = {
            headers[index]: row[index] if index < len(row) else None
            for index in range(len(headers))
        }
        if all(value is None for value in values.values()):
            continue
        tool_id = values["INVENTORY_ID"]
        tool_name = values["EUCTNAME"]
        filename = values["FILE_NAME"]
        filepath = values["FULLPATH"]
        if any(value in (None, "") for value in (tool_id, tool_name, filename, filepath)):
            raise InventoryValidationError(
                f"Row {row_number} lacks INVENTORY_ID, EUCTNAME, FILE_NAME, or FULLPATH"
            )
        records.append(
            InventoryRecord(
                tool_inventory_id=str(tool_id).strip(),
                tool_name=str(tool_name).strip(),
                inventory_filename=str(filename).strip(),
                stated_description=_optional_string(values["DESCRIPTION"]),
                filepath=Path(str(filepath).strip()),
                original_values=values,
            )
        )
    _validate_unique_identifiers_and_names(records)
    return records


def _optional_string(value: object) -> str | None:
    if value is None or str(value).strip() == "":
        return None
    return str(value)


def _validate_unique_identifiers_and_names(records: list[InventoryRecord]) -> None:
    seen_ids: dict[str, str] = {}
    seen_directories: dict[str, str] = {}
    for record in records:
        prior_name = seen_ids.get(record.tool_inventory_id)
        if prior_name is not None:
            raise InventoryValidationError(
                f"Duplicate INVENTORY_ID '{record.tool_inventory_id}' for "
                f"'{prior_name}' and '{record.tool_name}'"
            )
        seen_ids[record.tool_inventory_id] = record.tool_name

        directory = euc_directory_name(record.tool_name)
        prior_directory_name = seen_directories.get(directory.casefold())
        if prior_directory_name is not None:
            raise InventoryValidationError(
                "EUCTNAME values must produce unique workspace folders; "
                f"'{prior_directory_name}' and '{record.tool_name}' both map to '{directory}'"
            )
        seen_directories[directory.casefold()] = record.tool_name
